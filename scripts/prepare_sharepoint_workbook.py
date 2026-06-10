#!/usr/bin/env python3
"""Create a SharePoint/Graph-friendly copy of the customer workbook.

This script does not upload anything. It converts the local workbook copy into
named Excel tables with stable row metadata so the desktop app can use Graph
table row APIs reliably after the file is uploaded to SharePoint.
"""

from __future__ import annotations

import argparse
import shutil
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sharepoint_excel_sync import (
    COMPLETED_SHEET,
    MASTER_SHEET,
    OIL_GAS_SHEET,
    PROJECT_COLUMNS,
    PROJECT_TABLES,
)


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def normalize_headers(ws) -> list[str]:
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_values = [str(header).strip() if header else "" for header in headers]
    for column_name in PROJECT_COLUMNS:
        if column_name not in header_values:
            try:
                blank_idx = header_values.index("")
            except ValueError:
                blank_idx = -1
            if blank_idx >= 0:
                ws.cell(row=1, column=blank_idx + 1, value=column_name)
                header_values[blank_idx] = column_name
            else:
                ws.cell(row=1, column=ws.max_column + 1, value=column_name)
                header_values.append(column_name)
    for idx, header in enumerate(header_values, start=1):
        if not header:
            replacement = f"Unused Column {idx}"
            ws.cell(row=1, column=idx, value=replacement)
            header_values[idx - 1] = replacement
    return header_values


def ensure_metadata(ws, modified_by: str) -> None:
    headers = normalize_headers(ws)
    col_map = {header: idx + 1 for idx, header in enumerate(headers) if header}
    timestamp = utc_now()
    for row_idx in range(2, ws.max_row + 1):
        if not any(ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, ws.max_column + 1)):
            continue
        if not ws.cell(row=row_idx, column=col_map["Record ID"]).value:
            ws.cell(row=row_idx, column=col_map["Record ID"], value=uuid.uuid4().hex)
        if not ws.cell(row=row_idx, column=col_map["Last Modified UTC"]).value:
            ws.cell(row=row_idx, column=col_map["Last Modified UTC"], value=timestamp)
        if not ws.cell(row=row_idx, column=col_map["Last Modified By"]).value:
            ws.cell(row=row_idx, column=col_map["Last Modified By"], value=modified_by)


def remove_existing_table(ws, table_name: str) -> None:
    if table_name in ws.tables:
        del ws.tables[table_name]


def ensure_table(ws, table_name: str) -> None:
    remove_existing_table(ws, table_name)
    ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{max(ws.max_row, 2)}"
    table = Table(displayName=table_name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)


def prepare(source: Path, output: Path, modified_by: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if source.resolve() != output.resolve():
        shutil.copy2(source, output)
    wb = openpyxl.load_workbook(output)
    try:
        for sheet_name, table_name in PROJECT_TABLES.items():
            if sheet_name not in wb.sheetnames:
                if sheet_name == COMPLETED_SHEET:
                    wb.create_sheet(sheet_name)
                else:
                    raise KeyError(f"Workbook is missing required sheet: {sheet_name}")
            ws = wb[sheet_name]
            if ws.max_row == 1 and not any(cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))):
                for col_idx, header in enumerate(PROJECT_COLUMNS, start=1):
                    ws.cell(row=1, column=col_idx, value=header)
            ensure_metadata(ws, modified_by)
            ensure_table(ws, table_name)
        wb.save(output)
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare Customer Information.xlsx for SharePoint Graph sync.")
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("Reference Data/XLSX Workbooks/Customer Information.xlsx"),
        help="Path to the existing local workbook.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("Reference Data/XLSX Workbooks/Customer Information.sharepoint-ready.xlsx"),
        help="Path for the prepared workbook copy.",
    )
    parser.add_argument("--modified-by", default="Workbook preparation script")
    args = parser.parse_args()
    prepare(args.source, args.output, args.modified_by)
    print(f"Prepared workbook: {args.output}")
    print(f"Tables: {', '.join(PROJECT_TABLES.values())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
