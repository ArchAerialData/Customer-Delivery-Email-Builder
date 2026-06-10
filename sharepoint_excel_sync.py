import base64
import json
import os
import sys
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl


APP_NAME = "Email Builder"
CONFIG_ENV_VAR = "EMAIL_BUILDER_SHAREPOINT_CONFIG"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"
DEFAULT_SCOPES = ["User.Read", "Files.ReadWrite", "offline_access"]

MASTER_SHEET = "Construction"
LEGACY_MASTER_SHEET = "ABCD"
OIL_GAS_SHEET = "Oil & Gas"
COMPLETED_SHEET = "Completed Projects"

PROJECT_TABLES = {
    MASTER_SHEET: "ConstructionProjects",
    OIL_GAS_SHEET: "OilGasProjects",
    COMPLETED_SHEET: "CompletedProjects",
}

PROJECT_COLUMNS = [
    "Record ID",
    "Client",
    "Project Name",
    "Dropbox URLs",
    "Operations Contact Email",
    "Procore",
    "Internal URLs",
    "Last Modified UTC",
    "Last Modified By",
]

APP_TO_EXCEL = {
    "record_id": "Record ID",
    "client": "Client",
    "site": "Project Name",
    "link": "Dropbox URLs",
    "email": "Operations Contact Email",
    "procore": "Procore",
    "internal_url": "Internal URLs",
    "modified_utc": "Last Modified UTC",
    "modified_by": "Last Modified By",
}
EXCEL_TO_APP = {value: key for key, value in APP_TO_EXCEL.items()}

LOCAL_HEADER_ALIASES = {
    "client": "Client",
    "project name": "Project Name",
    "dropbox urls": "Dropbox URLs",
    "operations contact email": "Operations Contact Email",
    "procore": "Procore",
    "internal urls": "Internal URLs",
}


class SharePointSyncError(RuntimeError):
    """Base error for SharePoint workbook sync failures."""


class WorkbookFormatError(SharePointSyncError):
    """Raised when the SharePoint workbook is missing required tables/columns."""


class SharePointDependencyError(SharePointSyncError):
    """Raised when optional Graph dependencies are not installed."""


class ProjectConflictError(SharePointSyncError):
    def __init__(self, message: str, remote_row: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.remote_row = remote_row or {}


class GraphApiError(SharePointSyncError):
    def __init__(
        self,
        message: str,
        *,
        status_code: int | None = None,
        code: str | None = None,
        retry_after: int | None = None,
        details: Any = None,
    ) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.code = code
        self.retry_after = retry_after
        self.details = details


def _utc_timestamp() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _sheet_aliases(sheet_name: str) -> list[str]:
    if sheet_name == MASTER_SHEET:
        return [MASTER_SHEET, LEGACY_MASTER_SHEET]
    return [sheet_name]


def _normalizer(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _config_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "on"}
    return bool(value)


def _master_cell_value(row: tuple[Any, ...], idx: int | None) -> str:
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    return _normalizer(row[idx])


def encode_sharing_url(sharing_url: str) -> str:
    encoded = base64.urlsafe_b64encode(sharing_url.encode("utf-8")).decode("ascii").rstrip("=")
    return f"u!{encoded}"


def default_config_path() -> Path:
    override = os.environ.get(CONFIG_ENV_VAR)
    if override:
        return Path(override).expanduser()
    if os.name == "nt":
        base = Path(os.environ.get("APPDATA") or Path.home() / "AppData" / "Roaming")
        return base / APP_NAME / "sharepoint_config.json"
    if sys.platform == "darwin":
        return Path.home() / "Library" / "Application Support" / APP_NAME / "sharepoint_config.json"
    return Path(os.environ.get("XDG_CONFIG_HOME") or Path.home() / ".config") / "email-builder" / "sharepoint_config.json"


def default_token_cache_path(config_path: Path | None = None) -> Path:
    path = config_path or default_config_path()
    return path.with_name(".sharepoint_token_cache.json")


@dataclass
class SharePointWorkbookConfig:
    enabled: bool = False
    tenant_id: str = ""
    client_id: str = ""
    sharing_url: str = ""
    scopes: list[str] = field(default_factory=lambda: list(DEFAULT_SCOPES))
    config_path: Path = field(default_factory=default_config_path)
    token_cache_path: Path | None = None

    @property
    def authority(self) -> str:
        return f"https://login.microsoftonline.com/{self.tenant_id}"

    @property
    def is_ready(self) -> bool:
        return bool(self.enabled and self.tenant_id and self.client_id and self.sharing_url)

    @classmethod
    def load(cls, path: Path | None = None) -> "SharePointWorkbookConfig":
        config_path = path or default_config_path()
        if not config_path.exists():
            return cls(enabled=False, config_path=config_path, token_cache_path=default_token_cache_path(config_path))
        data = json.loads(config_path.read_text(encoding="utf-8"))
        scopes = data.get("scopes") or list(DEFAULT_SCOPES)
        if not isinstance(scopes, list):
            scopes = list(DEFAULT_SCOPES)
        token_cache = data.get("token_cache_path")
        return cls(
            enabled=_config_bool(data.get("enabled", False)),
            tenant_id=str(data.get("tenant_id") or "").strip(),
            client_id=str(data.get("client_id") or "").strip(),
            sharing_url=str(data.get("sharing_url") or "").strip(),
            scopes=[str(scope).strip() for scope in scopes if str(scope).strip()],
            config_path=config_path,
            token_cache_path=Path(token_cache).expanduser() if token_cache else default_token_cache_path(config_path),
        )

    def validate(self) -> None:
        if not self.enabled:
            return
        missing = []
        for key in ("tenant_id", "client_id", "sharing_url"):
            if not getattr(self, key):
                missing.append(key)
        if missing:
            joined = ", ".join(missing)
            raise SharePointSyncError(f"SharePoint sync is enabled but config is missing: {joined}")


class GraphTokenProvider:
    def __init__(self, config: SharePointWorkbookConfig) -> None:
        config.validate()
        try:
            import msal  # type: ignore
        except Exception as exc:  # pragma: no cover - dependency absent in local smoke environments
            raise SharePointDependencyError("Install msal to use SharePoint sync.") from exc
        self._msal = msal
        self._config = config
        self._cache_path = config.token_cache_path or default_token_cache_path(config.config_path)
        self._cache = msal.SerializableTokenCache()
        if self._cache_path.exists():
            try:
                self._cache.deserialize(self._cache_path.read_text(encoding="utf-8"))
            except Exception:
                self._cache = msal.SerializableTokenCache()
        self._app = msal.PublicClientApplication(
            client_id=config.client_id,
            authority=config.authority,
            token_cache=self._cache,
        )

    def _save_cache(self) -> None:
        if self._cache.has_state_changed:
            self._cache_path.parent.mkdir(parents=True, exist_ok=True)
            self._cache_path.write_text(self._cache.serialize(), encoding="utf-8")

    def acquire_token(self, *, interactive: bool = True) -> dict[str, Any]:
        accounts = self._app.get_accounts()
        result = None
        if accounts:
            result = self._app.acquire_token_silent(self._config.scopes, account=accounts[0])
        if not result and interactive:
            result = self._app.acquire_token_interactive(scopes=self._config.scopes, prompt="select_account")
        if not result or "access_token" not in result:
            detail = (result or {}).get("error_description") or (result or {}).get("error") or "No token returned."
            raise SharePointSyncError(f"Could not acquire Microsoft Graph token: {detail}")
        self._save_cache()
        return result

    def clear_cache(self) -> None:
        self._cache = self._msal.SerializableTokenCache()
        self._app.token_cache = self._cache
        self._cache_path.unlink(missing_ok=True)


class GraphWorkbookClient:
    def __init__(self, config: SharePointWorkbookConfig) -> None:
        try:
            import requests  # type: ignore
        except Exception as exc:  # pragma: no cover - dependency absent in local smoke environments
            raise SharePointDependencyError("Install requests to use SharePoint sync.") from exc
        self._requests = requests
        self._config = config
        self._token_provider = GraphTokenProvider(config)
        self._drive_item: dict[str, Any] | None = None
        self._session_id: str | None = None
        self._me: dict[str, Any] | None = None

    @property
    def drive_item(self) -> dict[str, Any] | None:
        return self._drive_item

    @property
    def web_url(self) -> str:
        item = self._drive_item or {}
        return str(item.get("webUrl") or self._config.sharing_url or "")

    def _headers(self, *, workbook_session: bool = False) -> dict[str, str]:
        token = self._token_provider.acquire_token()
        headers = {
            "Authorization": f"Bearer {token['access_token']}",
            "Accept": "application/json",
        }
        if workbook_session and self._session_id:
            headers["Workbook-Session-Id"] = self._session_id
        return headers

    def _request(
        self,
        method: str,
        url_or_path: str,
        *,
        json_body: Any = None,
        params: dict[str, Any] | None = None,
        workbook_session: bool = False,
        expected: tuple[int, ...] = (200,),
        retry: bool = True,
    ) -> Any:
        url = url_or_path if url_or_path.startswith("http") else f"{GRAPH_ROOT}{url_or_path}"
        attempts = 3 if retry and workbook_session else 2 if retry else 1
        session_refreshed = False
        for attempt in range(attempts):
            headers = self._headers(workbook_session=workbook_session)
            if json_body is not None:
                headers["Content-Type"] = "application/json"
            response = self._requests.request(
                method,
                url,
                headers=headers,
                json=json_body,
                params=params,
                timeout=45,
            )
            if response.status_code in expected:
                if response.status_code == 204 or not response.content:
                    return None
                return response.json()
            retry_after = _retry_after(response)
            if workbook_session and response.status_code == 404 and self._session_id and not session_refreshed:
                session_refreshed = True
                self._session_id = None
                self.create_session(persist_changes=True)
                continue
            if retry and attempt < attempts - 1 and response.status_code in {429, 503, 504}:
                time.sleep(retry_after if retry_after is not None else 2)
                continue
            raise _graph_error(response)
        raise GraphApiError("Graph request failed without a response.")

    def resolve_drive_item(self) -> dict[str, Any]:
        if self._drive_item:
            return self._drive_item
        encoded = encode_sharing_url(self._config.sharing_url)
        self._drive_item = self._request("GET", f"/shares/{encoded}/driveItem")
        return self._drive_item

    def create_session(self, *, persist_changes: bool = True) -> str:
        self.resolve_drive_item()
        data = self._request(
            "POST",
            self._workbook_path("/createSession"),
            json_body={"persistChanges": persist_changes},
            expected=(201,),
            retry=True,
        )
        self._session_id = str(data.get("id") or "")
        if not self._session_id:
            raise SharePointSyncError("Microsoft Graph did not return a workbook session ID.")
        return self._session_id

    def ensure_session(self) -> None:
        if not self._session_id:
            self.create_session(persist_changes=True)

    def _workbook_path(self, suffix: str = "") -> str:
        item = self.resolve_drive_item()
        parent = item.get("parentReference") or {}
        drive_id = parent.get("driveId")
        item_id = item.get("id")
        if not drive_id or not item_id:
            raise SharePointSyncError("Resolved SharePoint item is missing drive/item identifiers.")
        return f"/drives/{drive_id}/items/{item_id}/workbook{suffix}"

    def _table_path(self, table_name: str, suffix: str = "") -> str:
        encoded_table = quote(table_name, safe="")
        return self._workbook_path(f"/tables/{encoded_table}{suffix}")

    def get_current_user(self) -> str:
        if not self._me:
            self._me = self._request("GET", "/me", params={"$select": "displayName,userPrincipalName,mail"})
        return str(self._me.get("displayName") or self._me.get("mail") or self._me.get("userPrincipalName") or "")

    def list_worksheets(self) -> list[dict[str, Any]]:
        self.ensure_session()
        return self._paged_values(self._workbook_path("/worksheets"))

    def list_tables(self) -> list[dict[str, Any]]:
        self.ensure_session()
        return self._paged_values(self._workbook_path("/tables"))

    def list_table_columns(self, table_name: str) -> list[dict[str, Any]]:
        self.ensure_session()
        return self._paged_values(self._table_path(table_name, "/columns"))

    def list_table_rows(self, table_name: str) -> list[dict[str, Any]]:
        self.ensure_session()
        return self._paged_values(self._table_path(table_name, "/rows"), params={"$top": 500}, page_size=500)

    def _paged_values(
        self,
        path: str,
        *,
        params: dict[str, Any] | None = None,
        page_size: int | None = None,
    ) -> list[dict[str, Any]]:
        values: list[dict[str, Any]] = []
        next_url: str | None = path
        next_params = dict(params or {})
        while next_url:
            data = self._request("GET", next_url, params=next_params or None, workbook_session=True)
            page_values = data.get("value") or []
            values.extend(page_values)
            next_link = data.get("@odata.nextLink")
            if next_link:
                next_url = next_link
                next_params = {}
            elif page_size and len(page_values) == page_size:
                skip = int(next_params.get("$skip") or 0) + page_size
                next_url = path
                next_params = dict(params or {})
                next_params["$skip"] = skip
            else:
                next_url = None
        return values

    def add_table_row(self, table_name: str, values: list[Any]) -> dict[str, Any]:
        self.ensure_session()
        return self._request(
            "POST",
            self._table_path(table_name, "/rows"),
            json_body={"values": [values]},
            workbook_session=True,
            expected=(201, 200),
        )

    def update_table_row(self, table_name: str, index: int, values: list[Any]) -> dict[str, Any]:
        self.ensure_session()
        return self._request(
            "PATCH",
            self._table_path(table_name, f"/rows/{index}"),
            json_body={"values": [values]},
            workbook_session=True,
            expected=(200,),
        )

    def delete_table_row(self, table_name: str, index: int) -> None:
        self.ensure_session()
        self._request(
            "DELETE",
            self._table_path(table_name, f"/rows/{index}"),
            workbook_session=True,
            expected=(204,),
        )

    def sort_table(self, table_name: str, key_indexes: list[int]) -> None:
        fields = [{"key": key, "ascending": True} for key in key_indexes]
        self._request(
            "POST",
            self._table_path(table_name, "/sort/apply"),
            json_body={"fields": fields},
            workbook_session=True,
            expected=(204,),
        )

    def clear_token_cache(self) -> None:
        self._token_provider.clear_cache()


class SharePointProjectDataSource:
    is_remote = True

    def __init__(self, config: SharePointWorkbookConfig, client: GraphWorkbookClient | None = None) -> None:
        self.config = config
        self.client = client or GraphWorkbookClient(config)
        self._column_cache: dict[str, list[str]] = {}

    @property
    def description(self) -> str:
        return "SharePoint workbook"

    @property
    def web_url(self) -> str:
        return self.client.web_url

    def ensure_schema(self) -> None:
        self.validate_workbook()

    def validate_workbook(self) -> None:
        tables = {str(table.get("name")) for table in self.client.list_tables()}
        missing_tables = [table for table in PROJECT_TABLES.values() if table not in tables]
        if missing_tables:
            joined = ", ".join(missing_tables)
            raise WorkbookFormatError(f"SharePoint workbook is missing required Excel table(s): {joined}")
        missing_columns: dict[str, list[str]] = {}
        for table_name in PROJECT_TABLES.values():
            columns = self._columns_for_table(table_name)
            missing = [column for column in PROJECT_COLUMNS if column not in columns]
            if missing:
                missing_columns[table_name] = missing
        if missing_columns:
            details = "; ".join(f"{table}: {', '.join(cols)}" for table, cols in missing_columns.items())
            raise WorkbookFormatError(f"SharePoint workbook has missing required column(s): {details}")

    def _columns_for_table(self, table_name: str) -> list[str]:
        if table_name not in self._column_cache:
            columns = self.client.list_table_columns(table_name)
            self._column_cache[table_name] = [str(col.get("name") or "").strip() for col in columns]
        return self._column_cache[table_name]

    def _table_for_sheet(self, sheet_name: str) -> str:
        try:
            return PROJECT_TABLES[sheet_name]
        except KeyError as exc:
            raise WorkbookFormatError(f"No SharePoint table mapping is configured for sheet: {sheet_name}") from exc

    def _row_to_dict(self, sheet_name: str, table_name: str, row_payload: dict[str, Any]) -> dict[str, Any]:
        columns = self._columns_for_table(table_name)
        raw_values = row_payload.get("values") or [[]]
        values = raw_values[0] if raw_values and isinstance(raw_values[0], list) else raw_values
        by_column = {column: _normalizer(values[idx] if idx < len(values) else "") for idx, column in enumerate(columns)}
        record_id = by_column.get("Record ID") or f"{table_name}:{row_payload.get('index')}"
        row = {
            "row": record_id,
            "record_id": by_column.get("Record ID", ""),
            "source_index": row_payload.get("index"),
            "sheet": sheet_name,
        }
        for excel_name, app_key in EXCEL_TO_APP.items():
            row[app_key] = by_column.get(excel_name, "")
        return row

    def list_projects(self, sheet_name: str, allow_missing: bool = False) -> list[dict[str, Any]]:
        if allow_missing and sheet_name not in PROJECT_TABLES:
            return []
        table_name = self._table_for_sheet(sheet_name)
        rows = [
            self._row_to_dict(sheet_name, table_name, row)
            for row in self.client.list_table_rows(table_name)
        ]
        rows = [row for row in rows if row.get("client")]
        return sorted(rows, key=lambda row: (str(row.get("client", "")).lower(), str(row.get("site", "")).lower()))

    def _find_row(self, sheet_name: str, row_key: str) -> dict[str, Any]:
        table_name = self._table_for_sheet(sheet_name)
        wanted = str(row_key)
        for row in self.client.list_table_rows(table_name):
            mapped = self._row_to_dict(sheet_name, table_name, row)
            if str(mapped.get("row")) == wanted or str(mapped.get("record_id")) == wanted:
                return mapped
        raise SharePointSyncError(f"Could not find SharePoint row with Record ID: {row_key}")

    def _values_for_columns(
        self,
        table_name: str,
        values: dict[str, Any],
        *,
        record_id: str,
        modified_utc: str,
        modified_by: str,
    ) -> list[Any]:
        by_excel = {
            "Record ID": record_id,
            "Client": values.get("client", ""),
            "Project Name": values.get("site", ""),
            "Dropbox URLs": values.get("link", ""),
            "Operations Contact Email": values.get("email", ""),
            "Procore": values.get("procore", ""),
            "Internal URLs": values.get("internal_url", ""),
            "Last Modified UTC": modified_utc,
            "Last Modified By": modified_by,
        }
        return [by_excel.get(column, "") for column in self._columns_for_table(table_name)]

    def _check_conflict(self, remote_row: dict[str, Any], expected_modified_utc: str | None, force: bool) -> None:
        if force or not expected_modified_utc:
            return
        remote_modified = str(remote_row.get("modified_utc") or "")
        if remote_modified and remote_modified != str(expected_modified_utc):
            raise ProjectConflictError(
                "This row changed in SharePoint after you opened it.",
                remote_row=remote_row,
            )

    def save_project(
        self,
        sheet_name: str,
        values: dict[str, Any],
        row_key: str | int | None = None,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        table_name = self._table_for_sheet(sheet_name)
        modified_utc = _utc_timestamp()
        modified_by = self.client.get_current_user()
        if row_key is None:
            record_id = uuid.uuid4().hex
            row_values = self._values_for_columns(
                table_name,
                values,
                record_id=record_id,
                modified_utc=modified_utc,
                modified_by=modified_by,
            )
            self.client.add_table_row(table_name, row_values)
            return

        remote = self._find_row(sheet_name, str(row_key))
        self._check_conflict(remote, expected_modified_utc, force)
        record_id = str(remote.get("record_id") or row_key)
        row_values = self._values_for_columns(
            table_name,
            values,
            record_id=record_id,
            modified_utc=modified_utc,
            modified_by=modified_by,
        )
        self.client.update_table_row(table_name, int(remote["source_index"]), row_values)

    def delete_project(
        self,
        sheet_name: str,
        row_key: str | int,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        table_name = self._table_for_sheet(sheet_name)
        remote = self._find_row(sheet_name, str(row_key))
        self._check_conflict(remote, expected_modified_utc, force)
        self.client.delete_table_row(table_name, int(remote["source_index"]))

    def move_project_to_completed(
        self,
        sheet_name: str,
        row_key: str | int,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        if sheet_name == COMPLETED_SHEET:
            return
        source_table = self._table_for_sheet(sheet_name)
        completed_table = self._table_for_sheet(COMPLETED_SHEET)
        remote = self._find_row(sheet_name, str(row_key))
        self._check_conflict(remote, expected_modified_utc, force)
        modified_utc = _utc_timestamp()
        modified_by = self.client.get_current_user()
        row_values = self._values_for_columns(
            completed_table,
            remote,
            record_id=str(remote.get("record_id") or row_key),
            modified_utc=modified_utc,
            modified_by=modified_by,
        )
        self.client.add_table_row(completed_table, row_values)
        self.client.delete_table_row(source_table, int(remote["source_index"]))

    def sort_sheet(self, sheet_name: str) -> None:
        # The UI displays sorted rows after every reload. Avoid changing the shared
        # workbook's current user-facing sort/filter state from the desktop app.
        return

    def clear_token_cache(self) -> None:
        self.client.clear_token_cache()


class LocalExcelProjectDataSource:
    is_remote = False
    description = "Local workbook"
    web_url = ""

    def __init__(self, workbook_path: Path, *, email_normalizer=None) -> None:
        self.workbook_path = Path(workbook_path)
        self.email_normalizer = email_normalizer or _normalizer

    def _resolve_sheet_name(self, wb, sheet_name: str, allow_missing: bool = False):
        for candidate in _sheet_aliases(sheet_name):
            if candidate in wb.sheetnames:
                return candidate
        if allow_missing:
            return None
        if sheet_name == MASTER_SHEET:
            expected = " or ".join(f'"{name}"' for name in _sheet_aliases(sheet_name))
            raise KeyError(f"Workbook is missing the Construction sheet. Expected {expected}.")
        raise KeyError(f'Workbook is missing the "{sheet_name}" sheet.')

    def _master_column_map(self, ws, ensure: bool = False):
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {}
        for idx, header in enumerate(headers, start=1):
            if not header:
                continue
            normalized = str(header).strip().lower()
            if normalized in LOCAL_HEADER_ALIASES and normalized not in col_map:
                col_map[normalized] = idx
        changed = False
        if ensure:
            for normalized, header in LOCAL_HEADER_ALIASES.items():
                if normalized not in col_map:
                    idx = ws.max_column + 1
                    ws.cell(row=1, column=idx, value=header)
                    col_map[normalized] = idx
                    changed = True
        return col_map, changed

    def _get_or_create_sheet(self, wb, sheet_name: str):
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=True)
        if resolved_sheet:
            ws = wb[resolved_sheet]
            if not any(cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))):
                headers = list(LOCAL_HEADER_ALIASES.values())
                source_sheet_name = self._resolve_sheet_name(wb, MASTER_SHEET, allow_missing=True)
                if source_sheet_name:
                    source_headers = [cell.value for cell in next(wb[source_sheet_name].iter_rows(min_row=1, max_row=1))]
                    if any(source_headers):
                        headers = source_headers
                for col_idx, header in enumerate(headers, start=1):
                    ws.cell(row=1, column=col_idx, value=header)
            self._master_column_map(ws, ensure=True)
            return ws

        ws = wb.create_sheet(sheet_name)
        source_sheet_name = self._resolve_sheet_name(wb, MASTER_SHEET, allow_missing=True)
        headers = list(LOCAL_HEADER_ALIASES.values())
        if source_sheet_name:
            source_headers = [cell.value for cell in next(wb[source_sheet_name].iter_rows(min_row=1, max_row=1))]
            if any(source_headers):
                headers = source_headers
        ws.append(headers)
        return ws

    def _build_master_row(self, ws, values: dict[str, Any]):
        col_map, _ = self._master_column_map(ws, ensure=True)
        row = [""] * max(col_map.values())
        for key, value in values.items():
            col_idx = col_map.get(key)
            if col_idx is not None:
                row[col_idx - 1] = value
        return row

    def ensure_schema(self) -> None:
        if not self.workbook_path.exists():
            return
        wb = openpyxl.load_workbook(self.workbook_path)
        changed = False
        try:
            for sheet_name in (MASTER_SHEET, OIL_GAS_SHEET, COMPLETED_SHEET):
                resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=True)
                if not resolved_sheet:
                    continue
                headers_before = [cell.value for cell in next(wb[resolved_sheet].iter_rows(min_row=1, max_row=1))]
                ws = self._get_or_create_sheet(wb, sheet_name)
                _, sheet_changed = self._master_column_map(ws, ensure=True)
                headers_after = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                changed = changed or sheet_changed or headers_after != headers_before
            if changed:
                wb.save(self.workbook_path)
        finally:
            wb.close()

    def list_projects(self, sheet_name: str, allow_missing: bool = False) -> list[dict[str, Any]]:
        wb = openpyxl.load_workbook(self.workbook_path, data_only=True)
        try:
            resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=allow_missing)
            if resolved_sheet is None:
                return []
            ws = wb[resolved_sheet]
            col_map, _ = self._master_column_map(ws, ensure=False)
            client_idx = col_map.get("client", 1) - 1
            site_idx = col_map.get("project name", 2) - 1
            link_idx = col_map.get("dropbox urls", 3) - 1
            email_idx = col_map.get("operations contact email", 4) - 1
            procore_idx = col_map.get("procore")
            internal_url_idx = col_map.get("internal urls")
            procore_idx = procore_idx - 1 if procore_idx is not None else None
            internal_url_idx = internal_url_idx - 1 if internal_url_idx is not None else None
            rows = []
            for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                client = _master_cell_value(row, client_idx)
                site = _master_cell_value(row, site_idx)
                link = _master_cell_value(row, link_idx)
                email = _master_cell_value(row, email_idx)
                procore = _master_cell_value(row, procore_idx)
                internal_url = _master_cell_value(row, internal_url_idx)
                if client:
                    rows.append({
                        "row": idx,
                        "client": client,
                        "site": site,
                        "link": link,
                        "email": self.email_normalizer(email),
                        "procore": procore,
                        "internal_url": internal_url,
                        "record_id": "",
                        "modified_utc": "",
                        "modified_by": "",
                    })
            return rows
        finally:
            wb.close()

    def sort_sheet(self, sheet_name: str) -> None:
        wb = openpyxl.load_workbook(self.workbook_path)
        try:
            resolved_sheet = self._resolve_sheet_name(
                wb,
                sheet_name,
                allow_missing=(sheet_name in (OIL_GAS_SHEET, COMPLETED_SHEET)),
            )
            if resolved_sheet is None:
                return
            ws = wb[resolved_sheet]
            col_map, _ = self._master_column_map(ws, ensure=True)
            rows = list(ws.iter_rows(min_row=2, values_only=True))

            def key(row):
                client = _master_cell_value(row, col_map.get("client", 1) - 1)
                site = _master_cell_value(row, col_map.get("project name", 2) - 1)
                return (str(client).strip().lower(), str(site).strip().lower())

            rows_sorted = sorted(rows, key=key)
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for row in rows_sorted:
                ws.append(list(row))
            wb.save(self.workbook_path)
        finally:
            wb.close()

    def save_project(
        self,
        sheet_name: str,
        values: dict[str, Any],
        row_key: str | int | None = None,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        wb = openpyxl.load_workbook(self.workbook_path)
        try:
            if row_key is None:
                ws = self._get_or_create_sheet(wb, sheet_name)
                ws.append(self._build_master_row(ws, {
                    "client": values.get("client", ""),
                    "project name": values.get("site", ""),
                    "dropbox urls": values.get("link", ""),
                    "operations contact email": values.get("email", ""),
                    "procore": values.get("procore", ""),
                    "internal urls": values.get("internal_url", ""),
                }))
            else:
                ws = wb[self._resolve_sheet_name(wb, sheet_name)]
                col_map, _ = self._master_column_map(ws, ensure=True)
                row_idx = int(row_key)
                ws.cell(row=row_idx, column=col_map["client"], value=values.get("client", ""))
                ws.cell(row=row_idx, column=col_map["project name"], value=values.get("site", ""))
                ws.cell(row=row_idx, column=col_map["dropbox urls"], value=values.get("link", ""))
                ws.cell(row=row_idx, column=col_map["operations contact email"], value=values.get("email", ""))
                ws.cell(row=row_idx, column=col_map["procore"], value=values.get("procore", ""))
                ws.cell(row=row_idx, column=col_map["internal urls"], value=values.get("internal_url", ""))
            wb.save(self.workbook_path)
        finally:
            wb.close()
        self.sort_sheet(sheet_name)

    def delete_project(
        self,
        sheet_name: str,
        row_key: str | int,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        wb = openpyxl.load_workbook(self.workbook_path)
        try:
            ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            ws.delete_rows(int(row_key), 1)
            wb.save(self.workbook_path)
        finally:
            wb.close()
        self.sort_sheet(sheet_name)

    def move_project_to_completed(
        self,
        sheet_name: str,
        row_key: str | int,
        *,
        expected_modified_utc: str | None = None,
        force: bool = False,
    ) -> None:
        wb = openpyxl.load_workbook(self.workbook_path)
        try:
            source_ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            completed_ws = self._get_or_create_sheet(wb, COMPLETED_SHEET)
            row_idx = int(row_key)
            row_values = [
                source_ws.cell(row=row_idx, column=col_idx).value
                for col_idx in range(1, source_ws.max_column + 1)
            ]
            completed_ws.append(row_values)
            source_ws.delete_rows(row_idx, 1)
            wb.save(self.workbook_path)
        finally:
            wb.close()
        self.sort_sheet(sheet_name)
        self.sort_sheet(COMPLETED_SHEET)


def _retry_after(response) -> int | None:
    value = response.headers.get("Retry-After")
    if not value:
        return None
    try:
        return max(0, int(value))
    except ValueError:
        return None


def _graph_error(response) -> GraphApiError:
    retry_after = _retry_after(response)
    code = None
    message = response.text
    details = None
    try:
        details = response.json()
        error = details.get("error") or {}
        code = error.get("code")
        message = error.get("message") or message
    except Exception:
        details = response.text
    return GraphApiError(
        f"Microsoft Graph request failed ({response.status_code}): {message}",
        status_code=response.status_code,
        code=code,
        retry_after=retry_after,
        details=details,
    )
