import datetime as _dt
import calendar as _cal
import base64
import mimetypes
import tempfile
import uuid
import re
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from urllib.parse import unquote
from pathlib import Path
import tkinter as tk
from tkinter import messagebox

try:
    import ttkbootstrap as ttk
except Exception:  # pragma: no cover
    from tkinter import ttk

import openpyxl

try:
    import win32com.client as win32
except Exception:  # pragma: no cover
    win32 = None

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "2026_02_03"
MASTER_FILE = DATA_DIR / "Client-Site-URL-Email.xlsx"
MASTER_SHEET = "ABCD"
CC_LISTS_FILE = BASE_DIR / "Internal CC Lists.xlsx"
CC_LISTS_SHEET = "Sheet1"
SIGNATURE_NAME = "Main Bottom"
EMAIL_TEMPLATES_DIR = BASE_DIR / "Email Templates"

MEDIA_OPTIONS = [
    "photos",
    "photos and video",
    "photos, video, and 2D mapping",
]

MEDIA_SUBJECT = {
    "photos": "Photos",
    "photos and video": "Photos and Video",
    "photos, video, and 2D mapping": "Photos, Video, and 2D Mapping",
}

MEDIA_BODY = {
    "photos": "photos",
    "photos and video": "photos and video",
    "photos, video, and 2D mapping": "photos, video, and 2D mapping",
}


def _ordinal(n: int) -> str:
    if 10 <= (n % 100) <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def _normalize_key(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    cleaned = "".join(ch.lower() if ch.isalnum() else " " for ch in value)
    return " ".join(cleaned.split())



def format_date(date_obj: _dt.date) -> str:
    return date_obj.strftime(f"%B {_ordinal(date_obj.day)}, %Y")


def default_date() -> _dt.date:
    today = _dt.date.today()
    if today.weekday() == 0:  # Monday
        return today - _dt.timedelta(days=3)
    return today - _dt.timedelta(days=1)


class CalendarPopup(tk.Toplevel):
    def __init__(self, master, initial_date, on_select):
        super().__init__(master)
        self.title("Select Date")
        self.resizable(False, False)
        self.on_select = on_select
        self.year = initial_date.year
        self.month = initial_date.month

        header = ttk.Frame(self)
        header.pack(padx=10, pady=8, fill="x")

        ttk.Button(header, text="<", width=3, command=self._prev_month).pack(side="left")
        self.month_label = ttk.Label(header, text="")
        self.month_label.pack(side="left", expand=True)
        ttk.Button(header, text=">", width=3, command=self._next_month).pack(side="right")

        self.grid_frame = ttk.Frame(self)
        self.grid_frame.pack(padx=10, pady=8)

        self._render()

    def _prev_month(self):
        if self.month == 1:
            self.month = 12
            self.year -= 1
        else:
            self.month -= 1
        self._render()

    def _next_month(self):
        if self.month == 12:
            self.month = 1
            self.year += 1
        else:
            self.month += 1
        self._render()

    def _render(self):
        for child in self.grid_frame.winfo_children():
            child.destroy()

        self.month_label.config(text=f"{_cal.month_name[self.month]} {self.year}")

        for idx, day in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]):
            ttk.Label(self.grid_frame, text=day, width=4, anchor="center").grid(row=0, column=idx, padx=2, pady=2)

        month_data = _cal.monthcalendar(self.year, self.month)
        for r, week in enumerate(month_data, start=1):
            for c, day in enumerate(week):
                if day == 0:
                    ttk.Label(self.grid_frame, text="", width=4).grid(row=r, column=c, padx=2, pady=2)
                    continue
                ttk.Button(
                    self.grid_frame,
                    text=str(day),
                    width=4,
                    command=lambda d=day: self._select_day(d),
                ).grid(row=r, column=c, padx=2, pady=2)

    def _select_day(self, day):
        chosen = _dt.date(self.year, self.month, day)
        self.on_select(chosen)
        self.destroy()


class EmailTemplateApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Aerial Progress Email Builder")
        self.root.geometry("1200x720")

        self.dropbox_data = []
        self.contacts_data = []
        self.cc_lists = {}
        self.signature_names = []
        self.email_templates = {}

        self._apply_theme()
        self._load_data()
        self._build_ui()
        self._refresh_clients()
        self._set_default_date()
        self._update_email_preview()

    def _load_data(self):
        self.dropbox_data = self._read_dropbox_data()
        self.contacts_data = self._read_contacts_data()
        self.cc_lists = self._read_cc_lists()
        self.signature_names = self._read_signature_names()
        self.email_templates = self._read_email_templates()

    def _read_email_templates(self):
        templates = {}
        if not EMAIL_TEMPLATES_DIR.exists():
            return templates
        for path in EMAIL_TEMPLATES_DIR.glob("*.eml"):
            name = path.stem
            try:
                msg = BytesParser(policy=policy.default).parse(path.open("rb"))
            except Exception:
                continue
            subject = msg.get("Subject", "")
            text = ""
            html = ""
            if msg.is_multipart():
                for part in msg.walk():
                    ctype = part.get_content_type()
                    if ctype == "text/plain" and not text:
                        text = part.get_content()
                    elif ctype == "text/html" and not html:
                        html = part.get_content()
            else:
                ctype = msg.get_content_type()
                if ctype == "text/plain":
                    text = msg.get_content()
                elif ctype == "text/html":
                    html = msg.get_content()
            templates[name] = {
                "subject": subject or "",
                "text": text or "",
                "html": html or "",
            }
        return templates

    def _read_dropbox_data(self):
        wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h}
        client_idx = col_map.get("client", 0)
        site_idx = col_map.get("project name", 1)
        link_idx = col_map.get("dropbox urls", 2)
        rows = []
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            client = (row[client_idx] or "").strip() if isinstance(row[client_idx], str) else row[client_idx]
            site = (row[site_idx] or "").strip() if isinstance(row[site_idx], str) else row[site_idx]
            link = (row[link_idx] or "").strip() if isinstance(row[link_idx], str) else row[link_idx]
            if client and site:
                rows.append({
                    "row": idx,
                    "client": client,
                    "site": site,
                    "link": link or "",
                })
        wb.close()
        return rows

    def _read_contacts_data(self):
        wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h}
        client_idx = col_map.get("client", 0)
        site_idx = col_map.get("project name", 1)
        email_idx = col_map.get("operations contact email", 3)
        rows = []
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            client = (row[client_idx] or "").strip() if isinstance(row[client_idx], str) else row[client_idx]
            site = (row[site_idx] or "").strip() if isinstance(row[site_idx], str) else row[site_idx]
            email = (row[email_idx] or "").strip() if isinstance(row[email_idx], str) else row[email_idx]
            if client and site:
                rows.append({
                    "row": idx,
                    "client": client,
                    "site": site,
                    "email": email or "",
                })
        wb.close()
        return rows

    def _read_cc_lists(self):
        wb = openpyxl.load_workbook(CC_LISTS_FILE, data_only=True)
        ws = wb[CC_LISTS_SHEET]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        lists = {}
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            emails = []
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
                email = row[0]
                if isinstance(email, str):
                    email = email.strip()
                if email:
                    emails.append(email)
            lists[str(header).strip()] = emails
        wb.close()
        return lists

    def _read_signature_names(self):
        sig_dir = Path.home() / "AppData" / "Roaming" / "Microsoft" / "Signatures"
        names = []
        if sig_dir.exists():
            for path in sig_dir.glob("*.htm"):
                name = path.stem
                if name not in names:
                    names.append(name)
        return sorted(names)

    def _extract_inline_images(self, html, base_dir=None):
        assets = []
        def _write_asset(mime, data):
            data = re.sub(r"\s+", "", data)
            cid = f"sig-{uuid.uuid4().hex}"
            ext = {"png": ".png", "jpg": ".jpg", "jpeg": ".jpg", "gif": ".gif"}.get(mime, ".png")
            if not hasattr(self, "_sig_temp_dir") or not self._sig_temp_dir:
                self._sig_temp_dir = tempfile.mkdtemp(prefix="signature_images_")
            path = Path(self._sig_temp_dir) / f"{cid}{ext}"
            path.write_bytes(base64.b64decode(data))
            assets.append({"cid": cid, "path": str(path)})
            return cid

        # Handle src="data:image/...;base64,..." and src='...'
        def _replace_quoted(match):
            quote = match.group(1)
            mime = match.group(2).lower()
            data = match.group(3)
            try:
                cid = _write_asset(mime, data)
            except Exception:
                return match.group(0)
            return f'src={quote}cid:{cid}{quote}'

        pattern_quoted = re.compile(
            r'src=(["\'])data:image/([a-zA-Z0-9+]+);base64,([^"\']+)\1',
            flags=re.IGNORECASE | re.DOTALL,
        )
        html = pattern_quoted.sub(_replace_quoted, html)

        # Handle src=data:image/...;base64,... (unquoted)
        def _replace_unquoted(match):
            mime = match.group(1).lower()
            data = match.group(2)
            try:
                cid = _write_asset(mime, data)
            except Exception:
                return match.group(0)
            return f"src=cid:{cid}"

        pattern_unquoted = re.compile(
            r'src=data:image/([a-zA-Z0-9+]+);base64,([^\\s>]+)',
            flags=re.IGNORECASE | re.DOTALL,
        )
        html = pattern_unquoted.sub(_replace_unquoted, html)
        # Handle relative file references from Outlook signature folders
        def _replace_file_src(match):
            quote = match.group(1)
            src = match.group(2)
            lower = src.lower()
            if lower.startswith("cid:") or lower.startswith("data:") or lower.startswith("http"):
                return match.group(0)
            if not base_dir:
                return match.group(0)
            rel = unquote(src)
            path = (Path(base_dir) / rel).resolve()
            if not path.exists():
                return match.group(0)
            cid = f"sig-{uuid.uuid4().hex}"
            assets.append({"cid": cid, "path": str(path)})
            return f'src={quote}cid:{cid}{quote}'

        # img src="..."; also handles v:imagedata src="..."
        html = re.sub(
            r'src=(["\'])([^"\']+)\1',
            _replace_file_src,
            html,
            flags=re.IGNORECASE,
        )
        return html, assets

    def _load_signature_html(self, signature_name=None):
        sig_dir = Path.home() / "AppData" / "Roaming" / "Microsoft" / "Signatures"
        target = signature_name or SIGNATURE_NAME
        candidates = []
        if target:
            exact = sig_dir / f"{target}.htm"
            if exact.exists():
                candidates = [exact]
        if not candidates:
            candidates = sorted(sig_dir.glob(f"{target}*.htm"))
        if not candidates:
            candidates = sorted(sig_dir.glob("Main*.htm"))
        path = candidates[0] if candidates else None
        if not path or not path.exists():
            return "", []

        html = path.read_text(encoding="utf-8", errors="ignore")
        marker = "Thank you,"
        idx = html.find(marker)
        if idx != -1:
            html = html[idx:]
        return self._extract_inline_images(html, base_dir=path.parent)

    def _write_cc_lists(self):
        if CC_LISTS_FILE.exists():
            wb = openpyxl.load_workbook(CC_LISTS_FILE)
        else:
            wb = openpyxl.Workbook()
        if CC_LISTS_SHEET in wb.sheetnames:
            ws_old = wb[CC_LISTS_SHEET]
            wb.remove(ws_old)
        ws = wb.create_sheet(CC_LISTS_SHEET, 0)
        headers = list(self.cc_lists.keys())
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)
            for row_idx, email in enumerate(self.cc_lists[header], start=2):
                ws.cell(row=row_idx, column=col_idx, value=email)
        wb.save(CC_LISTS_FILE)
        wb.close()

    def _build_ui(self):
        style = ttk.Style()
        try:
            style.configure("LeftTab.TNotebook", tabposition="wn")
        except tk.TclError:
            pass

        notebook = ttk.Notebook(self.root, style="LeftTab.TNotebook")
        notebook.pack(fill="both", expand=True)

        self.compose_tab = ttk.Frame(notebook)
        self.dropbox_tab = ttk.Frame(notebook)
        self.contacts_tab = ttk.Frame(notebook)
        self.cc_tab = ttk.Frame(notebook)

        notebook.add(self.compose_tab, text="Compose")
        notebook.add(self.dropbox_tab, text="Dropbox Links")
        notebook.add(self.contacts_tab, text="Operations Contacts")
        notebook.add(self.cc_tab, text="Internal CC Lists")

        self._build_compose_tab()
        self._build_dropbox_tab()
        self._build_contacts_tab()
        self._build_cc_tab()

    def _apply_theme(self):
        style = ttk.Style()
        try:
            style.theme_use("darkly")
        except Exception:
            try:
                style.theme_use("clam")
            except tk.TclError:
                pass

        # Match the pyweb-local (Docker Desktop-inspired) palette
        colors = {
            "bg": "#1a1a1a",
            "panel": "#2a2a2a",
            "sidebar": "#1e1e1e",
            "border": "#404040",
            "text": "#ffffff",
            "text_secondary": "#b3b3b3",
            "text_muted": "#858585",
            "primary": "#314dec",
            "info": "#3a7ca5",
            "success": "#10b981",
            "warning": "#f59e0b",
            "error": "#ef4444",
        }

        self.root.configure(bg=colors["bg"])

        style.configure("TFrame", background=colors["bg"])
        style.configure("Card.TFrame", background=colors["panel"])
        style.configure("TLabel", background=colors["bg"], foreground=colors["text"], font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=colors["bg"], foreground=colors["text_secondary"])

        style.configure("TNotebook", background=colors["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", padding=(14, 10), background=colors["bg"], foreground=colors["text_secondary"])
        style.map("TNotebook.Tab", background=[("selected", colors["panel"])], foreground=[("selected", colors["text"])])

        style.configure("TEntry", fieldbackground=colors["panel"], background=colors["panel"], foreground=colors["text"])
        style.configure("TCombobox", fieldbackground=colors["panel"], background=colors["panel"], foreground=colors["text"])

        style.configure("TButton", background=colors["panel"], foreground=colors["text"], padding=(10, 6))
        style.map(
            "TButton",
            background=[("active", colors["primary"])],
            foreground=[("active", colors["text"])],
        )

        style.configure(
            "Treeview",
            background=colors["panel"],
            fieldbackground=colors["panel"],
            foreground=colors["text"],
            bordercolor=colors["border"],
            borderwidth=1,
        )
        style.configure(
            "Treeview.Heading",
            background=colors["sidebar"],
            foreground=colors["text_secondary"],
            font=("Segoe UI", 9, "bold"),
        )
        style.map("Treeview", background=[("selected", "#2f3d5a")], foreground=[("selected", colors["text"])])

        self._theme_colors = colors


    def _build_compose_tab(self):
        left = ttk.Frame(self.compose_tab)
        left.pack(side="left", fill="y", padx=20, pady=20)

        right = ttk.Frame(self.compose_tab)
        right.pack(side="right", fill="both", expand=True, padx=20, pady=20)

        ttk.Label(left, text="Client").pack(anchor="w")
        self.client_var = tk.StringVar()
        self.client_combo = ttk.Combobox(left, textvariable=self.client_var, state="readonly", width=40)
        self.client_combo.pack(anchor="w", pady=(0, 12))
        self.client_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_client_change())

        ttk.Label(left, text="Site").pack(anchor="w")
        self.site_var = tk.StringVar()
        self.site_combo = ttk.Combobox(left, textvariable=self.site_var, state="readonly", width=40)
        self.site_combo.pack(anchor="w", pady=(0, 12))
        self.site_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_site_change())

        ttk.Label(left, text="Media Type").pack(anchor="w")
        self.media_var = tk.StringVar(value=MEDIA_OPTIONS[1])
        self.media_combo = ttk.Combobox(left, textvariable=self.media_var, values=MEDIA_OPTIONS, state="readonly", width=40)
        self.media_combo.pack(anchor="w", pady=(0, 12))
        self.media_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        ttk.Label(left, text="Date").pack(anchor="w")
        date_frame = ttk.Frame(left)
        date_frame.pack(anchor="w", pady=(0, 12))
        self.date_var = tk.StringVar()
        self.date_entry = ttk.Entry(date_frame, textvariable=self.date_var, width=32, state="readonly")
        self.date_entry.pack(side="left")
        ttk.Button(date_frame, text="Pick", command=self._open_calendar).pack(side="left", padx=6)

        ttk.Label(left, text="To").pack(anchor="w")
        self.to_var = tk.StringVar()
        self.to_entry = ttk.Entry(left, textvariable=self.to_var, width=50)
        self.to_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(left, text="CC List").pack(anchor="w")
        self.cc_list_var = tk.StringVar()
        self.cc_list_combo = ttk.Combobox(left, textvariable=self.cc_list_var, state="readonly", width=40)
        self.cc_list_combo.pack(anchor="w", pady=(0, 6))
        self.cc_list_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_cc_list_change())

        ttk.Label(left, text="CC").pack(anchor="w")
        self.cc_var = tk.StringVar()
        self.cc_entry = ttk.Entry(left, textvariable=self.cc_var, width=50)
        self.cc_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(left, text="Signature").pack(anchor="w")
        self.signature_var = tk.StringVar()
        self.signature_combo = ttk.Combobox(left, textvariable=self.signature_var, state="readonly", width=40)
        self.signature_combo.pack(anchor="w", pady=(0, 12))
        self.signature_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        self.show_bcc_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(left, text="Show BCC", variable=self.show_bcc_var, command=self._toggle_bcc).pack(anchor="w")

        self.bcc_frame = ttk.Frame(left)
        ttk.Label(self.bcc_frame, text="BCC (optional)").pack(anchor="w")
        self.bcc_var = tk.StringVar()
        self.bcc_entry = ttk.Entry(self.bcc_frame, textvariable=self.bcc_var, width=50)
        self.bcc_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(left, text="Dropbox Link").pack(anchor="w")
        self.link_var = tk.StringVar()
        self.link_entry = ttk.Entry(left, textvariable=self.link_var, width=50)
        self.link_entry.pack(anchor="w", pady=(0, 12))

        ttk.Button(left, text="Create Outlook Draft", command=self._create_draft).pack(anchor="w", pady=(6, 0))

        ttk.Label(right, text="Subject").pack(anchor="w")
        self.subject_var = tk.StringVar()
        self.subject_entry = ttk.Entry(right, textvariable=self.subject_var, width=100)
        self.subject_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(right, text="Email Preview").pack(anchor="w")
        preview_frame = ttk.Frame(right)
        preview_frame.pack(fill="both", expand=True)
        self.preview = tk.Text(preview_frame, wrap="word", height=28)
        preview_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.preview.yview)
        self.preview.configure(yscrollcommand=preview_scroll.set)
        self.preview.pack(side="left", fill="both", expand=True)
        preview_scroll.pack(side="right", fill="y")
        self.preview.config(state="disabled")

        for var in [
            self.client_var,
            self.site_var,
            self.media_var,
            self.date_var,
            self.link_var,
            self.to_var,
            self.cc_var,
            self.bcc_var,
            self.signature_var,
            self.subject_var,
        ]:
            var.trace_add("write", lambda *_args: self._update_email_preview())

    def _build_dropbox_tab(self):
        frame = ttk.Frame(self.dropbox_tab)
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        self.dropbox_tree = ttk.Treeview(tree_frame, columns=("client", "site", "link"), show="headings")
        self.dropbox_tree.heading("client", text="Client")
        self.dropbox_tree.heading("site", text="Site")
        self.dropbox_tree.heading("link", text="Customer Dropbox URL")
        self.dropbox_tree.column("client", width=200)
        self.dropbox_tree.column("site", width=260)
        self.dropbox_tree.column("link", width=520)
        dropbox_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.dropbox_tree.yview)
        self.dropbox_tree.configure(yscrollcommand=dropbox_scroll.set)
        self.dropbox_tree.pack(side="left", fill="both", expand=True)
        dropbox_scroll.pack(side="right", fill="y")
        self.dropbox_tree.bind("<<TreeviewSelect>>", lambda _e: self._load_dropbox_selection())

        form = ttk.Frame(frame)
        form.pack(fill="x", pady=10)

        ttk.Label(form, text="Client").grid(row=0, column=0, sticky="w")
        ttk.Label(form, text="Site").grid(row=0, column=1, sticky="w")
        ttk.Label(form, text="Customer Dropbox URL").grid(row=0, column=2, sticky="w")

        self.db_client_var = tk.StringVar()
        self.db_site_var = tk.StringVar()
        self.db_link_var = tk.StringVar()

        ttk.Entry(form, textvariable=self.db_client_var, width=30).grid(row=1, column=0, sticky="w", padx=(0, 10))
        ttk.Entry(form, textvariable=self.db_site_var, width=40).grid(row=1, column=1, sticky="w", padx=(0, 10))
        ttk.Entry(form, textvariable=self.db_link_var, width=70).grid(row=1, column=2, sticky="w")

        btns = ttk.Frame(frame)
        btns.pack(fill="x", pady=4)
        ttk.Button(btns, text="Add", command=self._add_dropbox_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Update", command=self._update_dropbox_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Delete", command=self._delete_dropbox_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Reload", command=self._reload_dropbox_data).pack(side="left")

        self._refresh_dropbox_tree()

    def _build_contacts_tab(self):
        frame = ttk.Frame(self.contacts_tab)
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        self.contacts_tree = ttk.Treeview(tree_frame, columns=("client", "site", "email"), show="headings")
        self.contacts_tree.heading("client", text="Client")
        self.contacts_tree.heading("site", text="Site")
        self.contacts_tree.heading("email", text="Operations Contact Email")
        self.contacts_tree.column("client", width=200)
        self.contacts_tree.column("site", width=260)
        self.contacts_tree.column("email", width=520)
        contacts_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.contacts_tree.yview)
        self.contacts_tree.configure(yscrollcommand=contacts_scroll.set)
        self.contacts_tree.pack(side="left", fill="both", expand=True)
        contacts_scroll.pack(side="right", fill="y")
        self.contacts_tree.bind("<<TreeviewSelect>>", lambda _e: self._load_contact_selection())

        form = ttk.Frame(frame)
        form.pack(fill="x", pady=10)

        ttk.Label(form, text="Client").grid(row=0, column=0, sticky="w")
        ttk.Label(form, text="Site").grid(row=0, column=1, sticky="w")
        ttk.Label(form, text="Operations Contact Email").grid(row=0, column=2, sticky="w")

        self.ct_client_var = tk.StringVar()
        self.ct_site_var = tk.StringVar()
        self.ct_email_var = tk.StringVar()

        ttk.Entry(form, textvariable=self.ct_client_var, width=30).grid(row=1, column=0, sticky="w", padx=(0, 10))
        ttk.Entry(form, textvariable=self.ct_site_var, width=40).grid(row=1, column=1, sticky="w", padx=(0, 10))
        ttk.Entry(form, textvariable=self.ct_email_var, width=70).grid(row=1, column=2, sticky="w")

        btns = ttk.Frame(frame)
        btns.pack(fill="x", pady=4)
        ttk.Button(btns, text="Add", command=self._add_contact_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Update", command=self._update_contact_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Delete", command=self._delete_contact_entry).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Reload", command=self._reload_contacts_data).pack(side="left")

        self._refresh_contacts_tree()

    def _build_cc_tab(self):
        frame = ttk.Frame(self.cc_tab)
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        left = ttk.Frame(frame)
        left.pack(side="left", fill="y", padx=(0, 20))
        right = ttk.Frame(frame)
        right.pack(side="left", fill="both", expand=True)

        ttk.Label(left, text="CC Lists").pack(anchor="w")
        cc_list_frame = ttk.Frame(left)
        cc_list_frame.pack(anchor="w", pady=(4, 8), fill="y")
        self.cc_listbox = tk.Listbox(cc_list_frame, height=20, width=28)
        cc_list_scroll = ttk.Scrollbar(cc_list_frame, orient="vertical", command=self.cc_listbox.yview)
        self.cc_listbox.configure(yscrollcommand=cc_list_scroll.set)
        self.cc_listbox.pack(side="left", fill="y")
        cc_list_scroll.pack(side="right", fill="y")
        self.cc_listbox.bind("<<ListboxSelect>>", lambda _e: self._load_cc_list_selection())

        self.new_cc_list_var = tk.StringVar()
        ttk.Entry(left, textvariable=self.new_cc_list_var, width=28).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Create List", command=self._create_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Delete List", command=self._delete_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Reload", command=self._reload_cc_lists_data).pack(anchor="w")

        ttk.Label(right, text="Emails in Selected List").pack(anchor="w")
        cc_emails_frame = ttk.Frame(right)
        cc_emails_frame.pack(anchor="w", pady=(4, 8), fill="both", expand=True)
        self.cc_emails_listbox = tk.Listbox(cc_emails_frame, height=20)
        cc_emails_scroll = ttk.Scrollbar(cc_emails_frame, orient="vertical", command=self.cc_emails_listbox.yview)
        self.cc_emails_listbox.configure(yscrollcommand=cc_emails_scroll.set)
        self.cc_emails_listbox.pack(side="left", fill="both", expand=True)
        cc_emails_scroll.pack(side="right", fill="y")

        email_controls = ttk.Frame(right)
        email_controls.pack(fill="x")
        self.cc_email_var = tk.StringVar()
        ttk.Entry(email_controls, textvariable=self.cc_email_var, width=50).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Add Email", command=self._add_cc_email).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Remove Selected", command=self._remove_cc_email).pack(side="left")

        self._refresh_cc_listbox()

    def _refresh_clients(self):
        clients = sorted({row["client"] for row in self.dropbox_data})
        self.client_combo["values"] = clients
        if clients:
            self.client_var.set(clients[0])
            self._on_client_change()
        self._refresh_cc_lists()
        self._refresh_signature_list()

    def _refresh_cc_lists(self):
        list_names = sorted(self.cc_lists.keys())
        self.cc_list_combo["values"] = list_names
        default_name = "Commercial Construction"
        if default_name in self.cc_lists:
            self.cc_list_var.set(default_name)
        elif list_names:
            self.cc_list_var.set(list_names[0])
        else:
            self.cc_list_var.set("")
        self._on_cc_list_change()

    def _on_cc_list_change(self):
        list_name = self.cc_list_var.get()
        emails = self.cc_lists.get(list_name, [])
        self.cc_var.set("; ".join(emails))

    def _refresh_signature_list(self):
        if not hasattr(self, "signature_combo"):
            return
        names = self.signature_names
        self.signature_combo["values"] = names
        preferred = None
        for name in names:
            if name == SIGNATURE_NAME:
                preferred = name
                break
        if not preferred:
            for name in names:
                if name.startswith(SIGNATURE_NAME):
                    preferred = name
                    break
        if preferred:
            self.signature_var.set(preferred)
        elif names:
            self.signature_var.set(names[0])
        else:
            self.signature_var.set("")

    def _toggle_bcc(self):
        if self.show_bcc_var.get():
            self.bcc_frame.pack(anchor="w", pady=(0, 12))
        else:
            self.bcc_frame.pack_forget()
            self.bcc_var.set("")

    def _refresh_cc_listbox(self):
        if not hasattr(self, "cc_listbox"):
            return
        self.cc_listbox.delete(0, "end")
        for name in sorted(self.cc_lists.keys()):
            self.cc_listbox.insert("end", name)
        self._load_cc_list_selection()

    def _load_cc_list_selection(self):
        if not hasattr(self, "cc_listbox"):
            return
        selection = self.cc_listbox.curselection()
        if not selection:
            self.cc_emails_listbox.delete(0, "end")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        self.cc_emails_listbox.delete(0, "end")
        for email in emails:
            self.cc_emails_listbox.insert("end", email)

    def _create_cc_list(self):
        name = self.new_cc_list_var.get().strip()
        if not name:
            messagebox.showwarning("Missing Data", "List name is required.")
            return
        if name in self.cc_lists:
            messagebox.showwarning("Already Exists", "That list already exists.")
            return
        self.cc_lists[name] = []
        self._write_cc_lists()
        self.new_cc_list_var.set("")
        self._reload_cc_lists_data()

    def _delete_cc_list(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to delete.")
            return
        list_name = self.cc_listbox.get(selection[0])
        if not messagebox.askyesno("Confirm Delete", f"Delete CC list '{list_name}'?"):
            return
        self.cc_lists.pop(list_name, None)
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _add_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to add to.")
            return
        email = self.cc_email_var.get().strip()
        if not email:
            messagebox.showwarning("Missing Data", "Email is required.")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        if email in emails:
            messagebox.showwarning("Already Exists", "That email is already in the list.")
            return
        emails.append(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self.cc_email_var.set("")
        self._reload_cc_lists_data()

    def _remove_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to remove from.")
            return
        list_name = self.cc_listbox.get(selection[0])
        email_selection = self.cc_emails_listbox.curselection()
        if not email_selection:
            messagebox.showwarning("Select Email", "Choose an email to remove.")
            return
        emails = self.cc_lists.get(list_name, [])
        for idx in reversed(email_selection):
            email = self.cc_emails_listbox.get(idx)
            if email in emails:
                emails.remove(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _reload_cc_lists_data(self):
        self.cc_lists = self._read_cc_lists()
        self._refresh_cc_listbox()
        self._refresh_cc_lists()

    def _on_client_change(self):
        client = self.client_var.get()
        sites = sorted({row["site"] for row in self.dropbox_data if row["client"] == client})
        self.site_combo["values"] = sites
        if sites:
            self.site_var.set(sites[0])
        else:
            self.site_var.set("")
        self._on_site_change()

    def _on_site_change(self):
        client = self.client_var.get()
        site = self.site_var.get()

        link = ""
        for row in self.dropbox_data:
            if row["client"] == client and row["site"] == site:
                link = row["link"]
                break
        if not link:
            c_key = _normalize_key(client)
            s_key = _normalize_key(site)
            for row in self.dropbox_data:
                if _normalize_key(row["client"]) == c_key and _normalize_key(row["site"]) == s_key:
                    link = row["link"]
                    break
        self.link_var.set(link)

        email = ""
        for row in self.contacts_data:
            if row["client"] == client and row["site"] == site:
                email = row["email"]
                break
        if not email:
            c_key = _normalize_key(client)
            s_key = _normalize_key(site)
            for row in self.contacts_data:
                if _normalize_key(row["client"]) == c_key and _normalize_key(row["site"]) == s_key:
                    email = row["email"]
                    break
        if email:
            self.to_var.set(email)

    def _set_default_date(self):
        self.selected_date = default_date()
        self.date_var.set(format_date(self.selected_date))

    def _open_calendar(self):
        CalendarPopup(self.root, self.selected_date, self._set_date)

    def _set_date(self, date_obj):
        self.selected_date = date_obj
        self.date_var.set(format_date(date_obj))
        self._update_email_preview()

    def _compose_subject(self):
        client = self.client_var.get().strip()
        site = self.site_var.get().strip()
        media = MEDIA_SUBJECT.get(self.media_var.get(), self.media_var.get())
        date_text = self.date_var.get().strip()
        template = self.email_templates.get(self.signature_var.get().strip())
        if template and template.get("subject"):
            try:
                return template["subject"].format(
                    client=client,
                    site=site,
                    media_subject=media,
                    media_body=MEDIA_BODY.get(self.media_var.get(), self.media_var.get()),
                    date=date_text,
                    dropbox_link=self.link_var.get().strip(),
                    to_email=self.to_var.get().strip(),
                    cc_email=self.cc_var.get().strip(),
                )
            except Exception:
                pass
        if site and media and date_text:
            return f"{site} - Aerial Progress {media} - {date_text}"
        return ""

    def _compose_body(self):
        client = self.client_var.get().strip()
        site = self.site_var.get().strip()
        media_body = MEDIA_BODY.get(self.media_var.get(), self.media_var.get())
        date_text = self.date_var.get().strip()
        link = self.link_var.get().strip()
        template = self.email_templates.get(self.signature_var.get().strip())
        if template and template.get("text"):
            try:
                rendered = template["text"].format(
                    client=client,
                    site=site,
                    media_subject=MEDIA_SUBJECT.get(self.media_var.get(), self.media_var.get()),
                    media_body=media_body,
                    date=date_text,
                    dropbox_link=link,
                    to_email=self.to_var.get().strip(),
                    cc_email=self.cc_var.get().strip(),
                )
                return rendered.replace("\n", "\r\n")
            except Exception:
                pass
        greeting = f"Good Day {client} Team,"
        main = (
            f"Our pilots captured progress {media_body} of the {site} site on {date_text}. "
            "Please use the link below to access your imagery."
        )
        body = (
            f"{greeting}\r\n\r\n"
            f"{main}\r\n"
            f"Dropbox Link: {link}\r\n\r\n"
            "Let us know if you have any questions, feedback, or if we can further assist you.\r\n\r\n"
        )
        return body

    def _compose_body_html(self):
        client = self.client_var.get().strip()
        site = self.site_var.get().strip()
        media_body = MEDIA_BODY.get(self.media_var.get(), self.media_var.get())
        date_text = self.date_var.get().strip()
        link = self.link_var.get().strip()
        template = self.email_templates.get(self.signature_var.get().strip())
        if template and template.get("text"):
            try:
                rendered = template["text"].format(
                    client=client,
                    site=site,
                    media_subject=MEDIA_SUBJECT.get(self.media_var.get(), self.media_var.get()),
                    media_body=media_body,
                    date=date_text,
                    dropbox_link=link,
                    to_email=self.to_var.get().strip(),
                    cc_email=self.cc_var.get().strip(),
                )
                def _text_to_html(text):
                    from html import escape
                    paragraphs = []
                    current = []
                    for line in text.splitlines():
                        if line.strip() == "":
                            if current:
                                paragraphs.append("<br>".join(current))
                                current = []
                            continue
                        escaped = escape(line)
                        if "Dropbox Link:" in line and link:
                            escaped = escaped.replace(escape(link), f'<a href="{escape(link)}">{escape(link)}</a>')
                        current.append(escaped)
                    if current:
                        paragraphs.append("<br>".join(current))
                    html_parts = [f"<p style=\"margin:0 0 10px 0;\">{p}</p>" for p in paragraphs]
                    return (
                        "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; line-height:1.35;\">"
                        + "".join(html_parts)
                        + "</div>"
                    )

                signature_html, signature_assets = self._load_signature_html(self.signature_var.get().strip() or None)
                body_html = _text_to_html(rendered)
                if signature_html:
                    return body_html + signature_html, signature_assets
                return body_html, []
            except Exception:
                pass
        signature_html, signature_assets = self._load_signature_html(self.signature_var.get().strip() or None)
        link_html = f'<a href="{link}">{link}</a>' if link else ""

        body_html = (
            "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; "
            "line-height:1.35;\">"
            f"<p style=\"margin:0 0 10px 0;\">Good Day {client} Team,</p>"
            f"<p style=\"margin:0 0 8px 0;\">Our pilots captured progress {media_body} of the {site} site on {date_text}. "
            "Please use the link below to access your imagery.</p>"
            f"<p style=\"margin:0 0 10px 0;\"><b>Dropbox Link:</b> {link_html}</p>"
            "<p style=\"margin:0 0 12px 0;\">Let us know if you have any questions, feedback, or if we can further assist you.</p>"
            ""
            "</div>"
        )

        if signature_html:
            return body_html + signature_html, signature_assets
        return body_html, []

    def _update_email_preview(self):
        subject = self._compose_subject()
        body = self._compose_body()
        if subject:
            self.subject_var.set(subject)
        self.preview.config(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("1.0", body)
        self.preview.config(state="disabled")

    def _create_draft(self):
        if win32 is None:
            messagebox.showerror("Outlook", "pywin32 is not available. Please install pywin32 to create drafts.")
            return

        subject = self.subject_var.get().strip()
        try:
            body_html, signature_assets = self._compose_body_html()
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to build email HTML: {exc}")
            return

        to_line = self.to_var.get().strip().replace(",", ";")
        cc_line = self.cc_var.get().strip().replace(",", ";")
        bcc_line = self.bcc_var.get().strip().replace(",", ";")
        if not subject or not to_line:
            messagebox.showwarning("Missing Data", "Please ensure a client, site, and To address are selected.")
            return

        try:
            outlook = win32.Dispatch("Outlook.Application")
            mail = outlook.CreateItem(0)
            mail.To = to_line
            if cc_line:
                mail.CC = cc_line
            if bcc_line:
                mail.BCC = bcc_line
            mail.Subject = subject
            mail.BodyFormat = 2  # olFormatHTML
            mail.HTMLBody = body_html
            for asset in signature_assets:
                try:
                    attachment = mail.Attachments.Add(asset["path"])
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3712001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3713001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x7FFE000B",
                        True,
                    )
                except Exception as exc:
                    continue
            # Re-apply HTMLBody after inline attachments to help Outlook resolve CIDs
            mail.HTMLBody = body_html
            mail.Save()  # draft in Drafts
            try:
                mapi = outlook.GetNamespace("MAPI")
                mapi.SendAndReceive(True)
            except Exception:
                pass
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to create Outlook draft: {exc}")

    def _refresh_dropbox_tree(self):
        self.dropbox_tree.delete(*self.dropbox_tree.get_children())
        for row in self.dropbox_data:
            self.dropbox_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["link"]))

    def _load_dropbox_selection(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.dropbox_data if r["row"] == row_id), None)
        if not row:
            return
        self.db_client_var.set(row["client"])
        self.db_site_var.set(row["site"])
        self.db_link_var.set(row["link"])

    def _add_dropbox_entry(self):
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.append([client, site, None, link])
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_dropbox_data()

    def _update_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.cell(row=row_id, column=1, value=client)
        ws.cell(row=row_id, column=2, value=site)
        ws.cell(row=row_id, column=3, value=link)
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_dropbox_data()

    def _delete_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected Dropbox entry?"):
            return
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.delete_rows(row_id, 1)
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_dropbox_data()

    def _reload_dropbox_data(self):
        self.dropbox_data = self._read_dropbox_data()
        self._refresh_dropbox_tree()
        self._refresh_clients()

    def _refresh_contacts_tree(self):
        self.contacts_tree.delete(*self.contacts_tree.get_children())
        for row in self.contacts_data:
            self.contacts_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["email"]))

    def _load_contact_selection(self):
        selected = self.contacts_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.contacts_data if r["row"] == row_id), None)
        if not row:
            return
        self.ct_client_var.set(row["client"])
        self.ct_site_var.set(row["site"])
        self.ct_email_var.set(row["email"])

    def _add_contact_entry(self):
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = self.ct_email_var.get().strip()
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.append([client, site, None, email])
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_contacts_data()

    def _update_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = self.ct_email_var.get().strip()
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.cell(row=row_id, column=1, value=client)
        ws.cell(row=row_id, column=2, value=site)
        ws.cell(row=row_id, column=4, value=email)
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_contacts_data()

    def _delete_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected contact entry?"):
            return
        wb = openpyxl.load_workbook(MASTER_FILE)
        ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb.active
        ws.delete_rows(row_id, 1)
        wb.save(MASTER_FILE)
        wb.close()
        self._reload_contacts_data()

    def _reload_contacts_data(self):
        self.contacts_data = self._read_contacts_data()
        self._refresh_contacts_tree()


if __name__ == "__main__":
    try:
        root = ttk.Window()
    except Exception:
        root = tk.Tk()
    app = EmailTemplateApp(root)
    root.mainloop()
