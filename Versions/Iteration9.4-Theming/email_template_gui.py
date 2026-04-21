import datetime as _dt
import calendar as _cal
import mimetypes
import uuid
import re
import sys
import logging
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from urllib.parse import unquote
from pathlib import Path
import tkinter as tk
from tkinter import messagebox

try:
    import ttkbootstrap as ttk
except Exception:  # pragma: no cover
    from tkinter import ttk

import openpyxl

try:
    import win32com.client as win32
except Exception:  # pragma: no cover
    win32 = None

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "2026_02_03"
MASTER_FILE = DATA_DIR / "Client-Site-URL-Email.xlsx"
MASTER_SHEET = "Construction"
LEGACY_MASTER_SHEET = "ABCD"
OIL_GAS_SHEET = "Oil & Gas"
CC_LISTS_FILE = BASE_DIR / "Internal CC Lists.xlsx"
CC_LISTS_SHEET = "Sheet1"
EMAIL_TEMPLATES_DIR = BASE_DIR / "Email Templates"
GRAPH_DIR = BASE_DIR / "Backend Graph"
GRAPH_SETTINGS = GRAPH_DIR / "graph_app_settings.json"
GRAPH_TOKEN_CACHE = GRAPH_DIR / ".graph_token_cache.json"

logger = logging.getLogger(__name__)

if GRAPH_DIR.exists():
    sys.path.insert(0, str(GRAPH_DIR))
try:
    from graph_client import GraphEmailClient  # type: ignore
except Exception:
    GraphEmailClient = None

MEDIA_OPTIONS = [
    "photos",
    "photos and video",
    "photos, video, and 2D mapping",
    "photos, and 2D mapping",
    "ROW Documentation Video",
]

MEDIA_SUBJECT = {
    "photos": "Aerial Progress Photos",
    "photos and video": "Aerial Progress Photos and Video",
    "photos, video, and 2D mapping": "Aerial Progress Photos, Video, and 2D Mapping",
    "photos, and 2D mapping": "Aerial Progress Photos and 2D Mapping",
    "ROW Documentation Video": "ROW Documentation Video",
}

MEDIA_BODY = {
    "photos": "progress photos",
    "photos and video": "progress photos and video",
    "photos, video, and 2D mapping": "progress photos, video, and 2D mapping",
    "photos, and 2D mapping": "progress photos and 2D mapping",
    "ROW Documentation Video": "ROW documentation video",
}


def _ordinal(n: int) -> str:
    if 10 <= (n % 100) <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def _normalize_key(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    cleaned = "".join(ch.lower() if ch.isalnum() else " " for ch in value)
    tokens = cleaned.split()
    expanded = []
    for tok in tokens:
        if tok == "hs":
            expanded.append("high")
            expanded.append("school")
        elif tok == "ms":
            expanded.append("middle")
            expanded.append("school")
        elif tok == "es":
            expanded.append("elementary")
            expanded.append("school")
        elif tok == "jh":
            expanded.append("junior")
            expanded.append("high")
        elif tok == "jhs":
            expanded.append("junior")
            expanded.append("high")
            expanded.append("school")
        elif tok == "jr":
            expanded.append("junior")
        elif tok == "jrhigh":
            expanded.append("junior")
            expanded.append("high")
        elif tok == "jrhighschool":
            expanded.append("junior")
            expanded.append("high")
            expanded.append("school")
        else:
            expanded.append(tok)
    return " ".join(expanded)


EMAIL_ADDRESS_RE = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@(?:[A-Za-z0-9-]+\.)+[A-Za-z]{2,}")


def _sheet_aliases(sheet_name: str):
    if sheet_name == MASTER_SHEET:
        return [MASTER_SHEET, LEGACY_MASTER_SHEET]
    return [sheet_name]


def normalize_email_list(raw) -> str:
    if raw is None:
        return ""
    if not isinstance(raw, str):
        raw = str(raw)

    matches = EMAIL_ADDRESS_RE.findall(raw)
    if matches:
        deduped = []
        seen = set()
        for email in matches:
            key = email.lower()
            if key in seen:
                continue
            seen.add(key)
            deduped.append(email)
        return ", ".join(deduped)

    cleaned = re.sub(r"[\r\n\t;]+", ",", raw)
    parts = [part.strip() for part in cleaned.split(",") if part.strip()]
    return ", ".join(parts)



def format_date(date_obj: _dt.date) -> str:
    return date_obj.strftime(f"%B {_ordinal(date_obj.day)}, %Y")


def default_date() -> _dt.date:
    today = _dt.date.today()
    if today.weekday() == 0:  # Monday
        return today - _dt.timedelta(days=3)
    return today - _dt.timedelta(days=1)


class CalendarPopup(tk.Toplevel):
    def __init__(self, master, initial_date, on_select):
        super().__init__(master)
        self.title("Select Date")
        self.resizable(False, False)
        self.on_select = on_select
        self.year = initial_date.year
        self.month = initial_date.month
        self.calendar = _cal.Calendar(firstweekday=_cal.SUNDAY)

        header = ttk.Frame(self)
        header.pack(padx=10, pady=8, fill="x")

        ttk.Button(header, text="<", width=3, command=self._prev_month).pack(side="left")
        self.month_label = ttk.Label(header, text="")
        self.month_label.pack(side="left", expand=True)
        ttk.Button(header, text=">", width=3, command=self._next_month).pack(side="right")

        self.grid_frame = ttk.Frame(self)
        self.grid_frame.pack(padx=10, pady=8)

        self._render()

    def _prev_month(self):
        if self.month == 1:
            self.month = 12
            self.year -= 1
        else:
            self.month -= 1
        self._render()

    def _next_month(self):
        if self.month == 12:
            self.month = 1
            self.year += 1
        else:
            self.month += 1
        self._render()

    def _render(self):
        for child in self.grid_frame.winfo_children():
            child.destroy()

        self.month_label.config(text=f"{_cal.month_name[self.month]} {self.year}")

        for idx, day in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
            ttk.Label(self.grid_frame, text=day, width=4, anchor="center").grid(row=0, column=idx, padx=2, pady=2)

        month_data = self.calendar.monthdayscalendar(self.year, self.month)
        for r, week in enumerate(month_data, start=1):
            for c, day in enumerate(week):
                if day == 0:
                    ttk.Label(self.grid_frame, text="", width=4).grid(row=r, column=c, padx=2, pady=2)
                    continue
                ttk.Button(
                    self.grid_frame,
                    text=str(day),
                    width=4,
                    command=lambda d=day: self._select_day(d),
                ).grid(row=r, column=c, padx=2, pady=2)

    def _select_day(self, day):
        chosen = _dt.date(self.year, self.month, day)
        self.on_select(chosen)
        self.destroy()


class EmailTemplateApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Customer Delivery Email Builder")
        self.root.geometry("1215x927")

        self.master_data = []
        self.dropbox_data = []
        self.contacts_data = []
        self.sheet_data = {}
        self.client_site_tabs = {}
        self.compose_sheet_name = MASTER_SHEET
        self.cc_lists = {}
        self.signature_names = []
        self.template_folder_names = []
        self.active_template_folder_name = ""
        self.email_templates = {}
        self._undo_stack = []
        self.graph_client = None

        self._apply_theme()
        self._load_data()
        self._init_graph_client()
        self._build_ui()
        self._refresh_clients()
        self._set_default_date()
        self._update_email_preview()

    def _load_data(self):
        self.sheet_data = {
            MASTER_SHEET: self._read_master_data(MASTER_SHEET),
            OIL_GAS_SHEET: self._read_master_data(OIL_GAS_SHEET, allow_missing=True),
        }
        self._sync_primary_master_views()
        self.cc_lists = self._read_cc_lists()
        self.signature_names = self._read_signature_names()
        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()
        self.email_templates = self._read_email_templates()

    def _init_graph_client(self):
        if GraphEmailClient is None:
            return
        if not GRAPH_SETTINGS.exists():
            return
        try:
            self.graph_client = GraphEmailClient(GRAPH_SETTINGS, GRAPH_TOKEN_CACHE)
        except Exception as exc:
            logger.warning("Graph client not initialized: %s", exc)
            self.graph_client = None

    def _read_template_folder_names(self):
        if not EMAIL_TEMPLATES_DIR.exists():
            return []
        return sorted(path.name for path in EMAIL_TEMPLATES_DIR.iterdir() if path.is_dir())

    def _choose_active_template_folder(self):
        names = list(self.template_folder_names)
        if self.active_template_folder_name in names:
            return self.active_template_folder_name
        return names[0] if names else ""

    def _selected_template_dir(self):
        folder_name = (self.active_template_folder_name or "").strip()
        if folder_name:
            path = EMAIL_TEMPLATES_DIR / folder_name
            if path.exists() and path.is_dir():
                return path
        return EMAIL_TEMPLATES_DIR

    def _read_email_templates(self):
        templates = {}
        template_dir = self._selected_template_dir()
        if not template_dir.exists():
            return templates
        for path in template_dir.glob("*.eml"):
            name = path.stem
            try:
                msg = BytesParser(policy=policy.default).parse(path.open("rb"))
            except Exception:
                continue
            subject = msg.get("Subject", "")
            text = ""
            html = ""
            if msg.is_multipart():
                for part in msg.walk():
                    ctype = part.get_content_type()
                    if ctype == "text/plain" and not text:
                        text = part.get_content()
                    elif ctype == "text/html" and not html:
                        html = part.get_content()
            else:
                ctype = msg.get_content_type()
                if ctype == "text/plain":
                    text = msg.get_content()
                elif ctype == "text/html":
                    html = msg.get_content()
            templates[name] = {
                "subject": subject or "",
                "text": text or "",
                "html": html or "",
                "path": path,
            }
        return templates

    def _resolve_sheet_name(self, wb, sheet_name: str, allow_missing: bool = False):
        for candidate in _sheet_aliases(sheet_name):
            if candidate in wb.sheetnames:
                return candidate
        if allow_missing:
            return None
        if sheet_name == MASTER_SHEET:
            expected = " or ".join(f'"{name}"' for name in _sheet_aliases(sheet_name))
            raise KeyError(f"Workbook is missing the Construction sheet. Expected {expected}.")
        raise KeyError(f'Workbook is missing the "{sheet_name}" sheet.')

    def _get_or_create_sheet(self, wb, sheet_name: str):
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=True)
        if resolved_sheet:
            return wb[resolved_sheet]

        ws = wb.create_sheet(sheet_name)
        source_sheet_name = self._resolve_sheet_name(wb, MASTER_SHEET, allow_missing=True)
        headers = ["Client", "Project Name", "Dropbox URLs", "Operations Contact Email"]
        if source_sheet_name:
            source_headers = [cell.value for cell in next(wb[source_sheet_name].iter_rows(min_row=1, max_row=1))]
            if any(source_headers):
                headers = source_headers
        ws.append(headers)
        return ws

    def _sync_primary_master_views(self):
        self.master_data = list(self.sheet_data.get(MASTER_SHEET, []))
        self.dropbox_data = [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self.master_data
        ]
        self.contacts_data = [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self.master_data
        ]

    def _compose_rows(self):
        return self.sheet_data.get(self.compose_sheet_name, [])

    def _compose_dropbox_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self._compose_rows()
        ]

    def _compose_contacts_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self._compose_rows()
        ]

    def _sheet_for_cc_list(self, list_name: str) -> str:
        normalized = (list_name or "").strip().lower()
        if normalized == OIL_GAS_SHEET.lower():
            return OIL_GAS_SHEET
        return MASTER_SHEET

    def _read_master_data(self, sheet_name: str = MASTER_SHEET, allow_missing: bool = False):
        wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=allow_missing)
        if resolved_sheet is None:
            wb.close()
            return []
        ws = wb[resolved_sheet]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h}
        client_idx = col_map.get("client", 0)
        site_idx = col_map.get("project name", 1)
        link_idx = col_map.get("dropbox urls", 2)
        email_idx = col_map.get("operations contact email", 3)
        rows = []
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            client = (row[client_idx] or "").strip() if isinstance(row[client_idx], str) else row[client_idx]
            site = (row[site_idx] or "").strip() if isinstance(row[site_idx], str) else row[site_idx]
            link = (row[link_idx] or "").strip() if isinstance(row[link_idx], str) else row[link_idx]
            email = (row[email_idx] or "").strip() if isinstance(row[email_idx], str) else row[email_idx]
            if client:
                rows.append({
                    "row": idx,
                    "client": client,
                    "site": site,
                    "link": link or "",
                    "email": normalize_email_list(email),
                })
        wb.close()
        return rows

    def _sort_master_sheet(self, sheet_name: str = MASTER_SHEET):
        wb = openpyxl.load_workbook(MASTER_FILE)
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=(sheet_name == OIL_GAS_SHEET))
        if resolved_sheet is None:
            wb.close()
            return
        ws = wb[resolved_sheet]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # Preserve rows but sort by Client then Project Name
        def _key(row):
            client = row[0] or ""
            site = row[1] or ""
            return (str(client).strip().lower(), str(site).strip().lower())
        rows_sorted = sorted(rows, key=_key)
        # Clear existing data rows
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        # Write sorted rows back
        for r in rows_sorted:
            ws.append(list(r))
        wb.save(MASTER_FILE)
        wb.close()

    def _read_dropbox_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self._read_master_data()
        ]

    def _read_contacts_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self._read_master_data()
        ]

    def _read_cc_lists(self):
        wb = openpyxl.load_workbook(CC_LISTS_FILE, data_only=True)
        ws = wb[CC_LISTS_SHEET]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        lists = {}
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            emails = []
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
                email = row[0]
                if isinstance(email, str):
                    email = email.strip()
                if email:
                    emails.append(email)
            lists[str(header).strip()] = emails
        wb.close()
        return lists

    def _read_signature_names(self):
        return []

    def _extract_inline_images(self, html, base_dir=None):
        assets = []
        # Handle relative file references from Outlook signature folders
        def _replace_file_src(match):
            quote = match.group(1)
            src = match.group(2)
            lower = src.lower()
            if lower.startswith("cid:") or lower.startswith("data:") or lower.startswith("http"):
                return match.group(0)
            if not base_dir:
                return match.group(0)
            rel = unquote(src)
            path = (Path(base_dir) / rel).resolve()
            if not path.exists():
                return match.group(0)
            cid = f"sig-{uuid.uuid4().hex}"
            assets.append({"cid": cid, "path": str(path)})
            return f'src={quote}cid:{cid}{quote}'

        # img src="..."; also handles v:imagedata src="..."
        html = re.sub(
            r'src=(["\'])([^"\']+)\1',
            _replace_file_src,
            html,
            flags=re.IGNORECASE,
        )
        return html, assets

    def _write_cc_lists(self):
        self._record_undo(CC_LISTS_FILE)
        if CC_LISTS_FILE.exists():
            wb = openpyxl.load_workbook(CC_LISTS_FILE)
        else:
            wb = openpyxl.Workbook()
        if CC_LISTS_SHEET in wb.sheetnames:
            ws_old = wb[CC_LISTS_SHEET]
            wb.remove(ws_old)
        ws = wb.create_sheet(CC_LISTS_SHEET, 0)
        headers = list(self.cc_lists.keys())
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)
            for row_idx, email in enumerate(self.cc_lists[header], start=2):
                ws.cell(row=row_idx, column=col_idx, value=email)
        wb.save(CC_LISTS_FILE)
        wb.close()

    def _record_undo(self, path):
        try:
            p = Path(path)
            if not p.exists():
                return
            self._undo_stack.append((p, p.read_bytes()))
            if len(self._undo_stack) > 1:
                self._undo_stack = self._undo_stack[-1:]
        except Exception:
            pass

    def _undo_last_change(self):
        if not self._undo_stack:
            if hasattr(self, "status_var"):
                self.status_var.set("Nothing to undo.")
            else:
                messagebox.showinfo("Undo", "Nothing to undo.")
            return
        path, data = self._undo_stack.pop()
        try:
            path.write_bytes(data)
            if hasattr(self, "status_var"):
                self.status_var.set(f"Undid last change: {path.name}")
            self._reload_master_data(sheet_name=MASTER_SHEET)
            self._reload_master_data(sheet_name=OIL_GAS_SHEET, refresh_clients=False)
            self._reload_cc_lists_data()
        except Exception as exc:
            messagebox.showerror("Undo Error", f"Failed to undo last change: {exc}")

    def _build_ui(self):
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=1)

        container = ttk.Frame(self.root)
        container.grid(row=0, column=0, sticky="nsew")
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(1, weight=1)

        sidebar = ttk.Frame(container, style="Sidebar.TFrame", width=220)
        sidebar.grid(row=0, column=0, sticky="ns")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(1, weight=1)

        content = ttk.Frame(container)
        content.grid(row=0, column=1, sticky="nsew")
        content.grid_rowconfigure(1, weight=1)
        content.grid_columnconfigure(0, weight=1)

        header = tk.Frame(content, bg=self.colors["brand"], highlightbackground=self.colors["border"], highlightthickness=1)
        header.grid(row=0, column=0, sticky="ew")
        tk.Label(
            header,
            text="Arch Aerial, LLC || Customer Delivery Email Builder",
            bg=self.colors["brand"],
            fg=self.colors["text"],
            font=("Segoe UI Semibold", 16),
            padx=18,
            pady=14,
        ).pack(side="left")
        tk.Label(
            header,
            text="Local",
            bg=self.colors["primary"],
            fg=self.colors["text"],
            font=("Segoe UI Semibold", 11),
            padx=12,
            pady=6,
        ).pack(side="right", padx=18, pady=10)

        tab_container = ttk.Frame(content)
        tab_container.grid(row=1, column=0, sticky="nsew")

        top_nav = ttk.Frame(sidebar, style="Sidebar.TFrame")
        top_nav.grid(row=0, column=0, sticky="ew", padx=12, pady=12)

        bottom_nav = ttk.Frame(sidebar, style="Sidebar.TFrame")
        bottom_nav.grid(row=2, column=0, sticky="ew", padx=12, pady=12)

        self.compose_tab = ttk.Frame(tab_container)
        self.master_tab = ttk.Frame(tab_container)
        self.oil_gas_tab = ttk.Frame(tab_container)
        self.cc_tab = ttk.Frame(tab_container)
        self.about_tab = ttk.Frame(tab_container)

        self.tabs = {
            "compose": self.compose_tab,
            "construction": self.master_tab,
            "oil_gas": self.oil_gas_tab,
            "cc": self.cc_tab,
            "about": self.about_tab,
        }
        for frame in self.tabs.values():
            frame.place(relx=0, rely=0, relwidth=1, relheight=1)

        self.nav_buttons = {}
        nav_specs = [
            ("compose", "Compose", top_nav),
            ("construction", "Construction", top_nav),
            ("oil_gas", "Oil & Gas", top_nav),
            ("cc", "Internal CC Lists", top_nav),
            ("about", "About", bottom_nav),
        ]
        for tab_name, label, parent in nav_specs:
            button = ttk.Button(
                parent,
                text=label,
                style="Nav.TButton",
                command=lambda name=tab_name: self._show_tab(name),
            )
            button.pack(fill="x", pady=4)
            self.nav_buttons[tab_name] = button

        self._build_compose_tab()
        self._build_master_tab()
        self._build_oil_gas_tab()
        self._build_cc_tab()
        self._build_about_tab()

        self.status_var = tk.StringVar(value="")
        status = ttk.Label(self.root, textvariable=self.status_var, style="Muted.TLabel", anchor="w")
        status.grid(row=1, column=0, sticky="ew")

        self._show_tab("compose")

    def _apply_theme(self):
        style = ttk.Style()
        try:
            style.theme_use("darkly")
        except Exception:
            try:
                style.theme_use("clam")
            except tk.TclError:
                pass

        # Match the pyweb-local (Docker Desktop-inspired) palette
        colors = {
            "bg": "#1a1a1a",
            "panel": "#2a2a2a",
            "sidebar": "#1e1e1e",
            "brand": "#012458",
            "border": "#404040",
            "text": "#ffffff",
            "text_secondary": "#b3b3b3",
            "text_muted": "#858585",
            "primary": "#314dec",
            "info": "#3a7ca5",
            "success": "#10b981",
            "warning": "#f59e0b",
            "error": "#ef4444",
        }
        self.colors = colors

        self.root.configure(bg=colors["bg"])

        style.configure("TFrame", background=colors["bg"])
        style.configure("Card.TFrame", background=colors["panel"])
        style.configure("Sidebar.TFrame", background=colors["sidebar"])
        style.configure("Panel.TFrame", background=colors["panel"])
        style.configure("TLabel", background=colors["bg"], foreground=colors["text"], font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=colors["bg"], foreground=colors["text_secondary"])
        style.configure("Panel.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI", 10))
        style.configure("SectionTitle.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI Semibold", 10))
        style.configure("CardTitle.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI Semibold", 11))

        style.configure("TNotebook", background=colors["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", padding=(14, 10), background=colors["bg"], foreground=colors["text_secondary"])
        style.map("TNotebook.Tab", background=[("selected", colors["panel"])], foreground=[("selected", colors["text"])])

        style.configure("TEntry", fieldbackground="#3a3a3a", background="#3a3a3a", foreground=colors["text"], bordercolor=colors["border"])
        style.configure("TCombobox", fieldbackground="#3a3a3a", background="#3a3a3a", foreground=colors["text"], bordercolor=colors["border"])

        style.configure("TButton", background=colors["panel"], foreground=colors["text"], padding=(10, 6))
        style.map(
            "TButton",
            background=[("active", colors["primary"])],
            foreground=[("active", colors["text"])],
        )
        style.configure("Accent.TButton", background=colors["primary"], foreground=colors["text"], padding=(12, 8))
        style.map("Accent.TButton", background=[("active", colors["brand"])], foreground=[("active", colors["text"])])
        style.configure(
            "Nav.TButton",
            background=colors["sidebar"],
            foreground=colors["text_secondary"],
            padding=(12, 10),
            anchor="w",
            relief="flat",
        )
        style.map(
            "Nav.TButton",
            background=[("active", colors["panel"])],
            foreground=[("active", colors["text"])],
        )
        style.configure(
            "ActiveNav.TButton",
            background=colors["primary"],
            foreground=colors["text"],
            padding=(12, 10),
            anchor="w",
            relief="flat",
        )
        try:
            style.configure(
                "TScrollbar",
                background=colors["panel"],
                troughcolor=colors["sidebar"],
                bordercolor=colors["border"],
                arrowcolor=colors["text_secondary"],
                relief="flat",
            )
            style.map(
                "TScrollbar",
                background=[("active", colors["primary"])],
                arrowcolor=[("active", colors["text"])],
            )
        except tk.TclError:
            pass

        style.configure(
            "Treeview",
            background=colors["panel"],
            fieldbackground=colors["panel"],
            foreground=colors["text"],
            bordercolor=colors["border"],
            borderwidth=1,
        )
        style.configure(
            "Treeview.Heading",
            background=colors["sidebar"],
            foreground=colors["text_secondary"],
            font=("Segoe UI", 9, "bold"),
        )
        style.map("Treeview", background=[("selected", "#2f3d5a")], foreground=[("selected", colors["text"])])
        try:
            style.configure("Vertical.TScrollbar", arrowsize=14)
        except tk.TclError:
            pass
        try:
            style.configure(
                "ClientSites.Vertical.TScrollbar",
                background="#555555",
                troughcolor="#444444",
                bordercolor="#444444",
                arrowcolor=colors["text_secondary"],
                relief="flat",
            )
            style.map(
                "ClientSites.Vertical.TScrollbar",
                background=[("active", "#555555")],
                arrowcolor=[("active", colors["text"])],
            )
        except tk.TclError:
            pass

        self._theme_colors = colors

    def _create_card(self, parent, title: str, subtitle: str | None = None):
        card = tk.Frame(
            parent,
            bg=self.colors["panel"],
            highlightbackground=self.colors["border"],
            highlightthickness=1,
            bd=0,
        )
        header = tk.Frame(card, bg=self.colors["brand"], bd=0)
        header.pack(fill="x")
        tk.Label(
            header,
            text=title,
            bg=self.colors["brand"],
            fg=self.colors["text"],
            font=("Segoe UI Semibold", 14),
            padx=16,
            pady=12,
        ).pack(side="left")
        if subtitle:
            tk.Label(
                header,
                text=subtitle,
                bg=self.colors["brand"],
                fg="#d5dcf7",
                font=("Segoe UI", 10),
                padx=16,
                pady=12,
            ).pack(side="right")
        body = ttk.Frame(card, style="Panel.TFrame")
        body.pack(fill="both", expand=True, padx=16, pady=16)
        return card, body

    def _style_text_widget(self, widget):
        widget.configure(
            bg="#303030",
            fg=self.colors["text"],
            insertbackground=self.colors["text"],
            selectbackground=self.colors["primary"],
            selectforeground=self.colors["text"],
            highlightbackground=self.colors["border"],
            highlightcolor=self.colors["primary"],
            relief="flat",
            borderwidth=1,
        )

    def _style_listbox_widget(self, widget):
        widget.configure(
            bg="#303030",
            fg=self.colors["text"],
            selectbackground=self.colors["primary"],
            selectforeground=self.colors["text"],
            highlightbackground=self.colors["border"],
            highlightcolor=self.colors["primary"],
            relief="flat",
            borderwidth=1,
        )


    def _build_compose_tab(self):
        wrap = ttk.Frame(self.compose_tab)
        wrap.pack(fill="both", expand=True, padx=16, pady=16)

        left_card, left = self._create_card(wrap, "Compose Email")
        left_card.pack(side="left", fill="y", padx=(0, 16), pady=4)

        right_card, right = self._create_card(wrap, "Preview")
        right_card.pack(side="left", fill="both", expand=True, pady=4)

        ttk.Label(left, text="Client", style="SectionTitle.TLabel").pack(anchor="w")
        self.client_var = tk.StringVar()
        self.client_combo = ttk.Combobox(left, textvariable=self.client_var, state="readonly", width=40)
        self.client_combo.pack(anchor="w", pady=(0, 12))
        self.client_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_client_change())

        ttk.Label(left, text="Site", style="SectionTitle.TLabel").pack(anchor="w")
        self.site_var = tk.StringVar()
        self.site_combo = ttk.Combobox(left, textvariable=self.site_var, state="readonly", width=40)
        self.site_combo.pack(anchor="w", pady=(0, 12))
        self.site_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_site_change())

        ttk.Label(left, text="Media Type", style="SectionTitle.TLabel").pack(anchor="w")
        self.media_var = tk.StringVar(value=MEDIA_OPTIONS[1])
        self.media_combo = ttk.Combobox(left, textvariable=self.media_var, values=MEDIA_OPTIONS, state="readonly", width=40)
        self.media_combo.pack(anchor="w", pady=(0, 12))
        self.media_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        ttk.Label(left, text="Date", style="SectionTitle.TLabel").pack(anchor="w")
        date_frame = ttk.Frame(left, style="Panel.TFrame")
        date_frame.pack(anchor="w", pady=(0, 12))
        self.date_var = tk.StringVar()
        self.date_entry = ttk.Entry(date_frame, textvariable=self.date_var, width=32, state="readonly")
        self.date_entry.pack(side="left")
        ttk.Button(date_frame, text="Pick", command=self._open_calendar).pack(side="left", padx=6)

        ttk.Label(left, text="To", style="SectionTitle.TLabel").pack(anchor="w")
        self.to_var = tk.StringVar()
        self.to_entry = ttk.Entry(left, textvariable=self.to_var, width=50)
        self.to_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(left, text="CC List", style="SectionTitle.TLabel").pack(anchor="w")
        self.cc_list_var = tk.StringVar()
        self.cc_list_combo = ttk.Combobox(left, textvariable=self.cc_list_var, state="readonly", width=40)
        self.cc_list_combo.pack(anchor="w", pady=(0, 6))
        self.cc_list_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_cc_list_change())

        ttk.Label(left, text="CC", style="SectionTitle.TLabel").pack(anchor="w")
        self.cc_var = tk.StringVar()
        self.cc_entry = ttk.Entry(left, textvariable=self.cc_var, width=50)
        self.cc_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(left, text="Body", style="SectionTitle.TLabel").pack(anchor="w")
        self.signature_var = tk.StringVar()
        self.signature_combo = ttk.Combobox(left, textvariable=self.signature_var, state="readonly", width=40)
        self.signature_combo.pack(anchor="w", pady=(0, 12))
        self.signature_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        options_row = ttk.Frame(left, style="Panel.TFrame")
        options_row.pack(anchor="w", fill="x", pady=(0, 2))

        self.template_folder_frame = ttk.Frame(options_row, style="Panel.TFrame")
        self.template_folder_frame.pack(side="left")
        self.template_folder_vars = {}
        self._refresh_template_folder_selector()

        ttk.Label(left, text="Dropbox Link", style="SectionTitle.TLabel").pack(anchor="w")
        self.link_var = tk.StringVar()
        self.link_entry = ttk.Entry(left, textvariable=self.link_var, width=50)
        self.link_entry.pack(anchor="w", pady=(0, 12))

        ttk.Button(left, text="Create Outlook Draft", style="Accent.TButton", command=self._create_draft).pack(anchor="w", pady=(6, 6))
        ttk.Button(left, text="Reload", command=self._reload_compose_tab).pack(anchor="w")

        ttk.Label(right, text="Subject", style="SectionTitle.TLabel").pack(anchor="w")
        self.subject_var = tk.StringVar()
        self.subject_entry = ttk.Entry(right, textvariable=self.subject_var, width=100)
        self.subject_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(right, text="Email Preview", style="SectionTitle.TLabel").pack(anchor="w")
        preview_frame = ttk.Frame(right, style="Panel.TFrame")
        preview_frame.pack(fill="both", expand=True)
        self.preview = tk.Text(preview_frame, wrap="word", height=28)
        preview_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.preview.yview)
        self.preview.configure(yscrollcommand=preview_scroll.set)
        self.preview.pack(side="left", fill="both", expand=True)
        preview_scroll.pack(side="right", fill="y")
        self._style_text_widget(self.preview)
        self.preview.config(state="disabled")

        for var in [
            self.client_var,
            self.site_var,
            self.media_var,
            self.date_var,
            self.link_var,
            self.to_var,
            self.cc_var,
            self.signature_var,
            self.subject_var,
        ]:
            var.trace_add("write", lambda *_args: self._update_email_preview())

    def _build_master_tab(self):
        self._build_client_sites_tab(self.master_tab, MASTER_SHEET)

    def _build_oil_gas_tab(self):
        self._build_client_sites_tab(self.oil_gas_tab, OIL_GAS_SHEET)

    def _build_client_sites_tab(self, parent, sheet_name: str):
        title = "Construction Directory" if sheet_name == MASTER_SHEET else "Oil & Gas Directory"
        card, frame = self._create_card(parent, title)
        card.pack(fill="both", expand=True, padx=16, pady=16)

        ttk.Label(frame, text="Saved Records", style="SectionTitle.TLabel").pack(anchor="w", pady=(0, 8))
        tree_frame = ttk.Frame(frame, style="Panel.TFrame")
        tree_frame.pack(fill="both", expand=True)

        tree = ttk.Treeview(
            tree_frame,
            columns=("client", "site", "email", "link"),
            show="headings",
        )
        tree.heading("client", text="Client")
        tree.heading("site", text="Site")
        tree.heading("email", text="Operations Contact Emails")
        tree.heading("link", text="Customer Dropbox URL")
        tree.column("client", width=180)
        tree.column("site", width=240)
        tree.column("email", width=320)
        tree.column("link", width=420)
        master_scroll = ttk.Scrollbar(
            tree_frame,
            orient="vertical",
            command=tree.yview,
            style="ClientSites.Vertical.TScrollbar",
        )
        tree.configure(yscrollcommand=master_scroll.set)
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        tree.grid(row=0, column=0, sticky="nsew")
        master_scroll.grid(row=0, column=1, sticky="ns")
        tree.bind("<<TreeviewSelect>>", lambda _e, sheet=sheet_name: self._load_master_selection(sheet))

        form = ttk.Frame(frame, style="Panel.TFrame")
        form.pack(fill="x", pady=10)

        ttk.Label(form, text="Client", style="SectionTitle.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(form, text="Site", style="SectionTitle.TLabel").grid(row=0, column=1, sticky="w")
        ttk.Label(form, text="Operations Contact Emails", style="SectionTitle.TLabel").grid(row=0, column=2, sticky="w")
        ttk.Label(form, text="Customer Dropbox URL", style="SectionTitle.TLabel").grid(row=0, column=3, sticky="w")

        client_var = tk.StringVar()
        site_var = tk.StringVar()
        email_var = tk.StringVar()
        link_var = tk.StringVar()

        ttk.Entry(form, textvariable=client_var, width=26).grid(row=1, column=0, sticky="w", padx=(0, 10))
        ttk.Entry(form, textvariable=site_var, width=32).grid(row=1, column=1, sticky="w", padx=(0, 10))
        email_frame = ttk.Frame(form, style="Panel.TFrame")
        email_frame.grid(row=1, column=2, sticky="nsew", padx=(0, 10))
        email_text = tk.Text(email_frame, width=40, height=1, wrap="word")
        email_scroll = ttk.Scrollbar(
            email_frame,
            orient="vertical",
            command=email_text.yview,
        )
        email_text.configure(yscrollcommand=email_scroll.set)
        email_text.pack(side="left", fill="both", expand=True)
        email_scroll.pack(side="right", fill="y")
        self._style_text_widget(email_text)
        email_scroll.pack_forget()
        email_text.bind("<FocusIn>", lambda _e, sheet=sheet_name: self._expand_master_email_editor(sheet))
        email_text.bind("<FocusOut>", lambda _e, sheet=sheet_name: self._collapse_master_email_editor(sheet))
        email_text.bind("<Control-v>", lambda e, sheet=sheet_name: self._paste_master_email_list(e, sheet))
        email_text.bind("<<Paste>>", lambda e, sheet=sheet_name: self._paste_master_email_list(e, sheet))
        ttk.Entry(form, textvariable=link_var, width=52).grid(row=1, column=3, sticky="w")
        form.grid_columnconfigure(2, weight=1)

        btns = ttk.Frame(frame, style="Panel.TFrame")
        btns.pack(fill="x", pady=4)
        ttk.Button(btns, text="Add", style="Accent.TButton", command=lambda sheet=sheet_name: self._add_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Update", command=lambda sheet=sheet_name: self._update_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Delete", command=lambda sheet=sheet_name: self._delete_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Reload", command=lambda sheet=sheet_name: self._reload_master_tab(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Undo Last Change", command=self._undo_last_change).pack(side="left")

        self.client_site_tabs[sheet_name] = {
            "tree": tree,
            "client_var": client_var,
            "site_var": site_var,
            "email_var": email_var,
            "link_var": link_var,
            "email_text": email_text,
            "email_scroll": email_scroll,
        }

        if sheet_name == MASTER_SHEET:
            self.master_tree = tree
            self.master_client_var = client_var
            self.master_site_var = site_var
            self.master_email_var = email_var
            self.master_link_var = link_var
            self.master_email_text = email_text
            self.master_email_scroll = email_scroll

        self._refresh_master_tree(sheet_name)

    def _build_cc_tab(self):
        outer = ttk.Frame(self.cc_tab)
        outer.pack(fill="both", expand=True, padx=16, pady=16)

        left_card, left = self._create_card(outer, "Internal CC Lists")
        left_card.pack(side="left", fill="y", padx=(0, 16))

        right_card, right = self._create_card(outer, "List Members")
        right_card.pack(side="left", fill="both", expand=True)

        ttk.Label(left, text="CC Lists", style="SectionTitle.TLabel").pack(anchor="w")
        cc_list_frame = ttk.Frame(left, style="Panel.TFrame")
        cc_list_frame.pack(anchor="w", pady=(4, 8), fill="y")
        self.cc_listbox = tk.Listbox(cc_list_frame, height=20, width=28)
        cc_list_scroll = ttk.Scrollbar(cc_list_frame, orient="vertical", command=self.cc_listbox.yview)
        self.cc_listbox.configure(yscrollcommand=cc_list_scroll.set)
        self.cc_listbox.pack(side="left", fill="y")
        cc_list_scroll.pack(side="right", fill="y")
        self._style_listbox_widget(self.cc_listbox)
        self.cc_listbox.bind("<<ListboxSelect>>", lambda _e: self._load_cc_list_selection())

        self.new_cc_list_var = tk.StringVar()
        ttk.Entry(left, textvariable=self.new_cc_list_var, width=28).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Create List", style="Accent.TButton", command=self._create_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Delete List", command=self._delete_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Reload", command=self._reload_cc_tab).pack(anchor="w")

        ttk.Label(right, text="Emails in Selected List", style="SectionTitle.TLabel").pack(anchor="w")
        cc_emails_frame = ttk.Frame(right, style="Panel.TFrame")
        cc_emails_frame.pack(anchor="w", pady=(4, 8), fill="both", expand=True)
        self.cc_emails_listbox = tk.Listbox(cc_emails_frame, height=20)
        cc_emails_scroll = ttk.Scrollbar(cc_emails_frame, orient="vertical", command=self.cc_emails_listbox.yview)
        self.cc_emails_listbox.configure(yscrollcommand=cc_emails_scroll.set)
        self.cc_emails_listbox.pack(side="left", fill="both", expand=True)
        cc_emails_scroll.pack(side="right", fill="y")
        self._style_listbox_widget(self.cc_emails_listbox)

        email_controls = ttk.Frame(right, style="Panel.TFrame")
        email_controls.pack(fill="x")
        self.cc_email_var = tk.StringVar()
        ttk.Entry(email_controls, textvariable=self.cc_email_var, width=50).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Add Email", style="Accent.TButton", command=self._add_cc_email).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Remove Selected", command=self._remove_cc_email).pack(side="left")

        self._refresh_cc_listbox()

    def _build_about_tab(self):
        card, frame = self._create_card(self.about_tab, "About")
        card.pack(fill="both", expand=True, padx=16, pady=16)

        ttk.Label(
            frame,
            text="Created on: February 4th, 2026\nBy: Steven McLaren, Data Technician",
            anchor="center",
            justify="center",
            style="CardTitle.TLabel",
        ).pack(expand=True, pady=24)

    def _show_tab(self, tab_name: str):
        frame = self.tabs.get(tab_name)
        if not frame:
            return
        frame.tkraise()
        for name, button in self.nav_buttons.items():
            button.configure(style="ActiveNav.TButton" if name == tab_name else "Nav.TButton")

    def _refresh_clients(self):
        clients = [""] + sorted({row["client"] for row in self._compose_dropbox_data()})
        self.client_combo["values"] = clients
        current = self.client_var.get().strip()
        if current in clients:
            self.client_var.set(current)
        else:
            self.client_var.set("")
        self._on_client_change()
        self._refresh_cc_lists()
        self._refresh_signature_list()

    def _refresh_cc_lists(self):
        list_names = sorted(self.cc_lists.keys())
        self.cc_list_combo["values"] = list_names
        default_name = "Commercial Construction"
        current = self.cc_list_var.get().strip()
        if current and current in self.cc_lists:
            self.cc_list_var.set(current)
        elif default_name in self.cc_lists:
            self.cc_list_var.set(default_name)
        elif list_names:
            self.cc_list_var.set(list_names[0])
        else:
            self.cc_list_var.set("")
        self._on_cc_list_change()

    def _on_cc_list_change(self):
        list_name = self.cc_list_var.get()
        emails = self.cc_lists.get(list_name, [])
        self.cc_var.set("; ".join(emails))
        next_sheet = self._sheet_for_cc_list(list_name)
        if next_sheet != self.compose_sheet_name:
            self.compose_sheet_name = next_sheet
            self._refresh_clients()
        else:
            self._on_client_change()

    def _refresh_signature_list(self):
        if not hasattr(self, "signature_combo"):
            return
        names = list(self.email_templates.keys())
        names = sorted(names)
        default_names = ["Main", "Main (smclaren@archaerial.com)"]
        current = self.signature_var.get().strip() if hasattr(self, "signature_var") else ""
        self.signature_combo["values"] = names
        if current and current in names:
            self.signature_var.set(current)
        elif default_names[0] in names:
            self.signature_var.set(default_names[0])
        elif default_names[1] in names:
            self.signature_var.set(default_names[1])
        elif names:
            self.signature_var.set(names[0])
        else:
            self.signature_var.set("")

    def _template_context(self):
        return {
            "client": self.client_var.get().strip(),
            "site": self.site_var.get().strip(),
            "media_subject": MEDIA_SUBJECT.get(self.media_var.get(), self.media_var.get()),
            "media_body": MEDIA_BODY.get(self.media_var.get(), self.media_var.get()),
            "date": self.date_var.get().strip(),
            "dropbox_link": self.link_var.get().strip(),
            "to_email": self.to_var.get().strip(),
            "cc_email": self.cc_var.get().strip(),
        }

    def _render_template_text(self, template_text, ctx):
        if not template_text:
            return ""

        def _replace(m):
            key = m.group(1)
            return str(ctx.get(key, m.group(0)))

        return re.sub(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", _replace, template_text)

    def _trim_leading_main_block(self, template_name: str, html: str) -> str:
        if not html:
            return html
        marker = "Our pilots captured {media_body} of the {site} site on {date}."
        if marker not in html:
            return html
        meta_idx = html.find("<meta")
        if meta_idx == -1:
            return html
        next_div = html.find("<div", meta_idx)
        if next_div == -1:
            return html
        return html[next_div:]

    def _apply_legacy_placeholders(self, content: str, ctx) -> str:
        if not content:
            return content
        text = content
        text = text.replace("February 00, 2026", ctx.get("date", ""))
        text = re.sub(r"\bCLIENT\b", ctx.get("client", ""), text)
        text = re.sub(r"\bSITE\b", ctx.get("site", ""), text)
        return text

    def _html_to_plain_text(self, html: str) -> str:
        text = re.sub(r"(?is)<(script|style).*?</\1>", "", html or "")
        text = re.sub(r"<\s*/\s*p\s*>", "\n\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<\s*br\s*/?>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        try:
            from html import unescape
            text = unescape(text)
        except Exception:
            pass
        return text.strip()

    def _refresh_template_folder_selector(self):
        if not hasattr(self, "template_folder_frame"):
            return

        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()

        for child in self.template_folder_frame.winfo_children():
            child.destroy()
        self.template_folder_vars = {}

        if not self.template_folder_names:
            return

        ttk.Label(self.template_folder_frame, text="Templates").pack(side="left", padx=(0, 8))
        for folder_name in self.template_folder_names:
            var = tk.BooleanVar(value=folder_name == self.active_template_folder_name)
            self.template_folder_vars[folder_name] = var
            ttk.Checkbutton(
                self.template_folder_frame,
                text=folder_name,
                variable=var,
                command=lambda name=folder_name: self._on_template_folder_toggle(name),
            ).pack(side="left", padx=(0, 8))

    def _on_template_folder_toggle(self, folder_name):
        selected_var = self.template_folder_vars.get(folder_name)
        if selected_var is None:
            return

        if not selected_var.get():
            selected_var.set(True)
            return

        for name, var in self.template_folder_vars.items():
            var.set(name == folder_name)

        if self.active_template_folder_name == folder_name:
            return

        self.active_template_folder_name = folder_name
        self.email_templates = self._read_email_templates()
        self._refresh_signature_list()
        self._update_email_preview()

    def _refresh_cc_listbox(self):
        if not hasattr(self, "cc_listbox"):
            return
        self.cc_listbox.delete(0, "end")
        for name in sorted(self.cc_lists.keys()):
            self.cc_listbox.insert("end", name)
        self._load_cc_list_selection()

    def _load_cc_list_selection(self):
        if not hasattr(self, "cc_listbox"):
            return
        selection = self.cc_listbox.curselection()
        if not selection:
            self.cc_emails_listbox.delete(0, "end")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        self.cc_emails_listbox.delete(0, "end")
        for email in emails:
            self.cc_emails_listbox.insert("end", email)

    def _create_cc_list(self):
        name = self.new_cc_list_var.get().strip()
        if not name:
            messagebox.showwarning("Missing Data", "List name is required.")
            return
        if name in self.cc_lists:
            messagebox.showwarning("Already Exists", "That list already exists.")
            return
        self.cc_lists[name] = []
        self._write_cc_lists()
        self.new_cc_list_var.set("")
        self._reload_cc_lists_data()

    def _delete_cc_list(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to delete.")
            return
        list_name = self.cc_listbox.get(selection[0])
        if not messagebox.askyesno("Confirm Delete", f"Delete CC list '{list_name}'?"):
            return
        self.cc_lists.pop(list_name, None)
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _add_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to add to.")
            return
        email = self.cc_email_var.get().strip()
        if not email:
            messagebox.showwarning("Missing Data", "Email is required.")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        if email in emails:
            messagebox.showwarning("Already Exists", "That email is already in the list.")
            return
        emails.append(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self.cc_email_var.set("")
        self._reload_cc_lists_data()

    def _remove_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to remove from.")
            return
        list_name = self.cc_listbox.get(selection[0])
        email_selection = self.cc_emails_listbox.curselection()
        if not email_selection:
            messagebox.showwarning("Select Email", "Choose an email to remove.")
            return
        emails = self.cc_lists.get(list_name, [])
        for idx in reversed(email_selection):
            email = self.cc_emails_listbox.get(idx)
            if email in emails:
                emails.remove(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _reload_cc_lists_data(self):
        self.cc_lists = self._read_cc_lists()
        self._refresh_cc_listbox()
        self._refresh_cc_lists()

    def _reload_compose_tab(self):
        # Reload only Compose-related data
        self.signature_names = self._read_signature_names()
        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()
        self.email_templates = self._read_email_templates()
        self.cc_lists = self._read_cc_lists()
        self._refresh_template_folder_selector()
        # Refresh data sources for client/site/email/link without touching master tab UI
        self._reload_master_data(sheet_name=MASTER_SHEET, refresh_clients=True, refresh_master_tree=False)
        self._reload_master_data(sheet_name=OIL_GAS_SHEET, refresh_clients=False, refresh_master_tree=False)
        self._refresh_cc_lists()
        self._reset_compose_state(refresh_clients=True, force_blank_client_site=True)

    def _reload_master_tab(self, sheet_name: str = MASTER_SHEET):
        # Reload only the requested client-sites tab data
        self._reload_master_data(sheet_name=sheet_name, refresh_clients=False, refresh_master_tree=True)
        self._reset_master_form(sheet_name)

    def _reload_cc_tab(self):
        # Reload only Internal CC Lists tab data
        self._reload_cc_lists_data()
        self._reset_cc_tab()

    def _on_client_change(self):
        client = self.client_var.get()
        compose_dropbox_data = self._compose_dropbox_data()
        if not client:
            self.site_combo["values"] = [""]
            self.site_var.set("")
            self.link_var.set("")
            self.to_var.set("")
            return
        sites = [""] + sorted({row["site"] for row in compose_dropbox_data if row["client"] == client})
        self.site_combo["values"] = sites
        current_site = self.site_var.get().strip()
        if current_site in sites:
            self.site_var.set(current_site)
        else:
            self.site_var.set("")
        self._on_site_change()

    def _on_site_change(self):
        client = self.client_var.get()
        site = self.site_var.get()
        compose_dropbox_data = self._compose_dropbox_data()
        compose_contacts_data = self._compose_contacts_data()
        if not client:
            self.link_var.set("")
            self.to_var.set("")
            return

        c_key = _normalize_key(client)

        def _client_rows(rows):
            exact = [r for r in rows if r["client"] == client]
            if exact:
                return sorted(exact, key=lambda r: str(r.get("site", "")))
            norm = [r for r in rows if _normalize_key(r["client"]) == c_key]
            return sorted(norm, key=lambda r: str(r.get("site", "")))

        link = ""
        if site:
            for row in compose_dropbox_data:
                if row["client"] == client and row["site"] == site:
                    link = row["link"]
                    break
            if not link:
                s_key = _normalize_key(site)
                for row in compose_dropbox_data:
                    if _normalize_key(row["client"]) == c_key and _normalize_key(row["site"]) == s_key:
                        link = row["link"]
                        break
        else:
            client_dropbox_rows = _client_rows(compose_dropbox_data)
            if client_dropbox_rows:
                link = client_dropbox_rows[0].get("link", "")
        self.link_var.set(link)

        email = ""
        if site:
            for row in compose_contacts_data:
                if row["client"] == client and row["site"] == site:
                    email = row["email"]
                    break
            if not email:
                s_key = _normalize_key(site)
                for row in compose_contacts_data:
                    if _normalize_key(row["client"]) == c_key and _normalize_key(row["site"]) == s_key:
                        email = row["email"]
                        break
        else:
            client_contact_rows = _client_rows(compose_contacts_data)
            if client_contact_rows:
                email = client_contact_rows[0].get("email", "")
        self.to_var.set(normalize_email_list(email))

    def _reset_compose_state(self, refresh_clients: bool = True, force_blank_client_site: bool = False):
        if not hasattr(self, "client_var"):
            return
        self.media_var.set(MEDIA_OPTIONS[1])
        self.subject_var.set("")
        self.to_var.set("")
        self.cc_var.set("")
        self.link_var.set("")
        self._set_default_date()
        if force_blank_client_site:
            self.client_var.set("")
            self.site_var.set("")
            self.to_var.set("")
            self.link_var.set("")
        if refresh_clients:
            self._refresh_clients()
        self._update_email_preview()

    def _reset_master_form(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        widgets["client_var"].set("")
        widgets["site_var"].set("")
        self._set_master_email_text("", sheet_name)
        widgets["link_var"].set("")
        tree = widgets["tree"]
        tree.selection_remove(tree.selection())

    def _reset_cc_tab(self):
        if hasattr(self, "new_cc_list_var"):
            self.new_cc_list_var.set("")
        if hasattr(self, "cc_email_var"):
            self.cc_email_var.set("")
        if hasattr(self, "cc_listbox"):
            self.cc_listbox.selection_clear(0, "end")
        if hasattr(self, "cc_emails_listbox"):
            self.cc_emails_listbox.delete(0, "end")

    def _set_default_date(self):
        self.selected_date = default_date()
        self.date_var.set(format_date(self.selected_date))

    def _open_calendar(self):
        CalendarPopup(self.root, self.selected_date, self._set_date)

    def _set_date(self, date_obj):
        self.selected_date = date_obj
        self.date_var.set(format_date(date_obj))
        self._update_email_preview()

    def _get_master_email_text(self, sheet_name: str = MASTER_SHEET) -> str:
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return ""
        email_text = widgets.get("email_text")
        if not email_text:
            return widgets["email_var"].get()
        return email_text.get("1.0", "end-1c")

    def _set_master_email_text(self, value: str, sheet_name: str = MASTER_SHEET):
        normalized = normalize_email_list(value)
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        widgets["email_var"].set(normalized)
        email_text = widgets.get("email_text")
        if email_text:
            email_text.delete("1.0", "end")
            if normalized:
                email_text.insert("1.0", normalized)

    def _expand_master_email_editor(self, sheet_name: str = MASTER_SHEET, _event=None):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        widgets["email_text"].configure(height=4)
        widgets["email_scroll"].pack(side="right", fill="y")

    def _collapse_master_email_editor(self, sheet_name: str = MASTER_SHEET, _event=None):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        self._set_master_email_text(self._get_master_email_text(sheet_name), sheet_name)
        widgets["email_text"].configure(height=1)
        widgets["email_scroll"].pack_forget()

    def _paste_email_list(self, event, entry, variable):
        try:
            raw = self.root.clipboard_get()
        except tk.TclError:
            return None

        normalized = normalize_email_list(raw)
        if not normalized:
            return "break"

        entry.insert("insert", normalized)
        variable.set(normalize_email_list(variable.get()))
        return "break"

    def _paste_master_email_list(self, event, sheet_name: str = MASTER_SHEET):
        try:
            raw = self.root.clipboard_get()
        except tk.TclError:
            return None

        normalized = normalize_email_list(raw)
        if not normalized:
            return "break"

        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return "break"
        widgets["email_text"].insert("insert", normalized)
        self._set_master_email_text(self._get_master_email_text(sheet_name), sheet_name)
        return "break"

    def _compose_subject(self):
        ctx = self._template_context()
        template = self.email_templates.get(self.signature_var.get().strip())
        if template and template.get("subject"):
            rendered = self._render_template_text(template["subject"], ctx)
            if rendered:
                return rendered
        if ctx["site"] and ctx["media_subject"] and ctx["date"]:
            return f"{ctx['site']} - {ctx['media_subject']} - {ctx['date']}"
        return ""

    def _compose_body(self):
        ctx = self._template_context()
        selected_name = self.signature_var.get().strip()
        template = self.email_templates.get(selected_name)
        if template:
            if template.get("html"):
                rendered = self._trim_leading_main_block(selected_name, template["html"])
                rendered = self._render_template_text(rendered, ctx)
                rendered = self._apply_legacy_placeholders(rendered, ctx)
                rendered = self._html_to_plain_text(rendered)
                if rendered:
                    return rendered.replace("\n", "\r\n")
            if template.get("text"):
                rendered = self._render_template_text(template["text"], ctx)
                if rendered:
                    return rendered.replace("\n", "\r\n")
        greeting = f"Good Day {ctx['client']} Team,"
        main = (
            f"Our pilots captured {ctx['media_body']} of the {ctx['site']} site on {ctx['date']}. "
            "Please use the link below to access your imagery."
        )
        body = (
            f"{greeting}\r\n\r\n"
            f"{main}\r\n"
            f"Dropbox Link: {ctx['dropbox_link']}\r\n\r\n"
            "Let us know if you have any questions, feedback, or if we can further assist you.\r\n\r\n"
        )
        return body

    def _compose_body_html(self):
        ctx = self._template_context()
        selected_name = self.signature_var.get().strip()
        template = self.email_templates.get(selected_name)
        if template and (template.get("html") or template.get("text")):
            def _ensure_assets(html):
                return self._extract_inline_images(html, base_dir=template.get("path").parent if template.get("path") else None)

            def _text_to_html(text):
                from html import escape
                link = ctx["dropbox_link"]
                paragraphs = []
                current = []
                for line in text.splitlines():
                    if line.strip() == "":
                        if current:
                            paragraphs.append("<br>".join(current))
                            current = []
                        continue
                    escaped = escape(line)
                    if "Dropbox Link:" in line and link:
                        escaped = escaped.replace(escape(link), f'<a href="{escape(link)}">{escape(link)}</a>')
                    current.append(escaped)
                if current:
                    paragraphs.append("<br>".join(current))
                html_parts = [f"<p style=\"margin:0 0 10px 0;\">{p}</p>" for p in paragraphs]
                return (
                    "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; line-height:1.35;\">"
                    + "".join(html_parts)
                    + "</div>"
                )

            if template.get("html"):
                rendered = self._trim_leading_main_block(selected_name, template["html"])
                rendered = self._render_template_text(rendered, ctx)
                rendered = self._apply_legacy_placeholders(rendered, ctx)
                if ctx["dropbox_link"]:
                    rendered = rendered.replace(
                        ctx["dropbox_link"],
                        f'<a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>',
                    )
                    rendered = re.sub(
                        r"(<b>\s*Dropbox Link:\s*</b>)(?!\s*<a\b)",
                        rf'\1 <a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>',
                        rendered,
                        flags=re.IGNORECASE,
                    )
                body_html, template_assets = _ensure_assets(rendered)
                return body_html, template_assets

            if template.get("text"):
                rendered = self._render_template_text(template["text"], ctx)
                body_html = _text_to_html(rendered)
                return body_html, []
        link_html = f'<a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>' if ctx["dropbox_link"] else ""

        body_html = (
            "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; "
            "line-height:1.35;\">"
            f"<p style=\"margin:0 0 10px 0;\">Good Day {ctx['client']} Team,</p>"
            f"<p style=\"margin:0 0 8px 0;\">Our pilots captured {ctx['media_body']} of the {ctx['site']} site on {ctx['date']}. "
            "Please use the link below to access your imagery.</p>"
            f"<p style=\"margin:0 0 10px 0;\"><b>Dropbox Link:</b> {link_html}</p>"
            "<p style=\"margin:0 0 12px 0;\">Let us know if you have any questions, feedback, or if we can further assist you.</p>"
            ""
            "</div>"
        )
        return body_html, []

    def _update_email_preview(self):
        subject = self._compose_subject()
        body = self._compose_body()
        if subject:
            self.subject_var.set(subject)
        self.preview.config(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("1.0", body)
        self.preview.config(state="disabled")

    def _parse_recipients(self, raw: str):
        normalized = normalize_email_list(raw)
        tokens = re.split(r"[;,]+", normalized)
        return [t.strip() for t in tokens if t.strip()]

    def _create_draft_graph(self, subject: str, body_html: str, signature_assets):
        if not self.graph_client:
            raise RuntimeError("Graph client is not configured.")

        to_list = self._parse_recipients(self.to_var.get().strip())
        cc_list = self._parse_recipients(self.cc_var.get().strip())

        if not subject or not to_list:
            raise ValueError("Please ensure a client, site, and To address are selected.")

        inline_attachments = []
        for asset in signature_assets or []:
            if not asset.get("path"):
                continue
            inline_attachments.append(
                {"path": asset["path"], "cid": asset.get("cid") or Path(asset["path"]).stem}
            )

        self.graph_client.create_draft(
            subject=subject,
            body_html=body_html,
            to=to_list,
            cc=cc_list,
            bcc=[],
            attachments=[],
            inline_attachments=inline_attachments,
        )

    def _create_draft(self):
        subject = self.subject_var.get().strip()
        try:
            body_html, signature_assets = self._compose_body_html()
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to build email HTML: {exc}")
            return

        if self.graph_client:
            try:
                self._create_draft_graph(subject, body_html, signature_assets)
                return
            except Exception as exc:
                messagebox.showerror("Graph Error", f"Failed to create Graph draft: {exc}")
                return

        if win32 is None:
            messagebox.showerror(
                "Outlook",
                "Graph is not configured and pywin32 is not available. "
                "Install pywin32 or configure Graph.",
            )
            return

        to_list = self._parse_recipients(self.to_var.get().strip())
        cc_list = self._parse_recipients(self.cc_var.get().strip())
        if not subject or not to_list:
            messagebox.showwarning("Missing Data", "Please ensure a client, site, and To address are selected.")
            return

        try:
            outlook = win32.Dispatch("Outlook.Application")
            mail = outlook.CreateItem(0)
            for addr in to_list:
                recipient = mail.Recipients.Add(addr)
                recipient.Type = 1
            for addr in cc_list:
                recipient = mail.Recipients.Add(addr)
                recipient.Type = 2
            unresolved = []
            if not mail.Recipients.ResolveAll():
                for recipient in mail.Recipients:
                    if not recipient.Resolved:
                        unresolved.append(recipient.Name)
                if unresolved:
                    messagebox.showwarning(
                        "Unresolved Recipients",
                        "These recipients could not be resolved in Outlook:\n"
                        + "\n".join(unresolved)
                        + "\n\nPlease verify the addresses before sending.",
                    )
            mail.Subject = subject
            mail.BodyFormat = 2  # olFormatHTML
            mail.HTMLBody = body_html
            for asset in signature_assets:
                try:
                    attachment = mail.Attachments.Add(asset["path"])
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3712001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3713001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x7FFE000B",
                        True,
                    )
                except Exception as exc:
                    continue
            # Re-apply HTMLBody after inline attachments to help Outlook resolve CIDs
            mail.HTMLBody = body_html
            mail.Save()  # draft in Drafts
            try:
                self._sync_outlook_now(outlook)
            except Exception:
                pass
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to create Outlook draft: {exc}")

    def _refresh_dropbox_tree(self):
        self.dropbox_tree.delete(*self.dropbox_tree.get_children())
        for row in self.dropbox_data:
            self.dropbox_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["link"]))

    def _refresh_master_tree(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        tree = widgets["tree"]
        tree.delete(*tree.get_children())
        for row in self.sheet_data.get(sheet_name, []):
            tree.insert(
                "",
                "end",
                iid=str(row["row"]),
                values=(row["client"], row["site"], row["email"], row["link"]),
            )

    def _load_dropbox_selection(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.dropbox_data if r["row"] == row_id), None)
        if not row:
            return
        self.db_client_var.set(row["client"])
        self.db_site_var.set(row["site"])
        self.db_link_var.set(row["link"])

    def _load_master_selection(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.sheet_data.get(sheet_name, []) if r["row"] == row_id), None)
        if not row:
            return
        widgets["client_var"].set(row["client"])
        widgets["site_var"].set(row["site"])
        self._set_master_email_text(row["email"], sheet_name)
        widgets["link_var"].set(row["link"])

    def _add_master_entry(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        client = widgets["client_var"].get().strip()
        site = widgets["site_var"].get().strip()
        email = normalize_email_list(self._get_master_email_text(sheet_name))
        link = widgets["link_var"].get().strip()
        if not client:
            messagebox.showwarning("Missing Data", "Client is required.")
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = self._get_or_create_sheet(wb, sheet_name)
            ws.append([client, site, link, email])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._reload_master_tab(sheet_name)
        except Exception as exc:
            messagebox.showerror("Master Data Error", f"Failed to add entry: {exc}")

    def _update_master_entry(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = widgets["client_var"].get().strip()
        site = widgets["site_var"].get().strip()
        email = normalize_email_list(self._get_master_email_text(sheet_name))
        link = widgets["link_var"].get().strip()
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            ws.cell(row=row_id, column=1, value=client)
            ws.cell(row=row_id, column=2, value=site)
            ws.cell(row=row_id, column=3, value=link)
            ws.cell(row=row_id, column=4, value=email)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._reload_master_tab(sheet_name)
        except Exception as exc:
            messagebox.showerror("Master Data Error", f"Failed to update entry: {exc}")

    def _delete_master_entry(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._reload_master_data(sheet_name=sheet_name)
        except Exception as exc:
            messagebox.showerror("Master Data Error", f"Failed to delete entry: {exc}")

    def _add_dropbox_entry(self):
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.append([client, site, link, ""])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to add entry: {exc}")

    def _update_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.cell(row=row_id, column=1, value=client)
            ws.cell(row=row_id, column=2, value=site)
            ws.cell(row=row_id, column=3, value=link)
            wb.save(MASTER_FILE)
            wb.close()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to update entry: {exc}")

    def _delete_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected Dropbox entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to delete entry: {exc}")

    def _reload_dropbox_data(self):
        self._reload_master_data(sheet_name=MASTER_SHEET)

    def _refresh_contacts_tree(self):
        self.contacts_tree.delete(*self.contacts_tree.get_children())
        for row in self.contacts_data:
            self.contacts_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["email"]))

    def _load_contact_selection(self):
        selected = self.contacts_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.contacts_data if r["row"] == row_id), None)
        if not row:
            return
        self.ct_client_var.set(row["client"])
        self.ct_site_var.set(row["site"])
        self.ct_email_var.set(normalize_email_list(row["email"]))

    def _add_contact_entry(self):
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = normalize_email_list(self.ct_email_var.get())
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.append([client, site, None, email])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to add entry: {exc}")

    def _update_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = normalize_email_list(self.ct_email_var.get())
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.cell(row=row_id, column=1, value=client)
            ws.cell(row=row_id, column=2, value=site)
            ws.cell(row=row_id, column=4, value=email)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to update entry: {exc}")

    def _delete_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected contact entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to delete entry: {exc}")

    def _reload_contacts_data(self):
        self._reload_master_data(sheet_name=MASTER_SHEET)

    def _reload_master_data(self, sheet_name: str = MASTER_SHEET, refresh_clients: bool = True, refresh_master_tree: bool = True):
        try:
            self._sort_master_sheet(sheet_name)
        except Exception:
            pass
        self.sheet_data[sheet_name] = self._read_master_data(sheet_name, allow_missing=(sheet_name == OIL_GAS_SHEET))
        if sheet_name == MASTER_SHEET:
            self._sync_primary_master_views()
        if refresh_master_tree:
            self._refresh_master_tree(sheet_name)
        if refresh_clients and sheet_name == MASTER_SHEET:
            self._refresh_clients()


if __name__ == "__main__":
    try:
        root = ttk.Window()
    except Exception:
        root = tk.Tk()
    app = EmailTemplateApp(root)
    root.mainloop()
