import tempfile
import unittest
from pathlib import Path

import openpyxl

from scripts.prepare_sharepoint_workbook import prepare
from sharepoint_excel_sync import COMPLETED_SHEET, MASTER_SHEET, PROJECT_COLUMNS, PROJECT_TABLES


def create_source_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.title = MASTER_SHEET
        ws.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email", "Procore", "Internal URLs"])
        ws.append(["Construction Client", "Construction Site", "https://dropbox/construction", "construction@example.com", "Yes", "https://internal/construction"])

        oil = wb.create_sheet("Oil & Gas")
        oil.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email", "Procore", "Internal URLs"])
        oil.append(["Oil Client", "Oil Site", "https://dropbox/oil", "oil@example.com", "No", "https://internal/oil"])

        completed = wb.create_sheet(COMPLETED_SHEET)
        completed.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email", "Procore", None])
        completed.append(["Completed Client", "Completed Site", "https://dropbox/completed", "completed@example.com", "Done", None])
        wb.save(path)
    finally:
        wb.close()


class PrepareSharePointWorkbookTests(unittest.TestCase):
    def test_prepare_creates_tables_required_columns_and_metadata_in_output_copy(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            output = Path(tmp) / "prepared.xlsx"
            create_source_workbook(source)

            prepare(source, output, "Unit Test")

            wb = openpyxl.load_workbook(output, data_only=True)
            try:
                for sheet_name, table_name in PROJECT_TABLES.items():
                    ws = wb[sheet_name]
                    self.assertIn(table_name, ws.tables)
                    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                    self.assertFalse(any(header in (None, "") for header in headers))
                    for required in PROJECT_COLUMNS:
                        self.assertIn(required, headers)
                    record_id_col = headers.index("Record ID")
                    modified_col = headers.index("Last Modified UTC")
                    modified_by_col = headers.index("Last Modified By")
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if any(row):
                            self.assertTrue(row[record_id_col])
                            self.assertTrue(row[modified_col])
                            self.assertEqual(row[modified_by_col], "Unit Test")
            finally:
                wb.close()

    def test_prepare_does_not_mutate_source_when_output_differs(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            output = Path(tmp) / "prepared.xlsx"
            create_source_workbook(source)

            prepare(source, output, "Unit Test")

            wb = openpyxl.load_workbook(source, data_only=True)
            try:
                headers = [cell.value for cell in next(wb[MASTER_SHEET].iter_rows(min_row=1, max_row=1))]
                self.assertNotIn("Record ID", headers)
                self.assertEqual(list(wb[MASTER_SHEET].tables.keys()), [])
            finally:
                wb.close()

    def test_prepare_creates_missing_completed_sheet_with_required_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            output = Path(tmp) / "prepared.xlsx"
            create_source_workbook(source)
            wb = openpyxl.load_workbook(source)
            try:
                del wb[COMPLETED_SHEET]
                wb.save(source)
            finally:
                wb.close()

            prepare(source, output, "Unit Test")

            prepared = openpyxl.load_workbook(output)
            try:
                self.assertIn(COMPLETED_SHEET, prepared.sheetnames)
                ws = prepared[COMPLETED_SHEET]
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                for required in PROJECT_COLUMNS:
                    self.assertIn(required, headers)
                self.assertIn(PROJECT_TABLES[COMPLETED_SHEET], ws.tables)
            finally:
                prepared.close()


if __name__ == "__main__":
    unittest.main()
