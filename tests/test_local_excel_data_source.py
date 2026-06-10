import tempfile
import unittest
from pathlib import Path

import openpyxl

from sharepoint_excel_sync import (
    COMPLETED_SHEET,
    LOCAL_HEADER_ALIASES,
    MASTER_SHEET,
    OIL_GAS_SHEET,
    LocalExcelProjectDataSource,
)


def create_local_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.title = MASTER_SHEET
        ws.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email"])
        ws.append(["Zulu Client", "Site Z", "https://dropbox/z", "z@example.com; extra@example.com"])
        ws.append(["Alpha Client", "Site A", "https://dropbox/a", "alpha@example.com"])
        ws.append(["", "Blank Client", "ignored", "ignored@example.com"])

        oil = wb.create_sheet(OIL_GAS_SHEET)
        oil.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email"])

        completed = wb.create_sheet(COMPLETED_SHEET)
        completed.append(["Client", "Project Name", "Dropbox URLs", "Operations Contact Email", "Procore"])
        wb.save(path)
    finally:
        wb.close()


class LocalExcelProjectDataSourceTests(unittest.TestCase):
    def test_ensure_schema_adds_missing_project_headers_without_removing_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Customer Information.xlsx"
            create_local_workbook(path)
            data_source = LocalExcelProjectDataSource(path)

            data_source.ensure_schema()

            wb = openpyxl.load_workbook(path, data_only=True)
            try:
                for sheet_name in (MASTER_SHEET, OIL_GAS_SHEET, COMPLETED_SHEET):
                    headers = [cell.value for cell in next(wb[sheet_name].iter_rows(min_row=1, max_row=1))]
                    for header in LOCAL_HEADER_ALIASES.values():
                        self.assertIn(header, headers)
                self.assertEqual(wb[MASTER_SHEET]["A2"].value, "Zulu Client")
            finally:
                wb.close()

    def test_list_projects_ignores_blank_clients_and_applies_email_normalizer(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Customer Information.xlsx"
            create_local_workbook(path)
            data_source = LocalExcelProjectDataSource(path, email_normalizer=lambda raw: f"normalized:{raw}")

            rows = data_source.list_projects(MASTER_SHEET)

            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["client"], "Zulu Client")
            self.assertEqual(rows[0]["email"], "normalized:z@example.com; extra@example.com")
            self.assertEqual(rows[0]["record_id"], "")
            self.assertEqual(rows[0]["modified_utc"], "")

    def test_add_update_and_delete_project_persist_expected_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Customer Information.xlsx"
            create_local_workbook(path)
            data_source = LocalExcelProjectDataSource(path)
            values = {
                "client": "Beta Client",
                "site": "Beta Site",
                "link": "https://dropbox/beta",
                "email": "beta@example.com",
                "procore": "Yes",
                "internal_url": "https://internal/beta",
            }

            data_source.save_project(MASTER_SHEET, values)
            rows = data_source.list_projects(MASTER_SHEET)
            beta = next(row for row in rows if row["client"] == "Beta Client")

            updated_values = dict(values)
            updated_values["site"] = "Beta Site Updated"
            data_source.save_project(MASTER_SHEET, updated_values, beta["row"])
            updated_rows = data_source.list_projects(MASTER_SHEET)
            updated = next(row for row in updated_rows if row["client"] == "Beta Client")
            self.assertEqual(updated["site"], "Beta Site Updated")
            self.assertEqual(updated["link"], "https://dropbox/beta")
            self.assertEqual(updated["procore"], "Yes")
            self.assertEqual(updated["internal_url"], "https://internal/beta")

            data_source.delete_project(MASTER_SHEET, updated["row"])
            remaining = data_source.list_projects(MASTER_SHEET)
            self.assertNotIn("Beta Client", {row["client"] for row in remaining})

    def test_move_project_to_completed_removes_source_row_and_preserves_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Customer Information.xlsx"
            create_local_workbook(path)
            data_source = LocalExcelProjectDataSource(path)
            values = {
                "client": "Pipeline Client",
                "site": "Pipeline Site",
                "link": "https://dropbox/pipeline",
                "email": "pipeline@example.com",
                "procore": "Pipeline",
                "internal_url": "https://internal/pipeline",
            }
            data_source.save_project(OIL_GAS_SHEET, values)
            source_row = next(row for row in data_source.list_projects(OIL_GAS_SHEET) if row["client"] == "Pipeline Client")

            data_source.move_project_to_completed(OIL_GAS_SHEET, source_row["row"])

            self.assertNotIn("Pipeline Client", {row["client"] for row in data_source.list_projects(OIL_GAS_SHEET)})
            completed = next(row for row in data_source.list_projects(COMPLETED_SHEET) if row["client"] == "Pipeline Client")
            self.assertEqual(completed["site"], "Pipeline Site")
            self.assertEqual(completed["link"], "https://dropbox/pipeline")
            self.assertEqual(completed["email"], "pipeline@example.com")


if __name__ == "__main__":
    unittest.main()
