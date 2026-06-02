import datetime as _dt
import calendar as _cal
import mimetypes
import uuid
import re
import sys
import logging
import webbrowser
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from urllib.parse import unquote
from pathlib import Path
import tkinter as tk
from tkinter import messagebox

try:
    import ttkbootstrap as ttk
except Exception:  # pragma: no cover
    from tkinter import ttk

import openpyxl

try:
    import win32com.client as win32
except Exception:  # pragma: no cover
    win32 = None

def _app_base_dir() -> Path:
    """Return the folder that should hold app resources at runtime."""
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        if exe_dir.parent.name.lower() == "dist":
            # Local builds run from dist\Email Builder, but live data belongs in the project Reference Data folder.
            project_dir = exe_dir.parent.parent
            if (project_dir / "Reference Data").exists():
                return project_dir
        return exe_dir
    return Path(__file__).resolve().parent


BASE_DIR = _app_base_dir()


def _resource_path(*parts: str) -> Path:
    """Prefer the consolidated Reference Data folder, with legacy fallback."""
    ref_path = BASE_DIR / "Reference Data" / Path(*parts)
    if ref_path.exists():
        return ref_path
    return BASE_DIR / Path(*parts)


DATA_DIR = _resource_path("2026_02_03")
XLSX_WORKBOOKS_DIR = _resource_path("XLSX Workbooks")
MASTER_FILE = _resource_path("XLSX Workbooks", "Customer Information.xlsx")
if not MASTER_FILE.exists():
    MASTER_FILE = _resource_path("XLSX Workbooks", "Client-Site-URL-Email.xlsx")
if not MASTER_FILE.exists():
    MASTER_FILE = DATA_DIR / "Customer Information.xlsx"
if not MASTER_FILE.exists():
    MASTER_FILE = DATA_DIR / "Client-Site-URL-Email.xlsx"
MASTER_SHEET = "Construction"
LEGACY_MASTER_SHEET = "ABCD"
OIL_GAS_SHEET = "Oil & Gas"
COMPLETED_SHEET = "Completed Projects"
CC_LISTS_FILE = _resource_path("XLSX Workbooks", "Internal CC Lists.xlsx")
if not CC_LISTS_FILE.exists():
    CC_LISTS_FILE = _resource_path("Internal CC Lists.xlsx")
LEGACY_CC_LISTS_SHEET = "Sheet1"
DEFAULT_CC_LIST_NAMES = ["Commercial Construction", OIL_GAS_SHEET]
EMAIL_TEMPLATES_DIR = _resource_path("Email Templates")
GRAPH_DIR = _resource_path("Backend Graph")
GRAPH_SETTINGS = GRAPH_DIR / "graph_app_settings.json"
GRAPH_TOKEN_CACHE = GRAPH_DIR / ".graph_token_cache.json"

logger = logging.getLogger(__name__)

if GRAPH_DIR.exists():
    sys.path.insert(0, str(GRAPH_DIR))
try:
    from graph_client import GraphEmailClient  # type: ignore
except Exception:
    GraphEmailClient = None

MEDIA_OPTIONS = [
    "photos",
    "photos and video",
    "photos, video, and 2D mapping",
    "photos, and 2D mapping",
    "ROW Documentation Video",
]

MEDIA_SUBJECT = {
    "photos": "Aerial Progress Photos",
    "photos and video": "Aerial Progress Photos and Video",
    "photos, video, and 2D mapping": "Aerial Progress Photos, Video, and 2D Mapping",
    "photos, and 2D mapping": "Aerial Progress Photos and 2D Mapping",
    "ROW Documentation Video": "ROW Documentation Video",
}

MEDIA_BODY = {
    "photos": "progress photos",
    "photos and video": "progress photos and video",
    "photos, video, and 2D mapping": "progress photos, video, and 2D mapping",
    "photos, and 2D mapping": "progress photos and 2D mapping",
    "ROW Documentation Video": "ROW documentation video",
}

MASTER_HEADERS = {
    "client": "Client",
    "project name": "Project Name",
    "dropbox urls": "Dropbox URLs",
    "operations contact email": "Operations Contact Email",
    "procore": "Procore",
    "internal urls": "Internal URLs",
}


def _ordinal(n: int) -> str:
    if 10 <= (n % 100) <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def _normalize_key(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    cleaned = "".join(ch.lower() if ch.isalnum() else " " for ch in value)
    tokens = cleaned.split()
    expanded = []
    for tok in tokens:
        if tok == "hs":
            expanded.append("high")
            expanded.append("school")
        elif tok == "ms":
            expanded.append("middle")
            expanded.append("school")
        elif tok == "es":
            expanded.append("elementary")
            expanded.append("school")
        elif tok == "jh":
            expanded.append("junior")
            expanded.append("high")
        elif tok == "jhs":
            expanded.append("junior")
            expanded.append("high")
            expanded.append("school")
        elif tok == "jr":
            expanded.append("junior")
        elif tok == "jrhigh":
            expanded.append("junior")
            expanded.append("high")
        elif tok == "jrhighschool":
            expanded.append("junior")
            expanded.append("high")
            expanded.append("school")
        else:
            expanded.append(tok)
    return " ".join(expanded)


EMAIL_ADDRESS_RE = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@(?:[A-Za-z0-9-]+\.)+[A-Za-z]{2,}")
MOUSE_WHEEL_LINES = 4


def _sheet_aliases(sheet_name: str):
    if sheet_name == MASTER_SHEET:
        return [MASTER_SHEET, LEGACY_MASTER_SHEET]
    return [sheet_name]


def _master_cell_value(row, idx: int):
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    value = row[idx]
    if isinstance(value, str):
        return value.strip()
    return value or ""


def normalize_email_list(raw) -> str:
    if raw is None:
        return ""
    if not isinstance(raw, str):
        raw = str(raw)

    matches = EMAIL_ADDRESS_RE.findall(raw)
    if matches:
        deduped = []
        seen = set()
        for email in matches:
            key = email.lower()
            if key in seen:
                continue
            seen.add(key)
            deduped.append(email)
        return ", ".join(deduped)

    cleaned = re.sub(r"[\r\n\t;]+", ",", raw)
    parts = [part.strip() for part in cleaned.split(",") if part.strip()]
    return ", ".join(parts)



def format_date(date_obj: _dt.date) -> str:
    return date_obj.strftime(f"%B {_ordinal(date_obj.day)}, %Y")


def default_date() -> _dt.date:
    today = _dt.date.today()
    if today.weekday() == 0:  # Monday
        return today - _dt.timedelta(days=3)
    return today - _dt.timedelta(days=1)


class CalendarPopup(tk.Toplevel):
    def __init__(self, master, initial_date, on_select):
        super().__init__(master)
        self.title("Select Date")
        self.resizable(False, False)
        self.on_select = on_select
        self.year = initial_date.year
        self.month = initial_date.month
        self.calendar = _cal.Calendar(firstweekday=_cal.SUNDAY)

        header = ttk.Frame(self)
        header.pack(padx=10, pady=8, fill="x")

        ttk.Button(header, text="<", width=3, command=self._prev_month).pack(side="left")
        self.month_label = ttk.Label(header, text="")
        self.month_label.pack(side="left", expand=True)
        ttk.Button(header, text=">", width=3, command=self._next_month).pack(side="right")

        self.grid_frame = ttk.Frame(self)
        self.grid_frame.pack(padx=10, pady=8)

        self._render()

    def _prev_month(self):
        if self.month == 1:
            self.month = 12
            self.year -= 1
        else:
            self.month -= 1
        self._render()

    def _next_month(self):
        if self.month == 12:
            self.month = 1
            self.year += 1
        else:
            self.month += 1
        self._render()

    def _render(self):
        for child in self.grid_frame.winfo_children():
            child.destroy()

        self.month_label.config(text=f"{_cal.month_name[self.month]} {self.year}")

        for idx, day in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
            ttk.Label(self.grid_frame, text=day, width=4, anchor="center").grid(row=0, column=idx, padx=2, pady=2)

        month_data = self.calendar.monthdayscalendar(self.year, self.month)
        for r, week in enumerate(month_data, start=1):
            for c, day in enumerate(week):
                if day == 0:
                    ttk.Label(self.grid_frame, text="", width=4).grid(row=r, column=c, padx=2, pady=2)
                    continue
                ttk.Button(
                    self.grid_frame,
                    text=str(day),
                    width=4,
                    command=lambda d=day: self._select_day(d),
                ).grid(row=r, column=c, padx=2, pady=2)

    def _select_day(self, day):
        chosen = _dt.date(self.year, self.month, day)
        self.on_select(chosen)
        self.destroy()


class EmailTemplateApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Customer Delivery Email Builder")
        self.root.geometry("1480x900")
        self.root.minsize(1180, 700)
        self.root.resizable(True, True)

        self.master_data = []
        self.dropbox_data = []
        self.contacts_data = []
        self.sheet_data = {}
        self.client_site_tabs = {}
        self.compose_sheet_name = MASTER_SHEET
        self.cc_lists = {}
        self.signature_names = []
        self.template_folder_names = []
        self.active_template_folder_name = ""
        self.email_templates = {}
        self._undo_stack = []
        self.graph_client = None

        self._apply_theme()
        self._ensure_master_workbook_schema()
        self._load_data()
        self._init_graph_client()
        self._build_ui()
        self._refresh_clients()
        self._set_default_date()
        self._update_email_preview()

    def _load_data(self):
        self.sheet_data = {
            MASTER_SHEET: self._read_master_data(MASTER_SHEET),
            OIL_GAS_SHEET: self._read_master_data(OIL_GAS_SHEET, allow_missing=True),
            COMPLETED_SHEET: self._read_master_data(COMPLETED_SHEET, allow_missing=True),
        }
        self._sync_primary_master_views()
        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()
        self.cc_lists = self._read_cc_lists()
        self.signature_names = self._read_signature_names()
        self.email_templates = self._read_email_templates()

    def _init_graph_client(self):
        if GraphEmailClient is None:
            return
        if not GRAPH_SETTINGS.exists():
            return
        try:
            self.graph_client = GraphEmailClient(GRAPH_SETTINGS, GRAPH_TOKEN_CACHE)
        except Exception as exc:
            logger.warning("Graph client not initialized: %s", exc)
            self.graph_client = None

    def _read_template_folder_names(self):
        if not EMAIL_TEMPLATES_DIR.exists():
            return []
        return sorted(path.name for path in EMAIL_TEMPLATES_DIR.iterdir() if path.is_dir())

    def _choose_active_template_folder(self):
        names = list(self.template_folder_names)
        if self.active_template_folder_name in names:
            return self.active_template_folder_name
        return names[0] if names else ""

    def _selected_template_dir(self):
        folder_name = (self.active_template_folder_name or "").strip()
        if folder_name:
            path = EMAIL_TEMPLATES_DIR / folder_name
            if path.exists() and path.is_dir():
                return path
        return EMAIL_TEMPLATES_DIR

    def _read_email_templates(self):
        templates = {}
        template_dir = self._selected_template_dir()
        if not template_dir.exists():
            return templates
        for path in template_dir.glob("*.eml"):
            name = path.stem
            try:
                msg = BytesParser(policy=policy.default).parse(path.open("rb"))
            except Exception:
                continue
            subject = msg.get("Subject", "")
            text = ""
            html = ""
            if msg.is_multipart():
                for part in msg.walk():
                    ctype = part.get_content_type()
                    if ctype == "text/plain" and not text:
                        text = part.get_content()
                    elif ctype == "text/html" and not html:
                        html = part.get_content()
            else:
                ctype = msg.get_content_type()
                if ctype == "text/plain":
                    text = msg.get_content()
                elif ctype == "text/html":
                    html = msg.get_content()
            templates[name] = {
                "subject": subject or "",
                "text": text or "",
                "html": html or "",
                "path": path,
            }
        return templates

    def _resolve_sheet_name(self, wb, sheet_name: str, allow_missing: bool = False):
        for candidate in _sheet_aliases(sheet_name):
            if candidate in wb.sheetnames:
                return candidate
        if allow_missing:
            return None
        if sheet_name == MASTER_SHEET:
            expected = " or ".join(f'"{name}"' for name in _sheet_aliases(sheet_name))
            raise KeyError(f"Workbook is missing the Construction sheet. Expected {expected}.")
        raise KeyError(f'Workbook is missing the "{sheet_name}" sheet.')

    def _get_or_create_sheet(self, wb, sheet_name: str):
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=True)
        if resolved_sheet:
            ws = wb[resolved_sheet]
            if not any(cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))):
                headers = list(MASTER_HEADERS.values())
                source_sheet_name = self._resolve_sheet_name(wb, MASTER_SHEET, allow_missing=True)
                if source_sheet_name:
                    source_headers = [cell.value for cell in next(wb[source_sheet_name].iter_rows(min_row=1, max_row=1))]
                    if any(source_headers):
                        headers = source_headers
                for col_idx, header in enumerate(headers, start=1):
                    ws.cell(row=1, column=col_idx, value=header)
            self._master_column_map(ws, ensure=True)
            return ws

        ws = wb.create_sheet(sheet_name)
        source_sheet_name = self._resolve_sheet_name(wb, MASTER_SHEET, allow_missing=True)
        headers = list(MASTER_HEADERS.values())
        if source_sheet_name:
            source_headers = [cell.value for cell in next(wb[source_sheet_name].iter_rows(min_row=1, max_row=1))]
            if any(source_headers):
                headers = source_headers
        ws.append(headers)
        return ws

    def _master_column_map(self, ws, ensure: bool = False):
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {}
        for idx, header in enumerate(headers, start=1):
            if not header:
                continue
            normalized = str(header).strip().lower()
            if normalized in MASTER_HEADERS and normalized not in col_map:
                col_map[normalized] = idx

        changed = False
        if ensure:
            for normalized, header in MASTER_HEADERS.items():
                if normalized not in col_map:
                    idx = ws.max_column + 1
                    ws.cell(row=1, column=idx, value=header)
                    col_map[normalized] = idx
                    changed = True
        return col_map, changed

    def _build_master_row(self, ws, values: dict):
        col_map, _ = self._master_column_map(ws, ensure=True)
        row = [""] * max(col_map.values())
        for key, value in values.items():
            col_idx = col_map.get(key)
            if col_idx is None:
                continue
            row[col_idx - 1] = value
        return row

    def _ensure_master_workbook_schema(self):
        if not MASTER_FILE.exists():
            return
        wb = openpyxl.load_workbook(MASTER_FILE)
        changed = False
        for sheet_name in (MASTER_SHEET, OIL_GAS_SHEET, COMPLETED_SHEET):
            resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=True)
            if not resolved_sheet:
                continue
            headers_before = [cell.value for cell in next(wb[resolved_sheet].iter_rows(min_row=1, max_row=1))]
            ws = self._get_or_create_sheet(wb, sheet_name)
            _, sheet_changed = self._master_column_map(ws, ensure=True)
            headers_after = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            changed = changed or sheet_changed or (not any(headers_before) and any(headers_after))
        if changed:
            wb.save(MASTER_FILE)
        wb.close()

    def _sync_primary_master_views(self):
        self.master_data = list(self.sheet_data.get(MASTER_SHEET, []))
        self.dropbox_data = [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self.master_data
        ]
        self.contacts_data = [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self.master_data
        ]

    def _compose_rows(self):
        return self.sheet_data.get(self.compose_sheet_name, [])

    def _compose_dropbox_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self._compose_rows()
        ]

    def _compose_contacts_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self._compose_rows()
        ]

    def _lookup_compose_field(self, field_name: str) -> str:
        client = self.client_var.get().strip()
        site = self.site_var.get().strip()
        rows = self._compose_rows()
        if not client:
            return ""

        client_key = _normalize_key(client)

        def _client_matches(row):
            return row["client"] == client or _normalize_key(row["client"]) == client_key

        matching_rows = [row for row in rows if _client_matches(row)]
        if not matching_rows:
            return ""

        if site:
            site_key = _normalize_key(site)
            for row in matching_rows:
                if row["client"] == client and row["site"] == site:
                    return str(row.get(field_name, "") or "")
            for row in matching_rows:
                if _normalize_key(row["site"]) == site_key:
                    return str(row.get(field_name, "") or "")

        matching_rows = sorted(matching_rows, key=lambda row: str(row.get("site", "")))
        return str(matching_rows[0].get(field_name, "") or "")

    def _sheet_for_cc_list(self, list_name: str) -> str:
        normalized = (list_name or "").strip().lower()
        if normalized == OIL_GAS_SHEET.lower():
            return OIL_GAS_SHEET
        return MASTER_SHEET

    def _cc_sheet_name(self, template_name: str = "") -> str:
        name = (template_name or self.active_template_folder_name or "").strip()
        return name or LEGACY_CC_LISTS_SHEET

    def _default_cc_lists(self):
        return {name: [] for name in DEFAULT_CC_LIST_NAMES}

    def _ensure_cc_sheet(self, wb, sheet_name: str):
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1), [])]
        existing = {str(header).strip() for header in headers if header}
        changed = False
        if not existing:
            for col_idx, header in enumerate(DEFAULT_CC_LIST_NAMES, start=1):
                ws.cell(row=1, column=col_idx, value=header)
            return ws, True
        for header in DEFAULT_CC_LIST_NAMES:
            if header not in existing:
                ws.cell(row=1, column=ws.max_column + 1, value=header)
                changed = True
        return ws, changed

    def _load_cc_workbook(self):
        if CC_LISTS_FILE.exists():
            wb = openpyxl.load_workbook(CC_LISTS_FILE)
        else:
            wb = openpyxl.Workbook()
        changed = False

        if wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1 and wb["Sheet"]["A1"].value is None:
            if LEGACY_CC_LISTS_SHEET == "Sheet":
                changed = True
            else:
                wb["Sheet"].title = LEGACY_CC_LISTS_SHEET
                changed = True

        for template_name in self.template_folder_names:
            _, sheet_changed = self._ensure_cc_sheet(wb, self._cc_sheet_name(template_name))
            changed = changed or sheet_changed

        active_sheet_name = self._cc_sheet_name()
        _, sheet_changed = self._ensure_cc_sheet(wb, active_sheet_name)
        changed = changed or sheet_changed

        if changed:
            wb.save(CC_LISTS_FILE)
        return wb

    def _read_master_data(self, sheet_name: str = MASTER_SHEET, allow_missing: bool = False):
        wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=allow_missing)
        if resolved_sheet is None:
            wb.close()
            return []
        ws = wb[resolved_sheet]
        col_map, _ = self._master_column_map(ws, ensure=False)
        client_idx = col_map.get("client", 1) - 1
        site_idx = col_map.get("project name", 2) - 1
        link_idx = col_map.get("dropbox urls", 3) - 1
        email_idx = col_map.get("operations contact email", 4) - 1
        procore_idx = col_map.get("procore")
        if procore_idx is not None:
            procore_idx -= 1
        internal_url_idx = col_map.get("internal urls")
        if internal_url_idx is not None:
            internal_url_idx -= 1
        rows = []
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            client = _master_cell_value(row, client_idx)
            site = _master_cell_value(row, site_idx)
            link = _master_cell_value(row, link_idx)
            email = _master_cell_value(row, email_idx)
            procore = _master_cell_value(row, procore_idx)
            internal_url = _master_cell_value(row, internal_url_idx)
            if client:
                rows.append({
                    "row": idx,
                    "client": client,
                    "site": site,
                    "link": link,
                    "email": normalize_email_list(email),
                    "procore": procore,
                    "internal_url": internal_url,
                })
        wb.close()
        return rows

    def _sort_master_sheet(self, sheet_name: str = MASTER_SHEET):
        wb = openpyxl.load_workbook(MASTER_FILE)
        resolved_sheet = self._resolve_sheet_name(wb, sheet_name, allow_missing=(sheet_name in (OIL_GAS_SHEET, COMPLETED_SHEET)))
        if resolved_sheet is None:
            wb.close()
            return
        ws = wb[resolved_sheet]
        col_map, _ = self._master_column_map(ws, ensure=True)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # Preserve rows but sort by Client then Project Name
        def _key(row):
            client = _master_cell_value(row, col_map.get("client", 1) - 1)
            site = _master_cell_value(row, col_map.get("project name", 2) - 1)
            return (str(client).strip().lower(), str(site).strip().lower())
        rows_sorted = sorted(rows, key=_key)
        # Clear existing data rows
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        # Write sorted rows back
        for r in rows_sorted:
            ws.append(list(r))
        wb.save(MASTER_FILE)
        wb.close()

    def _read_dropbox_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "link": row["link"]}
            for row in self._read_master_data()
        ]

    def _read_contacts_data(self):
        return [
            {"row": row["row"], "client": row["client"], "site": row["site"], "email": row["email"]}
            for row in self._read_master_data()
        ]

    def _read_cc_lists(self):
        wb = self._load_cc_workbook()
        sheet_name = self._cc_sheet_name()
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif LEGACY_CC_LISTS_SHEET in wb.sheetnames:
            ws = wb[LEGACY_CC_LISTS_SHEET]
        else:
            ws, _ = self._ensure_cc_sheet(wb, sheet_name)
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        lists = {}
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            emails = []
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
                email = row[0]
                if isinstance(email, str):
                    email = email.strip()
                if email:
                    emails.append(email)
            lists[str(header).strip()] = emails
        wb.close()
        for default_name in DEFAULT_CC_LIST_NAMES:
            lists.setdefault(default_name, [])
        return lists

    def _read_signature_names(self):
        return []

    def _extract_inline_images(self, html, base_dir=None):
        assets = []
        # Handle relative file references from Outlook signature folders
        def _replace_file_src(match):
            quote = match.group(1)
            src = match.group(2)
            lower = src.lower()
            if lower.startswith("cid:") or lower.startswith("data:") or lower.startswith("http"):
                return match.group(0)
            if not base_dir:
                return match.group(0)
            rel = unquote(src)
            path = (Path(base_dir) / rel).resolve()
            if not path.exists():
                return match.group(0)
            cid = f"sig-{uuid.uuid4().hex}"
            assets.append({"cid": cid, "path": str(path)})
            return f'src={quote}cid:{cid}{quote}'

        # img src="..."; also handles v:imagedata src="..."
        html = re.sub(
            r'src=(["\'])([^"\']+)\1',
            _replace_file_src,
            html,
            flags=re.IGNORECASE,
        )
        return html, assets

    def _write_cc_lists(self):
        self._record_undo(CC_LISTS_FILE)
        wb = self._load_cc_workbook()
        sheet_name = self._cc_sheet_name()
        if sheet_name in wb.sheetnames:
            ws_old = wb[sheet_name]
            wb.remove(ws_old)
        ws = wb.create_sheet(sheet_name)
        headers = list(self.cc_lists.keys())
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)
            for row_idx, email in enumerate(self.cc_lists[header], start=2):
                ws.cell(row=row_idx, column=col_idx, value=email)
        wb.save(CC_LISTS_FILE)
        wb.close()

    def _record_undo(self, path):
        try:
            p = Path(path)
            if not p.exists():
                return
            self._undo_stack.append((p, p.read_bytes()))
            if len(self._undo_stack) > 1:
                self._undo_stack = self._undo_stack[-1:]
        except Exception:
            pass

    def _undo_last_change(self):
        if not self._undo_stack:
            if hasattr(self, "status_var"):
                self.status_var.set("Nothing to undo.")
            else:
                messagebox.showinfo("Undo", "Nothing to undo.")
            return
        path, data = self._undo_stack.pop()
        try:
            path.write_bytes(data)
            if hasattr(self, "status_var"):
                self.status_var.set(f"Undid last change: {path.name}")
            self._reload_master_data(sheet_name=MASTER_SHEET)
            self._reload_master_data(sheet_name=OIL_GAS_SHEET, refresh_clients=False)
            self._reload_master_data(sheet_name=COMPLETED_SHEET, refresh_clients=False)
            self._reload_cc_lists_data()
        except Exception as exc:
            messagebox.showerror("Undo Error", f"Failed to undo last change: {exc}")

    def _build_ui(self):
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=1)

        container = ttk.Frame(self.root)
        container.grid(row=0, column=0, sticky="nsew")
        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(1, weight=1)

        sidebar = ttk.Frame(container, style="Sidebar.TFrame", width=220)
        sidebar.grid(row=0, column=0, sticky="ns")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(1, weight=1)

        content = ttk.Frame(container)
        content.grid(row=0, column=1, sticky="nsew")
        content.grid_rowconfigure(1, weight=1)
        content.grid_columnconfigure(0, weight=1)

        header = tk.Frame(content, bg=self.colors["brand"], highlightbackground=self.colors["border"], highlightthickness=1)
        header.grid(row=0, column=0, sticky="ew")
        tk.Label(
            header,
            text="Arch Aerial, LLC || Customer Delivery Email Builder",
            bg=self.colors["brand"],
            fg=self.colors["text"],
            font=("Segoe UI Semibold", 16),
            padx=18,
            pady=14,
        ).pack(side="left")
        tab_container = ttk.Frame(content)
        tab_container.grid(row=1, column=0, sticky="nsew")

        top_nav = ttk.Frame(sidebar, style="Sidebar.TFrame")
        top_nav.grid(row=0, column=0, sticky="ew", padx=12, pady=12)

        bottom_nav = ttk.Frame(sidebar, style="Sidebar.TFrame")
        bottom_nav.grid(row=2, column=0, sticky="ew", padx=12, pady=12)

        self.compose_tab = ttk.Frame(tab_container)
        self.master_tab = ttk.Frame(tab_container)
        self.oil_gas_tab = ttk.Frame(tab_container)
        self.cc_tab = ttk.Frame(tab_container)
        self.completed_tab = ttk.Frame(tab_container)
        self.about_tab = ttk.Frame(tab_container)

        self.tabs = {
            "compose": self.compose_tab,
            "construction": self.master_tab,
            "oil_gas": self.oil_gas_tab,
            "cc": self.cc_tab,
            "completed": self.completed_tab,
            "about": self.about_tab,
        }

        self.nav_buttons = {}
        nav_specs = [
            ("compose", "Compose", top_nav),
            ("construction", "Construction", top_nav),
            ("oil_gas", "Oil & Gas", top_nav),
            ("cc", "Internal CC Lists", top_nav),
            ("completed", "Completed", top_nav),
            ("about", "About", bottom_nav),
        ]
        for tab_name, label, parent in nav_specs:
            button = ttk.Button(
                parent,
                text=label,
                style="Nav.TButton",
                command=lambda name=tab_name: self._show_tab(name),
            )
            button.pack(fill="x", pady=4)
            self.nav_buttons[tab_name] = button

        self._build_compose_tab()
        self._build_master_tab()
        self._build_oil_gas_tab()
        self._build_cc_tab()
        self._build_completed_tab()
        self._build_about_tab()

        self.status_var = tk.StringVar(value="")
        status = ttk.Label(self.root, textvariable=self.status_var, style="Muted.TLabel", anchor="w")
        status.grid(row=1, column=0, sticky="ew")

        self._show_tab("compose")

    def _apply_theme(self):
        style = ttk.Style()
        try:
            style.theme_use("darkly")
        except Exception:
            try:
                style.theme_use("clam")
            except tk.TclError:
                pass

        # Match the pyweb-local (Docker Desktop-inspired) palette
        colors = {
            "bg": "#1a1a1a",
            "panel": "#2a2a2a",
            "sidebar": "#1e1e1e",
            "brand": "#012458",
            "border": "#404040",
            "text": "#ffffff",
            "text_secondary": "#b3b3b3",
            "text_muted": "#858585",
            "primary": "#314dec",
            "info": "#3a7ca5",
            "success": "#10b981",
            "warning": "#f59e0b",
            "error": "#ef4444",
        }
        self.colors = colors

        self.root.configure(bg=colors["bg"])

        style.configure("TFrame", background=colors["bg"])
        style.configure("Card.TFrame", background=colors["panel"])
        style.configure("Sidebar.TFrame", background=colors["sidebar"])
        style.configure("Panel.TFrame", background=colors["panel"])
        style.configure("TLabel", background=colors["bg"], foreground=colors["text"], font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=colors["bg"], foreground=colors["text_secondary"])
        style.configure("Panel.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI", 10))
        style.configure("SectionTitle.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI Semibold", 10))
        style.configure("CardTitle.TLabel", background=colors["panel"], foreground=colors["text"], font=("Segoe UI Semibold", 11))

        style.configure("TNotebook", background=colors["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", padding=(14, 10), background=colors["bg"], foreground=colors["text_secondary"])
        style.map("TNotebook.Tab", background=[("selected", colors["panel"])], foreground=[("selected", colors["text"])])

        style.configure("TEntry", fieldbackground="#3a3a3a", background="#3a3a3a", foreground=colors["text"], bordercolor=colors["border"])
        style.configure("TCombobox", fieldbackground="#3a3a3a", background="#3a3a3a", foreground=colors["text"], bordercolor=colors["border"])

        style.configure("TButton", background=colors["panel"], foreground=colors["text"], padding=(10, 6))
        style.map(
            "TButton",
            background=[("active", colors["primary"])],
            foreground=[("active", colors["text"])],
        )
        style.configure("Accent.TButton", background=colors["primary"], foreground=colors["text"], padding=(12, 8))
        style.map("Accent.TButton", background=[("active", colors["brand"])], foreground=[("active", colors["text"])])
        style.configure(
            "Nav.TButton",
            background=colors["sidebar"],
            foreground=colors["text_secondary"],
            padding=(12, 10),
            anchor="w",
            relief="flat",
        )
        style.map(
            "Nav.TButton",
            background=[("active", colors["panel"])],
            foreground=[("active", colors["text"])],
        )
        style.configure(
            "ActiveNav.TButton",
            background=colors["primary"],
            foreground=colors["text"],
            padding=(12, 10),
            anchor="w",
            relief="flat",
        )
        try:
            style.configure(
                "TScrollbar",
                background=colors["panel"],
                troughcolor=colors["sidebar"],
                bordercolor=colors["border"],
                arrowcolor=colors["text_secondary"],
                relief="flat",
            )
            style.map(
                "TScrollbar",
                background=[("active", colors["primary"])],
                arrowcolor=[("active", colors["text"])],
            )
        except tk.TclError:
            pass

        style.configure(
            "Treeview",
            background=colors["panel"],
            fieldbackground=colors["panel"],
            foreground=colors["text"],
            bordercolor=colors["border"],
            borderwidth=1,
        )
        style.configure(
            "Treeview.Heading",
            background=colors["sidebar"],
            foreground=colors["text_secondary"],
            font=("Segoe UI", 9, "bold"),
        )
        style.map("Treeview", background=[("selected", "#2f3d5a")], foreground=[("selected", colors["text"])])
        try:
            style.configure("Vertical.TScrollbar", arrowsize=14)
        except tk.TclError:
            pass
        try:
            style.configure(
                "ClientSites.Vertical.TScrollbar",
                background="#555555",
                troughcolor="#444444",
                bordercolor="#444444",
                arrowcolor=colors["text_secondary"],
                relief="flat",
            )
            style.map(
                "ClientSites.Vertical.TScrollbar",
                background=[("active", "#555555")],
                arrowcolor=[("active", colors["text"])],
            )
        except tk.TclError:
            pass

        self._theme_colors = colors

    def _create_card(self, parent, title: str, subtitle: str | None = None):
        card = tk.Frame(
            parent,
            bg=self.colors["panel"],
            highlightbackground=self.colors["border"],
            highlightthickness=1,
            bd=0,
        )
        header = tk.Frame(card, bg=self.colors["brand"], bd=0)
        header.pack(fill="x")
        tk.Label(
            header,
            text=title,
            bg=self.colors["brand"],
            fg=self.colors["text"],
            font=("Segoe UI Semibold", 14),
            padx=16,
            pady=12,
        ).pack(side="left")
        if subtitle:
            tk.Label(
                header,
                text=subtitle,
                bg=self.colors["brand"],
                fg="#d5dcf7",
                font=("Segoe UI", 10),
                padx=16,
                pady=12,
            ).pack(side="right")
        body = ttk.Frame(card, style="Panel.TFrame")
        body.pack(fill="both", expand=True, padx=16, pady=16)
        return card, body

    def _style_text_widget(self, widget):
        widget.configure(
            bg="#303030",
            fg=self.colors["text"],
            insertbackground=self.colors["text"],
            selectbackground=self.colors["primary"],
            selectforeground=self.colors["text"],
            highlightbackground=self.colors["border"],
            highlightcolor=self.colors["primary"],
            relief="flat",
            borderwidth=1,
        )
        self._bind_mousewheel_lines(widget)

    def _style_listbox_widget(self, widget):
        widget.configure(
            bg="#303030",
            fg=self.colors["text"],
            selectbackground=self.colors["primary"],
            selectforeground=self.colors["text"],
            highlightbackground=self.colors["border"],
            highlightcolor=self.colors["primary"],
            relief="flat",
            borderwidth=1,
        )
        self._bind_mousewheel_lines(widget)

    def _bind_mousewheel_lines(self, widget):
        widget.bind("<MouseWheel>", self._scroll_mousewheel_lines, add="+")
        widget.bind("<Button-4>", self._scroll_mousewheel_lines, add="+")
        widget.bind("<Button-5>", self._scroll_mousewheel_lines, add="+")

    def _scroll_mousewheel_lines(self, event):
        if getattr(event, "num", None) == 4:
            direction = -1
            steps = 1
        elif getattr(event, "num", None) == 5:
            direction = 1
            steps = 1
        else:
            delta = getattr(event, "delta", 0)
            if not delta:
                return None
            direction = -1 if delta > 0 else 1
            steps = max(1, abs(delta) // 120)
        event.widget.yview_scroll(direction * MOUSE_WHEEL_LINES * steps, "units")
        return "break"


    def _build_compose_tab(self):
        wrap = ttk.Frame(self.compose_tab)
        wrap.pack(fill="both", expand=True, padx=16, pady=16)

        left_card, left = self._create_card(wrap, "Compose Email")
        left_card.pack(side="left", fill="y", padx=(0, 16), pady=4)

        right_card, right = self._create_card(wrap, "Preview")
        right_card.pack(side="left", fill="both", expand=True, pady=4)

        ttk.Label(left, text="Client", style="SectionTitle.TLabel").pack(anchor="w")
        self.client_var = tk.StringVar()
        self.client_combo = ttk.Combobox(left, textvariable=self.client_var, state="readonly", width=40)
        self.client_combo.pack(anchor="w", pady=(0, 12))
        self.client_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_client_change())

        ttk.Label(left, text="Site", style="SectionTitle.TLabel").pack(anchor="w")
        self.site_var = tk.StringVar()
        self.site_combo = ttk.Combobox(left, textvariable=self.site_var, state="readonly", width=40)
        self.site_combo.pack(anchor="w", pady=(0, 12))
        self.site_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_site_change())

        ttk.Label(left, text="Media Type", style="SectionTitle.TLabel").pack(anchor="w")
        self.media_var = tk.StringVar(value=MEDIA_OPTIONS[1])
        self.media_combo = ttk.Combobox(left, textvariable=self.media_var, values=MEDIA_OPTIONS, state="readonly", width=40)
        self.media_combo.pack(anchor="w", pady=(0, 12))
        self.media_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        ttk.Label(left, text="Date", style="SectionTitle.TLabel").pack(anchor="w")
        date_frame = ttk.Frame(left, style="Panel.TFrame")
        date_frame.pack(anchor="w", pady=(0, 12))
        self.date_var = tk.StringVar()
        self.date_entry = ttk.Entry(date_frame, textvariable=self.date_var, width=32, state="readonly")
        self.date_entry.pack(side="left")
        ttk.Button(date_frame, text="Pick", command=self._open_calendar).pack(side="left", padx=6)

        ttk.Label(left, text="To", style="SectionTitle.TLabel").pack(anchor="w")
        self.to_var = tk.StringVar()
        to_frame = ttk.Frame(left, style="Panel.TFrame")
        to_frame.pack(anchor="w", fill="x", pady=(0, 12))
        self.to_text = tk.Text(to_frame, width=50, height=1, wrap="word")
        self.to_scroll = ttk.Scrollbar(
            to_frame,
            orient="vertical",
            command=self.to_text.yview,
        )
        self.to_text.configure(yscrollcommand=self.to_scroll.set)
        self.to_text.pack(side="left", fill="both", expand=True)
        self.to_scroll.pack(side="right", fill="y")
        self._style_text_widget(self.to_text)
        self.to_scroll.pack_forget()
        self.to_text.bind("<FocusIn>", self._expand_compose_to_editor)
        self.to_text.bind("<FocusOut>", self._collapse_compose_to_editor)
        self.to_text.bind("<KeyRelease>", self._sync_compose_to_var)
        self.to_text.bind("<Control-v>", self._paste_compose_to_list)
        self.to_text.bind("<<Paste>>", self._paste_compose_to_list)

        ttk.Label(left, text="CC List", style="SectionTitle.TLabel").pack(anchor="w")
        self.cc_list_var = tk.StringVar()
        self.cc_list_combo = ttk.Combobox(left, textvariable=self.cc_list_var, state="readonly", width=40)
        self.cc_list_combo.pack(anchor="w", pady=(0, 6))
        self.cc_list_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_cc_list_change())

        ttk.Label(left, text="CC", style="SectionTitle.TLabel").pack(anchor="w")
        self.cc_var = tk.StringVar()
        cc_frame = ttk.Frame(left, style="Panel.TFrame")
        cc_frame.pack(anchor="w", fill="x", pady=(0, 12))
        self.cc_text = tk.Text(cc_frame, width=50, height=1, wrap="word")
        self.cc_scroll = ttk.Scrollbar(
            cc_frame,
            orient="vertical",
            command=self.cc_text.yview,
        )
        self.cc_text.configure(yscrollcommand=self.cc_scroll.set)
        self.cc_text.pack(side="left", fill="both", expand=True)
        self.cc_scroll.pack(side="right", fill="y")
        self._style_text_widget(self.cc_text)
        self.cc_scroll.pack_forget()
        self.cc_text.bind("<FocusIn>", self._expand_compose_cc_editor)
        self.cc_text.bind("<FocusOut>", self._collapse_compose_cc_editor)
        self.cc_text.bind("<KeyRelease>", self._sync_compose_cc_var)
        self.cc_text.bind("<Control-v>", self._paste_compose_cc_list)
        self.cc_text.bind("<<Paste>>", self._paste_compose_cc_list)

        ttk.Label(left, text="Body", style="SectionTitle.TLabel").pack(anchor="w")
        self.signature_var = tk.StringVar()
        self.signature_combo = ttk.Combobox(left, textvariable=self.signature_var, state="readonly", width=40)
        self.signature_combo.pack(anchor="w", pady=(0, 12))
        self.signature_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_email_preview())

        options_row = ttk.Frame(left, style="Panel.TFrame")
        options_row.pack(anchor="w", fill="x", pady=(0, 2))

        self.template_folder_frame = ttk.Frame(options_row, style="Panel.TFrame")
        self.template_folder_frame.pack(side="left")
        self.template_folder_vars = {}
        self._refresh_template_folder_selector()

        ttk.Label(left, text="Dropbox Link", style="SectionTitle.TLabel").pack(anchor="w")
        self.link_var = tk.StringVar()
        self.link_entry = ttk.Entry(left, textvariable=self.link_var, width=50)
        self.link_entry.pack(anchor="w", pady=(0, 12))

        ttk.Button(left, text="Create Outlook Draft", style="Accent.TButton", command=self._create_draft).pack(anchor="w", pady=(6, 6))
        ttk.Button(left, text="Reload", command=self._reload_compose_tab).pack(anchor="w")

        ttk.Label(right, text="Subject", style="SectionTitle.TLabel").pack(anchor="w")
        self.subject_var = tk.StringVar()
        self.subject_entry = ttk.Entry(right, textvariable=self.subject_var, width=100)
        self.subject_entry.pack(anchor="w", pady=(0, 12))

        ttk.Label(right, text="Email Preview", style="SectionTitle.TLabel").pack(anchor="w")
        preview_frame = ttk.Frame(right, style="Panel.TFrame")
        preview_frame.pack(fill="both", expand=True)
        self.preview = tk.Text(preview_frame, wrap="word", height=28)
        preview_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.preview.yview)
        self.preview.configure(yscrollcommand=preview_scroll.set)
        self.preview.pack(side="left", fill="both", expand=True)
        preview_scroll.pack(side="right", fill="y")
        self._style_text_widget(self.preview)
        self.preview.config(state="disabled")

        for var in [
            self.client_var,
            self.site_var,
            self.media_var,
            self.date_var,
            self.link_var,
            self.to_var,
            self.cc_var,
            self.signature_var,
            self.subject_var,
        ]:
            var.trace_add("write", lambda *_args: self._update_email_preview())

    def _build_master_tab(self):
        self._build_client_sites_tab(self.master_tab, MASTER_SHEET)

    def _build_oil_gas_tab(self):
        self._build_client_sites_tab(self.oil_gas_tab, OIL_GAS_SHEET)

    def _build_completed_tab(self):
        self._build_client_sites_tab(self.completed_tab, COMPLETED_SHEET)

    def _build_client_sites_tab(self, parent, sheet_name: str):
        titles = {
            MASTER_SHEET: "Construction Directory",
            OIL_GAS_SHEET: "Oil & Gas Directory",
            COMPLETED_SHEET: "Completed Projects",
        }
        title = titles.get(sheet_name, f"{sheet_name} Directory")
        card, frame = self._create_card(parent, title)
        card.pack(fill="both", expand=True, padx=16, pady=16)

        ttk.Label(frame, text="Saved Records", style="SectionTitle.TLabel").pack(anchor="w", pady=(0, 8))
        tree_frame = ttk.Frame(frame, style="Panel.TFrame")
        tree_frame.pack(fill="both", expand=True)

        tree = ttk.Treeview(
            tree_frame,
            columns=("client", "site"),
            show="headings",
        )
        tree.heading("client", text="Client")
        tree.heading("site", text="Site")
        tree.column("client", width=340, anchor="w")
        tree.column("site", width=440, anchor="w")
        tree.column("client", stretch=True)
        tree.column("site", stretch=True)
        master_scroll = ttk.Scrollbar(
            tree_frame,
            orient="vertical",
            command=tree.yview,
            style="ClientSites.Vertical.TScrollbar",
        )
        tree.configure(yscrollcommand=master_scroll.set)
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        tree.grid(row=0, column=0, sticky="nsew")
        master_scroll.grid(row=0, column=1, sticky="ns")
        self._bind_mousewheel_lines(tree)
        tree.bind("<Double-1>", lambda _e, sheet=sheet_name: self._update_master_entry(sheet))

        btns = ttk.Frame(frame, style="Panel.TFrame")
        btns.pack(fill="x", pady=4)
        ttk.Button(btns, text="Add New", style="Accent.TButton", command=lambda sheet=sheet_name: self._add_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Edit", command=lambda sheet=sheet_name: self._update_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Delete", command=lambda sheet=sheet_name: self._delete_master_entry(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Reload", command=lambda sheet=sheet_name: self._reload_master_tab(sheet)).pack(side="left", padx=(0, 6))
        ttk.Button(btns, text="Undo Last Change", command=self._undo_last_change).pack(side="left", padx=(0, 6))
        if sheet_name != COMPLETED_SHEET:
            ttk.Button(btns, text="Move To Completed", command=lambda sheet=sheet_name: self._move_selected_master_entry_to_completed(sheet)).pack(side="left")

        self.client_site_tabs[sheet_name] = {
            "tree": tree,
        }

        if sheet_name == MASTER_SHEET:
            self.master_tree = tree

        self._refresh_master_tree(sheet_name)

    def _build_cc_tab(self):
        outer = ttk.Frame(self.cc_tab)
        outer.pack(fill="both", expand=True, padx=16, pady=16)

        left_card, left = self._create_card(outer, "Internal CC Lists")
        left_card.pack(side="left", fill="y", padx=(0, 16))

        right_card, right = self._create_card(outer, "List Members")
        right_card.pack(side="left", fill="both", expand=True)

        ttk.Label(left, text="CC Lists", style="SectionTitle.TLabel").pack(anchor="w")
        self.cc_template_label = ttk.Label(left, text="")
        self.cc_template_label.pack(anchor="w", pady=(0, 8))
        cc_list_frame = ttk.Frame(left, style="Panel.TFrame")
        cc_list_frame.pack(anchor="w", pady=(4, 8), fill="y")
        self.cc_listbox = tk.Listbox(cc_list_frame, height=20, width=28)
        cc_list_scroll = ttk.Scrollbar(cc_list_frame, orient="vertical", command=self.cc_listbox.yview)
        self.cc_listbox.configure(yscrollcommand=cc_list_scroll.set)
        self.cc_listbox.pack(side="left", fill="y")
        cc_list_scroll.pack(side="right", fill="y")
        self._style_listbox_widget(self.cc_listbox)
        self.cc_listbox.bind("<<ListboxSelect>>", lambda _e: self._load_cc_list_selection())

        self.new_cc_list_var = tk.StringVar()
        ttk.Entry(left, textvariable=self.new_cc_list_var, width=28).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Create List", style="Accent.TButton", command=self._create_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Edit List", command=self._edit_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Delete List", command=self._delete_cc_list).pack(anchor="w", pady=(0, 6))
        ttk.Button(left, text="Reload", command=self._reload_cc_tab).pack(anchor="w")

        ttk.Label(right, text="Emails in Selected List", style="SectionTitle.TLabel").pack(anchor="w")
        cc_emails_frame = ttk.Frame(right, style="Panel.TFrame")
        cc_emails_frame.pack(anchor="w", pady=(4, 8), fill="both", expand=True)
        self.cc_emails_listbox = tk.Listbox(cc_emails_frame, height=20)
        cc_emails_scroll = ttk.Scrollbar(cc_emails_frame, orient="vertical", command=self.cc_emails_listbox.yview)
        self.cc_emails_listbox.configure(yscrollcommand=cc_emails_scroll.set)
        self.cc_emails_listbox.pack(side="left", fill="both", expand=True)
        cc_emails_scroll.pack(side="right", fill="y")
        self._style_listbox_widget(self.cc_emails_listbox)

        email_controls = ttk.Frame(right, style="Panel.TFrame")
        email_controls.pack(fill="x")
        self.cc_email_var = tk.StringVar()
        ttk.Entry(email_controls, textvariable=self.cc_email_var, width=50).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Add Email", style="Accent.TButton", command=self._add_cc_email).pack(side="left", padx=(0, 6))
        ttk.Button(email_controls, text="Remove Selected", command=self._remove_cc_email).pack(side="left")

        self._refresh_cc_listbox()

    def _build_about_tab(self):
        card, frame = self._create_card(self.about_tab, "About")
        card.pack(fill="both", expand=True, padx=16, pady=16)

        ttk.Label(
            frame,
            text="Created on: February 4th, 2026\nBy: Steven McLaren, Data Technician",
            anchor="center",
            justify="center",
            style="CardTitle.TLabel",
        ).pack(expand=True, pady=24)

    def _show_tab(self, tab_name: str):
        frame = self.tabs.get(tab_name)
        if not frame:
            return
        for name, tab_frame in self.tabs.items():
            if name == tab_name:
                tab_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
                tab_frame.lift()
            else:
                tab_frame.place_forget()
        for name, button in self.nav_buttons.items():
            button.configure(style="ActiveNav.TButton" if name == tab_name else "Nav.TButton")

    def _refresh_clients(self):
        clients = [""] + sorted({row["client"] for row in self._compose_dropbox_data()})
        self.client_combo["values"] = clients
        current = self.client_var.get().strip()
        if current in clients:
            self.client_var.set(current)
        else:
            self.client_var.set("")
        self._on_client_change()
        self._refresh_cc_lists()
        self._refresh_signature_list()

    def _refresh_cc_lists(self):
        list_names = sorted(self.cc_lists.keys())
        self.cc_list_combo["values"] = list_names
        default_name = "Commercial Construction"
        current = self.cc_list_var.get().strip()
        if current and current in self.cc_lists:
            self.cc_list_var.set(current)
        elif default_name in self.cc_lists:
            self.cc_list_var.set(default_name)
        elif list_names:
            self.cc_list_var.set(list_names[0])
        else:
            self.cc_list_var.set("")
        self._on_cc_list_change()

    def _on_cc_list_change(self):
        list_name = self.cc_list_var.get()
        emails = self.cc_lists.get(list_name, [])
        self._set_compose_cc_text("; ".join(emails))
        next_sheet = self._sheet_for_cc_list(list_name)
        if next_sheet != self.compose_sheet_name:
            self.compose_sheet_name = next_sheet
            self._refresh_clients()
        else:
            self._on_client_change()

    def _refresh_signature_list(self):
        if not hasattr(self, "signature_combo"):
            return
        names = list(self.email_templates.keys())
        names = sorted(names)
        default_names = ["Main", "Main (smclaren@archaerial.com)"]
        current = self.signature_var.get().strip() if hasattr(self, "signature_var") else ""
        self.signature_combo["values"] = names
        if current and current in names:
            self.signature_var.set(current)
        elif default_names[0] in names:
            self.signature_var.set(default_names[0])
        elif default_names[1] in names:
            self.signature_var.set(default_names[1])
        elif names:
            self.signature_var.set(names[0])
        else:
            self.signature_var.set("")

    def _template_context(self):
        return {
            "client": self.client_var.get().strip(),
            "site": self.site_var.get().strip(),
            "media_subject": MEDIA_SUBJECT.get(self.media_var.get(), self.media_var.get()),
            "media_body": MEDIA_BODY.get(self.media_var.get(), self.media_var.get()),
            "date": self.date_var.get().strip(),
            "dropbox_link": self.link_var.get().strip(),
            "to_email": self.to_var.get().strip(),
            "cc_email": self.cc_var.get().strip(),
        }

    def _render_template_text(self, template_text, ctx):
        if not template_text:
            return ""

        def _replace(m):
            key = m.group(1)
            return str(ctx.get(key, m.group(0)))

        return re.sub(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", _replace, template_text)

    def _trim_leading_main_block(self, template_name: str, html: str) -> str:
        if not html:
            return html
        marker = "Our pilots captured {media_body} of the {site} site on {date}."
        if marker not in html:
            return html
        meta_idx = html.find("<meta")
        if meta_idx == -1:
            return html
        next_div = html.find("<div", meta_idx)
        if next_div == -1:
            return html
        return html[next_div:]

    def _apply_legacy_placeholders(self, content: str, ctx) -> str:
        if not content:
            return content
        text = content
        text = text.replace("February 00, 2026", ctx.get("date", ""))
        text = re.sub(r"\bCLIENT\b", ctx.get("client", ""), text)
        text = re.sub(r"\bSITE\b", ctx.get("site", ""), text)
        return text

    def _html_to_plain_text(self, html: str) -> str:
        text = re.sub(r"(?is)<(script|style).*?</\1>", "", html or "")
        text = re.sub(r"<\s*/\s*p\s*>", "\n\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<\s*br\s*/?>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        try:
            from html import unescape
            text = unescape(text)
        except Exception:
            pass
        return text.strip()

    def _refresh_template_folder_selector(self):
        if not hasattr(self, "template_folder_frame"):
            return

        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()

        for child in self.template_folder_frame.winfo_children():
            child.destroy()
        self.template_folder_vars = {}

        if not self.template_folder_names:
            return

        ttk.Label(self.template_folder_frame, text="Templates").pack(side="left", padx=(0, 8))
        for folder_name in self.template_folder_names:
            var = tk.BooleanVar(value=folder_name == self.active_template_folder_name)
            self.template_folder_vars[folder_name] = var
            ttk.Checkbutton(
                self.template_folder_frame,
                text=folder_name,
                variable=var,
                command=lambda name=folder_name: self._on_template_folder_toggle(name),
            ).pack(side="left", padx=(0, 8))

    def _on_template_folder_toggle(self, folder_name):
        selected_var = self.template_folder_vars.get(folder_name)
        if selected_var is None:
            return

        if not selected_var.get():
            selected_var.set(True)
            return

        for name, var in self.template_folder_vars.items():
            var.set(name == folder_name)

        if self.active_template_folder_name == folder_name:
            return

        self.active_template_folder_name = folder_name
        self.cc_lists = self._read_cc_lists()
        self.email_templates = self._read_email_templates()
        self._refresh_cc_listbox()
        self._refresh_cc_lists()
        self._refresh_signature_list()
        self._update_email_preview()

    def _refresh_cc_listbox(self):
        if not hasattr(self, "cc_listbox"):
            return
        if hasattr(self, "cc_template_label"):
            template_name = self.active_template_folder_name or "Default"
            self.cc_template_label.configure(text=f"Template: {template_name}")
        current_selection = None
        selection = self.cc_listbox.curselection()
        if selection:
            current_selection = self.cc_listbox.get(selection[0])
        self.cc_listbox.delete(0, "end")
        names = sorted(self.cc_lists.keys())
        selected_index = None
        for idx, name in enumerate(names):
            self.cc_listbox.insert("end", name)
            if name == current_selection:
                selected_index = idx
        if selected_index is None and names:
            selected_index = 0
        if selected_index is not None:
            self.cc_listbox.selection_set(selected_index)
        self._load_cc_list_selection()

    def _load_cc_list_selection(self):
        if not hasattr(self, "cc_listbox"):
            return
        selection = self.cc_listbox.curselection()
        if not selection:
            self.cc_emails_listbox.delete(0, "end")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        self.cc_emails_listbox.delete(0, "end")
        for email in emails:
            self.cc_emails_listbox.insert("end", email)

    def _create_cc_list(self):
        name = self.new_cc_list_var.get().strip()
        if not name:
            messagebox.showwarning("Missing Data", "List name is required.")
            return
        if name in self.cc_lists:
            messagebox.showwarning("Already Exists", "That list already exists.")
            return
        self.cc_lists[name] = []
        self._write_cc_lists()
        self.new_cc_list_var.set("")
        self._reload_cc_lists_data()

    def _parse_cc_email_text(self, raw) -> list[str]:
        normalized = normalize_email_list(raw)
        return [email.strip() for email in normalized.split(",") if email.strip()]

    def _edit_cc_list(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to edit.")
            return
        original_name = self.cc_listbox.get(selection[0])
        self._open_cc_list_window(original_name)

    def _open_cc_list_window(self, original_name: str):
        dialog = tk.Toplevel(self.root)
        dialog.title("Edit CC List")
        dialog.geometry("720x500")
        dialog.minsize(620, 440)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.configure(bg=self.colors["bg"])

        shell = ttk.Frame(dialog, style="Panel.TFrame")
        shell.pack(fill="both", expand=True, padx=18, pady=18)
        shell.grid_columnconfigure(1, weight=1)
        shell.grid_rowconfigure(1, weight=1)

        ttk.Label(shell, text="List Name", style="SectionTitle.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        name_var = tk.StringVar(value=original_name)
        name_entry = ttk.Entry(shell, textvariable=name_var)
        name_entry.grid(row=0, column=1, sticky="ew", pady=(0, 10))

        ttk.Label(shell, text="Emails", style="SectionTitle.TLabel").grid(row=1, column=0, sticky="nw", padx=(0, 12), pady=(0, 10))
        email_frame = ttk.Frame(shell, style="Panel.TFrame")
        email_frame.grid(row=1, column=1, sticky="nsew", pady=(0, 10))
        email_frame.grid_rowconfigure(0, weight=1)
        email_frame.grid_columnconfigure(0, weight=1)
        email_text = tk.Text(email_frame, height=10, wrap="word")
        email_scroll = ttk.Scrollbar(email_frame, orient="vertical", command=email_text.yview)
        email_text.configure(yscrollcommand=email_scroll.set)
        email_text.grid(row=0, column=0, sticky="nsew")
        email_scroll.grid(row=0, column=1, sticky="ns")
        self._style_text_widget(email_text)
        email_text.insert("1.0", ", ".join(self.cc_lists.get(original_name, [])))

        def paste_email_list(event):
            try:
                raw = dialog.clipboard_get()
            except tk.TclError:
                return None
            normalized = normalize_email_list(raw)
            if not normalized:
                return "break"
            email_text.insert("insert", normalized)
            return "break"

        email_text.bind("<Control-v>", paste_email_list)
        email_text.bind("<<Paste>>", paste_email_list)

        btns = ttk.Frame(shell, style="Panel.TFrame")
        btns.grid(row=2, column=0, columnspan=2, sticky="e", pady=(8, 0))

        def save():
            new_name = name_var.get().strip()
            if not new_name:
                messagebox.showwarning("Missing Data", "List name is required.", parent=dialog)
                return
            if new_name != original_name and new_name in self.cc_lists:
                messagebox.showwarning("Already Exists", "That list already exists.", parent=dialog)
                return
            emails = self._parse_cc_email_text(email_text.get("1.0", "end-1c"))
            updated_lists = {}
            for list_name, existing_emails in self.cc_lists.items():
                if list_name == original_name:
                    updated_lists[new_name] = emails
                else:
                    updated_lists[list_name] = existing_emails
            self.cc_lists = updated_lists
            try:
                self._write_cc_lists()
            except Exception as exc:
                messagebox.showerror("CC List Error", f"Failed to save list: {exc}", parent=dialog)
                return
            dialog.destroy()
            self._reload_cc_lists_data()
            self._select_cc_list(new_name)

        ttk.Button(btns, text="Cancel", command=dialog.destroy).pack(side="right", padx=(6, 0))
        ttk.Button(btns, text="Save", style="Accent.TButton", command=save).pack(side="right")
        dialog.bind("<Escape>", lambda _e: dialog.destroy())
        dialog.bind("<Control-s>", lambda _e: save())
        name_entry.focus_set()
        dialog.wait_window()

    def _select_cc_list(self, list_name: str):
        if not hasattr(self, "cc_listbox"):
            return
        for idx in range(self.cc_listbox.size()):
            if self.cc_listbox.get(idx) == list_name:
                self.cc_listbox.selection_clear(0, "end")
                self.cc_listbox.selection_set(idx)
                self.cc_listbox.see(idx)
                self._load_cc_list_selection()
                return

    def _delete_cc_list(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to delete.")
            return
        list_name = self.cc_listbox.get(selection[0])
        if not messagebox.askyesno("Confirm Delete", f"Delete CC list '{list_name}'?"):
            return
        self.cc_lists.pop(list_name, None)
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _add_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to add to.")
            return
        email = self.cc_email_var.get().strip()
        if not email:
            messagebox.showwarning("Missing Data", "Email is required.")
            return
        list_name = self.cc_listbox.get(selection[0])
        emails = self.cc_lists.get(list_name, [])
        if email in emails:
            messagebox.showwarning("Already Exists", "That email is already in the list.")
            return
        emails.append(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self.cc_email_var.set("")
        self._reload_cc_lists_data()

    def _remove_cc_email(self):
        selection = self.cc_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select List", "Choose a list to remove from.")
            return
        list_name = self.cc_listbox.get(selection[0])
        email_selection = self.cc_emails_listbox.curselection()
        if not email_selection:
            messagebox.showwarning("Select Email", "Choose an email to remove.")
            return
        emails = self.cc_lists.get(list_name, [])
        for idx in reversed(email_selection):
            email = self.cc_emails_listbox.get(idx)
            if email in emails:
                emails.remove(email)
        self.cc_lists[list_name] = emails
        self._write_cc_lists()
        self._reload_cc_lists_data()

    def _reload_cc_lists_data(self):
        self.cc_lists = self._read_cc_lists()
        self._refresh_cc_listbox()
        self._refresh_cc_lists()

    def _reload_compose_tab(self):
        # Reload only Compose-related data
        self.signature_names = self._read_signature_names()
        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()
        self.cc_lists = self._read_cc_lists()
        self.email_templates = self._read_email_templates()
        self._refresh_template_folder_selector()
        # Refresh data sources for client/site/email/link without touching master tab UI
        self._reload_master_data(sheet_name=MASTER_SHEET, refresh_clients=True, refresh_master_tree=False)
        self._reload_master_data(sheet_name=OIL_GAS_SHEET, refresh_clients=False, refresh_master_tree=False)
        self._refresh_cc_lists()
        self._reset_compose_state(refresh_clients=True, force_blank_client_site=True)

    def _reload_master_tab(self, sheet_name: str = MASTER_SHEET):
        # Reload only the requested client-sites tab data
        self._reload_master_data(
            sheet_name=sheet_name,
            refresh_clients=(sheet_name == self.compose_sheet_name),
            refresh_master_tree=True,
        )
        self._reset_master_form(sheet_name)

    def _reload_cc_tab(self):
        # Reload only Internal CC Lists tab data
        self.template_folder_names = self._read_template_folder_names()
        self.active_template_folder_name = self._choose_active_template_folder()
        self._reload_cc_lists_data()
        self._reset_cc_tab()

    def _on_client_change(self):
        client = self.client_var.get()
        compose_dropbox_data = self._compose_dropbox_data()
        if not client:
            self.site_combo["values"] = [""]
            self.site_var.set("")
            self.link_var.set("")
            self._set_compose_to_text("")
            return
        sites = [""] + sorted({row["site"] for row in compose_dropbox_data if row["client"] == client})
        self.site_combo["values"] = sites
        current_site = self.site_var.get().strip()
        if current_site in sites:
            self.site_var.set(current_site)
        else:
            self.site_var.set("")
        self._on_site_change()

    def _on_site_change(self):
        if not self.client_var.get().strip():
            self.link_var.set("")
            self._set_compose_to_text("")
            return
        self.link_var.set(self._lookup_compose_field("link"))
        self._set_compose_to_text(self._lookup_compose_field("email"))

    def _handle_post_draft_procore(self):
        procore_link = self._lookup_compose_field("procore")
        if not procore_link:
            return
        try:
            opened = webbrowser.open(procore_link, new=2)
            if not opened:
                raise RuntimeError("No browser handler is available.")
        except Exception:
            messagebox.showinfo(
                "Procore Link",
                f"The draft was created.\n\nOpen this Procore link for the selected site:\n{procore_link}",
            )

    def _reset_compose_state(self, refresh_clients: bool = True, force_blank_client_site: bool = False):
        if not hasattr(self, "client_var"):
            return
        self.media_var.set(MEDIA_OPTIONS[1])
        self.subject_var.set("")
        self._set_compose_to_text("")
        self._set_compose_cc_text("")
        self.link_var.set("")
        self._set_default_date()
        if force_blank_client_site:
            self.client_var.set("")
            self.site_var.set("")
            self._set_compose_to_text("")
            self.link_var.set("")
        if refresh_clients:
            self._refresh_clients()
        self._update_email_preview()

    def _reset_master_form(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        tree = widgets["tree"]
        tree.selection_remove(tree.selection())

    def _reset_cc_tab(self):
        if hasattr(self, "new_cc_list_var"):
            self.new_cc_list_var.set("")
        if hasattr(self, "cc_email_var"):
            self.cc_email_var.set("")
        if hasattr(self, "cc_listbox"):
            self.cc_listbox.selection_clear(0, "end")
        if hasattr(self, "cc_emails_listbox"):
            self.cc_emails_listbox.delete(0, "end")

    def _set_default_date(self):
        self.selected_date = default_date()
        self.date_var.set(format_date(self.selected_date))

    def _open_calendar(self):
        CalendarPopup(self.root, self.selected_date, self._set_date)

    def _set_date(self, date_obj):
        self.selected_date = date_obj
        self.date_var.set(format_date(date_obj))
        self._update_email_preview()

    def _paste_email_list(self, event, entry, variable):
        try:
            raw = self.root.clipboard_get()
        except tk.TclError:
            return None

        normalized = normalize_email_list(raw)
        if not normalized:
            return "break"

        entry.insert("insert", normalized)
        variable.set(normalize_email_list(variable.get()))
        return "break"

    def _get_compose_cc_text(self) -> str:
        if not hasattr(self, "cc_text"):
            return self.cc_var.get()
        return self.cc_text.get("1.0", "end-1c")

    def _get_compose_to_text(self) -> str:
        if not hasattr(self, "to_text"):
            return self.to_var.get()
        return self.to_text.get("1.0", "end-1c")

    def _set_compose_to_text(self, value: str):
        normalized = normalize_email_list(value)
        self.to_var.set(normalized)
        if hasattr(self, "to_text"):
            self.to_text.delete("1.0", "end")
            if normalized:
                self.to_text.insert("1.0", normalized)

    def _sync_compose_to_var(self, _event=None):
        self.to_var.set(normalize_email_list(self._get_compose_to_text()))

    def _expand_compose_to_editor(self, _event=None):
        if not hasattr(self, "to_text"):
            return
        self.to_text.configure(height=4)
        self.to_scroll.pack(side="right", fill="y")

    def _collapse_compose_to_editor(self, _event=None):
        if not hasattr(self, "to_text"):
            return
        self._set_compose_to_text(self._get_compose_to_text())
        self.to_text.configure(height=1)
        self.to_scroll.pack_forget()

    def _paste_compose_to_list(self, _event=None):
        try:
            raw = self.root.clipboard_get()
        except tk.TclError:
            return None

        normalized = normalize_email_list(raw)
        if not normalized:
            return "break"

        self.to_text.insert("insert", normalized)
        self._set_compose_to_text(self._get_compose_to_text())
        return "break"

    def _set_compose_cc_text(self, value: str):
        normalized = normalize_email_list(value)
        self.cc_var.set(normalized)
        if hasattr(self, "cc_text"):
            self.cc_text.delete("1.0", "end")
            if normalized:
                self.cc_text.insert("1.0", normalized)

    def _sync_compose_cc_var(self, _event=None):
        self.cc_var.set(normalize_email_list(self._get_compose_cc_text()))

    def _expand_compose_cc_editor(self, _event=None):
        if not hasattr(self, "cc_text"):
            return
        self.cc_text.configure(height=4)
        self.cc_scroll.pack(side="right", fill="y")

    def _collapse_compose_cc_editor(self, _event=None):
        if not hasattr(self, "cc_text"):
            return
        self._set_compose_cc_text(self._get_compose_cc_text())
        self.cc_text.configure(height=1)
        self.cc_scroll.pack_forget()

    def _paste_compose_cc_list(self, _event=None):
        try:
            raw = self.root.clipboard_get()
        except tk.TclError:
            return None

        normalized = normalize_email_list(raw)
        if not normalized:
            return "break"

        self.cc_text.insert("insert", normalized)
        self._set_compose_cc_text(self._get_compose_cc_text())
        return "break"

    def _compose_subject(self):
        ctx = self._template_context()
        template = self.email_templates.get(self.signature_var.get().strip())
        if template and template.get("subject"):
            rendered = self._render_template_text(template["subject"], ctx)
            if rendered:
                return rendered
        if ctx["site"] and ctx["media_subject"] and ctx["date"]:
            return f"{ctx['site']} - {ctx['media_subject']} - {ctx['date']}"
        return ""

    def _compose_body(self):
        ctx = self._template_context()
        selected_name = self.signature_var.get().strip()
        template = self.email_templates.get(selected_name)
        if template:
            if template.get("html"):
                rendered = self._trim_leading_main_block(selected_name, template["html"])
                rendered = self._render_template_text(rendered, ctx)
                rendered = self._apply_legacy_placeholders(rendered, ctx)
                rendered = self._html_to_plain_text(rendered)
                if rendered:
                    return rendered.replace("\n", "\r\n")
            if template.get("text"):
                rendered = self._render_template_text(template["text"], ctx)
                if rendered:
                    return rendered.replace("\n", "\r\n")
        greeting = f"Good Day {ctx['client']} Team,"
        main = (
            f"Our pilots captured {ctx['media_body']} of the {ctx['site']} site on {ctx['date']}. "
            "Please use the link below to access your imagery."
        )
        body = (
            f"{greeting}\r\n\r\n"
            f"{main}\r\n"
            f"Dropbox Link: {ctx['dropbox_link']}\r\n\r\n"
            "Let us know if you have any questions, feedback, or if we can further assist you.\r\n\r\n"
        )
        return body

    def _compose_body_html(self):
        ctx = self._template_context()
        selected_name = self.signature_var.get().strip()
        template = self.email_templates.get(selected_name)
        if template and (template.get("html") or template.get("text")):
            def _ensure_assets(html):
                return self._extract_inline_images(html, base_dir=template.get("path").parent if template.get("path") else None)

            def _text_to_html(text):
                from html import escape
                link = ctx["dropbox_link"]
                paragraphs = []
                current = []
                for line in text.splitlines():
                    if line.strip() == "":
                        if current:
                            paragraphs.append("<br>".join(current))
                            current = []
                        continue
                    escaped = escape(line)
                    if "Dropbox Link:" in line and link:
                        escaped = escaped.replace(escape(link), f'<a href="{escape(link)}">{escape(link)}</a>')
                    current.append(escaped)
                if current:
                    paragraphs.append("<br>".join(current))
                html_parts = [f"<p style=\"margin:0 0 10px 0;\">{p}</p>" for p in paragraphs]
                return (
                    "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; line-height:1.35;\">"
                    + "".join(html_parts)
                    + "</div>"
                )

            if template.get("html"):
                rendered = self._trim_leading_main_block(selected_name, template["html"])
                rendered = self._render_template_text(rendered, ctx)
                rendered = self._apply_legacy_placeholders(rendered, ctx)
                if ctx["dropbox_link"]:
                    rendered = rendered.replace(
                        ctx["dropbox_link"],
                        f'<a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>',
                    )
                    rendered = re.sub(
                        r"(<b>\s*Dropbox Link:\s*</b>)(?!\s*<a\b)",
                        rf'\1 <a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>',
                        rendered,
                        flags=re.IGNORECASE,
                    )
                body_html, template_assets = _ensure_assets(rendered)
                return body_html, template_assets

            if template.get("text"):
                rendered = self._render_template_text(template["text"], ctx)
                body_html = _text_to_html(rendered)
                return body_html, []
        link_html = f'<a href="{ctx["dropbox_link"]}">{ctx["dropbox_link"]}</a>' if ctx["dropbox_link"] else ""

        body_html = (
            "<div style=\"font-family:Arial,Helvetica,sans-serif; font-size:11pt; color:#000; "
            "line-height:1.35;\">"
            f"<p style=\"margin:0 0 10px 0;\">Good Day {ctx['client']} Team,</p>"
            f"<p style=\"margin:0 0 8px 0;\">Our pilots captured {ctx['media_body']} of the {ctx['site']} site on {ctx['date']}. "
            "Please use the link below to access your imagery.</p>"
            f"<p style=\"margin:0 0 10px 0;\"><b>Dropbox Link:</b> {link_html}</p>"
            "<p style=\"margin:0 0 12px 0;\">Let us know if you have any questions, feedback, or if we can further assist you.</p>"
            ""
            "</div>"
        )
        return body_html, []

    def _update_email_preview(self):
        subject = self._compose_subject()
        body = self._compose_body()
        if subject:
            self.subject_var.set(subject)
        self.preview.config(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("1.0", body)
        self.preview.config(state="disabled")

    def _parse_recipients(self, raw: str):
        normalized = normalize_email_list(raw)
        tokens = re.split(r"[;,]+", normalized)
        return [t.strip() for t in tokens if t.strip()]

    def _create_draft_graph(self, subject: str, body_html: str, signature_assets):
        if not self.graph_client:
            raise RuntimeError("Graph client is not configured.")

        to_list = self._parse_recipients(self.to_var.get().strip())
        cc_list = self._parse_recipients(self.cc_var.get().strip())

        if not subject or not to_list:
            raise ValueError("Please ensure a client, site, and To address are selected.")

        inline_attachments = []
        for asset in signature_assets or []:
            if not asset.get("path"):
                continue
            inline_attachments.append(
                {"path": asset["path"], "cid": asset.get("cid") or Path(asset["path"]).stem}
            )

        self.graph_client.create_draft(
            subject=subject,
            body_html=body_html,
            to=to_list,
            cc=cc_list,
            bcc=[],
            attachments=[],
            inline_attachments=inline_attachments,
        )

    def _create_draft(self):
        subject = self.subject_var.get().strip()
        try:
            body_html, signature_assets = self._compose_body_html()
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to build email HTML: {exc}")
            return

        if self.graph_client:
            try:
                self._create_draft_graph(subject, body_html, signature_assets)
                self._handle_post_draft_procore()
                return
            except Exception as exc:
                messagebox.showerror("Graph Error", f"Failed to create Graph draft: {exc}")
                return

        if win32 is None:
            messagebox.showerror(
                "Outlook",
                "Graph is not configured and pywin32 is not available. "
                "Install pywin32 or configure Graph.",
            )
            return

        to_list = self._parse_recipients(self.to_var.get().strip())
        cc_list = self._parse_recipients(self.cc_var.get().strip())
        if not subject or not to_list:
            messagebox.showwarning("Missing Data", "Please ensure a client, site, and To address are selected.")
            return

        try:
            outlook = win32.Dispatch("Outlook.Application")
            mail = outlook.CreateItem(0)
            for addr in to_list:
                recipient = mail.Recipients.Add(addr)
                recipient.Type = 1
            for addr in cc_list:
                recipient = mail.Recipients.Add(addr)
                recipient.Type = 2
            mail.Subject = subject
            mail.BodyFormat = 2  # olFormatHTML
            mail.HTMLBody = body_html
            for asset in signature_assets:
                try:
                    attachment = mail.Attachments.Add(asset["path"])
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3712001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x3713001F",
                        asset["cid"],
                    )
                    attachment.PropertyAccessor.SetProperty(
                        "http://schemas.microsoft.com/mapi/proptag/0x7FFE000B",
                        True,
                    )
                except Exception as exc:
                    continue
            # Re-apply HTMLBody after inline attachments to help Outlook resolve CIDs
            mail.HTMLBody = body_html
            mail.Save()  # draft in Drafts
            try:
                self._sync_outlook_now(outlook)
            except Exception:
                pass
            self._handle_post_draft_procore()
        except Exception as exc:
            messagebox.showerror("Outlook Error", f"Failed to create Outlook draft: {exc}")

    def _refresh_dropbox_tree(self):
        self.dropbox_tree.delete(*self.dropbox_tree.get_children())
        for row in self.dropbox_data:
            self.dropbox_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["link"]))

    def _refresh_master_tree(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        tree = widgets["tree"]
        tree.delete(*tree.get_children())
        for row in self.sheet_data.get(sheet_name, []):
            tree.insert(
                "",
                "end",
                iid=str(row["row"]),
                values=(row["client"], row["site"]),
            )

    def _load_dropbox_selection(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.dropbox_data if r["row"] == row_id), None)
        if not row:
            return
        self.db_client_var.set(row["client"])
        self.db_site_var.set(row["site"])
        self.db_link_var.set(row["link"])

    def _load_master_selection(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            return
        self._update_master_entry(sheet_name)

    def _add_master_entry(self, sheet_name: str = MASTER_SHEET):
        self._open_master_entry_window(sheet_name)

    def _update_master_entry(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to edit.")
            return
        row_id = int(selected[0])
        row = next((r for r in self.sheet_data.get(sheet_name, []) if r["row"] == row_id), None)
        if not row:
            messagebox.showwarning("Select Row", "The selected row could not be found.")
            return
        self._open_master_entry_window(sheet_name, row)

    def _open_master_entry_window(self, sheet_name: str = MASTER_SHEET, row: dict | None = None):
        is_edit = row is not None
        title = "Edit Client Data" if is_edit else "Add New Client Data"
        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.geometry("720x570")
        dialog.minsize(620, 500)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.configure(bg=self.colors["bg"])

        shell = ttk.Frame(dialog, style="Panel.TFrame")
        shell.pack(fill="both", expand=True, padx=18, pady=18)
        shell.grid_columnconfigure(1, weight=1)
        shell.grid_rowconfigure(2, weight=1)

        ttk.Label(shell, text="Client", style="SectionTitle.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        client_var = tk.StringVar(value=row["client"] if row else "")
        client_entry = ttk.Entry(shell, textvariable=client_var)
        client_entry.grid(row=0, column=1, sticky="ew", pady=(0, 10))

        ttk.Label(shell, text="Site", style="SectionTitle.TLabel").grid(row=1, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        site_var = tk.StringVar(value=row["site"] if row else "")
        ttk.Entry(shell, textvariable=site_var).grid(row=1, column=1, sticky="ew", pady=(0, 10))

        ttk.Label(shell, text="Operations Contact Emails", style="SectionTitle.TLabel").grid(row=2, column=0, sticky="nw", padx=(0, 12), pady=(0, 10))
        email_frame = ttk.Frame(shell, style="Panel.TFrame")
        email_frame.grid(row=2, column=1, sticky="nsew", pady=(0, 10))
        email_frame.grid_rowconfigure(0, weight=1)
        email_frame.grid_columnconfigure(0, weight=1)
        email_text = tk.Text(email_frame, height=7, wrap="word")
        email_scroll = ttk.Scrollbar(email_frame, orient="vertical", command=email_text.yview)
        email_text.configure(yscrollcommand=email_scroll.set)
        email_text.grid(row=0, column=0, sticky="nsew")
        email_scroll.grid(row=0, column=1, sticky="ns")
        self._style_text_widget(email_text)
        if row and row.get("email"):
            email_text.insert("1.0", row["email"])

        def paste_email_list(event):
            try:
                raw = dialog.clipboard_get()
            except tk.TclError:
                return None
            normalized = normalize_email_list(raw)
            if not normalized:
                return "break"
            email_text.insert("insert", normalized)
            return "break"

        email_text.bind("<Control-v>", paste_email_list)
        email_text.bind("<<Paste>>", paste_email_list)

        ttk.Label(shell, text="Customer Dropbox URL", style="SectionTitle.TLabel").grid(row=3, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        link_var = tk.StringVar(value=row["link"] if row else "")
        ttk.Entry(shell, textvariable=link_var).grid(row=3, column=1, sticky="ew", pady=(0, 10))

        ttk.Label(shell, text="Procore", style="SectionTitle.TLabel").grid(row=4, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        procore_var = tk.StringVar(value=row.get("procore", "") if row else "")
        ttk.Entry(shell, textvariable=procore_var).grid(row=4, column=1, sticky="ew", pady=(0, 10))

        ttk.Label(shell, text="Internal Dropbox URL", style="SectionTitle.TLabel").grid(row=5, column=0, sticky="w", padx=(0, 12), pady=(0, 10))
        internal_url_var = tk.StringVar(value=row.get("internal_url", "") if row else "")
        ttk.Entry(shell, textvariable=internal_url_var).grid(row=5, column=1, sticky="ew", pady=(0, 10))

        btns = ttk.Frame(shell, style="Panel.TFrame")
        btns.grid(row=6, column=0, columnspan=2, sticky="ew", pady=(8, 0))

        def save():
            values = {
                "client": client_var.get().strip(),
                "site": site_var.get().strip(),
                "email": normalize_email_list(email_text.get("1.0", "end-1c")),
                "link": link_var.get().strip(),
                "procore": procore_var.get().strip(),
                "internal_url": internal_url_var.get().strip(),
            }
            if not values["client"]:
                messagebox.showwarning("Missing Data", "Client is required.", parent=dialog)
                return
            try:
                self._save_master_entry(sheet_name, values, row["row"] if row else None)
            except Exception as exc:
                messagebox.showerror("Master Data Error", f"Failed to save entry: {exc}", parent=dialog)
                return
            dialog.destroy()

        ttk.Button(btns, text="Cancel", command=dialog.destroy).pack(side="right", padx=(6, 0))
        ttk.Button(btns, text="Save", style="Accent.TButton", command=save).pack(side="right")
        dialog.bind("<Escape>", lambda _e: dialog.destroy())
        dialog.bind("<Control-s>", lambda _e: save())
        client_entry.focus_set()
        dialog.wait_window()

    def _save_master_entry(self, sheet_name: str, values: dict, row_id: int | None = None):
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            if row_id is None:
                ws = self._get_or_create_sheet(wb, sheet_name)
                ws.append(
                    self._build_master_row(
                        ws,
                        {
                            "client": values["client"],
                            "project name": values["site"],
                            "dropbox urls": values["link"],
                            "operations contact email": values["email"],
                            "procore": values["procore"],
                            "internal urls": values["internal_url"],
                        },
                    )
                )
            else:
                ws = wb[self._resolve_sheet_name(wb, sheet_name)]
                col_map, _ = self._master_column_map(ws, ensure=True)
                ws.cell(row=row_id, column=col_map["client"], value=values["client"])
                ws.cell(row=row_id, column=col_map["project name"], value=values["site"])
                ws.cell(row=row_id, column=col_map["dropbox urls"], value=values["link"])
                ws.cell(row=row_id, column=col_map["operations contact email"], value=values["email"])
                ws.cell(row=row_id, column=col_map["procore"], value=values["procore"])
                ws.cell(row=row_id, column=col_map["internal urls"], value=values["internal_url"])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._reload_master_tab(sheet_name)
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _delete_master_entry(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._reload_master_data(sheet_name=sheet_name)
        except Exception as exc:
            messagebox.showerror("Master Data Error", f"Failed to delete entry: {exc}")

    def _move_master_entry_to_completed(self, sheet_name: str, row_id: int):
        wb = None
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            source_ws = wb[self._resolve_sheet_name(wb, sheet_name)]
            completed_ws = self._get_or_create_sheet(wb, COMPLETED_SHEET)
            row_values = [
                source_ws.cell(row=row_id, column=col_idx).value
                for col_idx in range(1, source_ws.max_column + 1)
            ]
            completed_ws.append(row_values)
            source_ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet(sheet_name)
            self._sort_master_sheet(COMPLETED_SHEET)
            self._reload_master_data(sheet_name=sheet_name, refresh_clients=(sheet_name == self.compose_sheet_name))
            self._reload_master_data(sheet_name=COMPLETED_SHEET, refresh_clients=False)
            if hasattr(self, "status_var"):
                self.status_var.set(f"Moved entry to {COMPLETED_SHEET}.")
        finally:
            try:
                if wb:
                    wb.close()
            except Exception:
                pass

    def _move_selected_master_entry_to_completed(self, sheet_name: str = MASTER_SHEET):
        widgets = self.client_site_tabs.get(sheet_name)
        if not widgets:
            return
        selected = widgets["tree"].selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to move.")
            return
        row_id = int(selected[0])
        try:
            self._move_master_entry_to_completed(sheet_name, row_id)
        except Exception as exc:
            messagebox.showerror("Master Data Error", f"Failed to move entry: {exc}")

    def _add_dropbox_entry(self):
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.append([client, site, link, ""])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to add entry: {exc}")

    def _update_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.db_client_var.get().strip()
        site = self.db_site_var.get().strip()
        link = self.db_link_var.get().strip()
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.cell(row=row_id, column=1, value=client)
            ws.cell(row=row_id, column=2, value=site)
            ws.cell(row=row_id, column=3, value=link)
            wb.save(MASTER_FILE)
            wb.close()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to update entry: {exc}")

    def _delete_dropbox_entry(self):
        selected = self.dropbox_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected Dropbox entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._reload_dropbox_data()
        except Exception as exc:
            messagebox.showerror("Dropbox Error", f"Failed to delete entry: {exc}")

    def _reload_dropbox_data(self):
        self._reload_master_data(sheet_name=MASTER_SHEET)

    def _refresh_contacts_tree(self):
        self.contacts_tree.delete(*self.contacts_tree.get_children())
        for row in self.contacts_data:
            self.contacts_tree.insert("", "end", iid=str(row["row"]), values=(row["client"], row["site"], row["email"]))

    def _load_contact_selection(self):
        selected = self.contacts_tree.selection()
        if not selected:
            return
        row_id = int(selected[0])
        row = next((r for r in self.contacts_data if r["row"] == row_id), None)
        if not row:
            return
        self.ct_client_var.set(row["client"])
        self.ct_site_var.set(row["site"])
        self.ct_email_var.set(normalize_email_list(row["email"]))

    def _add_contact_entry(self):
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = normalize_email_list(self.ct_email_var.get())
        if not client or not site:
            messagebox.showwarning("Missing Data", "Client and Site are required.")
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.append([client, site, None, email])
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to add entry: {exc}")

    def _update_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to update.")
            return
        row_id = int(selected[0])
        client = self.ct_client_var.get().strip()
        site = self.ct_site_var.get().strip()
        email = normalize_email_list(self.ct_email_var.get())
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.cell(row=row_id, column=1, value=client)
            ws.cell(row=row_id, column=2, value=site)
            ws.cell(row=row_id, column=4, value=email)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to update entry: {exc}")

    def _delete_contact_entry(self):
        selected = self.contacts_tree.selection()
        if not selected:
            messagebox.showwarning("Select Row", "Choose a row to delete.")
            return
        row_id = int(selected[0])
        if not messagebox.askyesno("Confirm Delete", "Delete selected contact entry?"):
            return
        try:
            self._record_undo(MASTER_FILE)
            wb = openpyxl.load_workbook(MASTER_FILE)
            ws = wb[self._resolve_sheet_name(wb, MASTER_SHEET)]
            ws.delete_rows(row_id, 1)
            wb.save(MASTER_FILE)
            wb.close()
            self._sort_master_sheet()
            self._reload_contacts_data()
        except Exception as exc:
            messagebox.showerror("Contacts Error", f"Failed to delete entry: {exc}")

    def _reload_contacts_data(self):
        self._reload_master_data(sheet_name=MASTER_SHEET)

    def _reload_master_data(self, sheet_name: str = MASTER_SHEET, refresh_clients: bool = True, refresh_master_tree: bool = True):
        try:
            self._sort_master_sheet(sheet_name)
        except Exception:
            pass
        self.sheet_data[sheet_name] = self._read_master_data(sheet_name, allow_missing=(sheet_name in (OIL_GAS_SHEET, COMPLETED_SHEET)))
        if sheet_name == MASTER_SHEET:
            self._sync_primary_master_views()
        if refresh_master_tree:
            self._refresh_master_tree(sheet_name)
        if refresh_clients:
            self._refresh_clients()


if __name__ == "__main__":
    try:
        root = ttk.Window()
    except Exception:
        root = tk.Tk()
    app = EmailTemplateApp(root)
    root.mainloop()
